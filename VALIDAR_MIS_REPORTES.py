import sys,types,importlib.util,tempfile,unittest,json,io
from pathlib import Path
from datetime import datetime,timedelta
from unittest.mock import patch
from concurrent.futures import ThreadPoolExecutor
from flask import Flask,g
from jinja2 import ChoiceLoader,DictLoader,FileSystemLoader
from openpyxl import load_workbook
root=Path(__file__).resolve().parent
# Isolated database and fake authentication; no application data is accessed.
auth=types.ModuleType('auth_store')
auth.ROLE_ADMIN='DESARROLLADOR';auth.ROLE_CHIEF='JEFE DE DISEÑO';auth.ROLE_COLLABORATOR='DISEÑADOR'
for name in ['create_user','get_user_by_id','list_users','audit_event','reset_password','set_user_status','unlock_user']:setattr(auth,name,lambda *a,**kw:None)
sys.modules['auth_store']=auth
permissions=types.ModuleType('permisos');permissions.MODULES={'proyectos':'Seguimiento','oportunidades':'Proyectos'};sys.modules['permisos']=permissions
pkg=types.ModuleType('reportes');pkg.__path__=[str(root/'reportes')];sys.modules['reportes']=pkg
spec=importlib.util.spec_from_file_location('reportes.db',root/'reportes/db.py');db=importlib.util.module_from_spec(spec);sys.modules['reportes.db']=db;spec.loader.exec_module(db)
from reportes import routes
class Tests(unittest.TestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup);db.DB_FILE=Path(self.tmp.name)/'reports.db'
  self.clock=patch.object(db,'now',return_value=datetime(2026,10,2,16,tzinfo=db.BOGOTA));self.clock.start();self.addCleanup(self.clock.stop)
  self.work=patch.object(db,'is_workday',side_effect=lambda day,p:day.weekday() in p['weekdays']);self.work.start();self.addCleanup(self.work.stop)
  db.ensure_database();self.user={'id':'u','name':'Diseñador','role':auth.ROLE_COLLABORATOR};self.other={'id':'other','name':'Otro'};self.chief={'id':'chief','is_admin':True}
  config={'start':'06:00','end':'15:00','breaks':'12:00-13:00','minimum_reports':2,'weekdays':[0,1,2,3,4],'modules':[]}
  db.save_profile('u','chief','',config,self.chief);db.save_profile('other','chief','',config,self.chief)
  with db.connection() as c:c.execute("UPDATE profile_history SET effective_date='2026-10-01'")
  self.raw={'day':'2026-10-02','start':'06:00','end':'09:00','description':'Preparar modelo tridimensional','project':'3968','equipment':'Tolva','equipment_number':'01','activity':'Modelo 3D','stage':'ETAPA 2','workplace':'NUCLEO'}
 def app(self):
  app=Flask(__name__);app.secret_key='test';app.config['TESTING']=True;app.register_blueprint(routes.bp)
  app.jinja_loader=ChoiceLoader([DictLoader({'reportes/base.html':'{% block report_title %}{% endblock %}{% block report_content %}{% endblock %}{% block scripts %}{% endblock %}'}),FileSystemLoader(str(root/'templates'))])
  @app.before_request
  def current():g.current_user=self.user
  @app.context_processor
  def context():return {'current_user':self.user,'csrf_token':'test'}
  return app
 def test_plan_and_complete_atomic(self):
  task=db.plan_personal_task(self.raw,self.user);self.assertEqual(db.day_summary('u',db.today())['count'],0)
  report=db.submit_report({**self.raw,'pending_id':task},self.user)
  self.assertEqual(db.personal_tasks('u',db.today())[0]['report_id'],report)
  self.assertEqual(db.report(report)['minutes'],180);self.assertEqual(json.loads(db.report(report)['details'])['equipment'],'Tolva')
  with self.assertRaises(ValueError):db.submit_report({**self.raw,'pending_id':task},self.user)
 def test_concurrent_completion_once(self):
  task=db.plan_personal_task(self.raw,self.user)
  def run(_):
   try:db.submit_report({**self.raw,'pending_id':task},self.user);return True
   except ValueError:return False
  with ThreadPoolExecutor(2) as pool:self.assertEqual(sum(pool.map(run,[1,2])),1)
 def test_owner_and_cancel(self):
  task=db.plan_personal_task(self.raw,self.user)
  with self.assertRaises(ValueError):db.cancel_personal_task(task,self.other)
  with self.assertRaises(ValueError):db.submit_report({**self.raw,'pending_id':task},self.other)
  db.cancel_personal_task(task,self.user);self.assertEqual(db.personal_tasks('u',db.today())[0]['status'],'CANCELADA')
 def test_future_plan_not_future_report(self):
  with patch.object(db,'now',return_value=datetime(2026,10,2,6,tzinfo=db.BOGOTA)):
   task=db.plan_personal_task(self.raw,self.user)
   with self.assertRaises(ValueError):db.submit_report({**self.raw,'pending_id':task},self.user)
   self.assertEqual(db.personal_tasks('u',db.today())[0]['status'],'PENDIENTE')
 def test_overlap_hours_and_breaks(self):
  db.plan_personal_task(self.raw,self.user)
  with self.assertRaises(ValueError):db.plan_personal_task({**self.raw,'start':'08:00','end':'10:00'},self.user)
  with self.assertRaises(ValueError):db.plan_personal_task({**self.raw,'start':'05:00'},self.user)
  with self.assertRaises(ValueError):db.plan_personal_task({**self.raw,'start':'12:00','end':'13:00'},self.user)
  r=db.submit_report({**self.raw,'start':'11:00','end':'14:00'},self.user);self.assertEqual(db.report(r)['minutes'],120)
 def test_analysis_zero_days_returned_and_average(self):
  a=db.personal_analysis('u',db.parse_day('2026-10-01'),db.today());self.assertEqual(a['average'],0)
  rid=db.submit_report(self.raw,self.user);a=db.personal_analysis('u',db.parse_day('2026-10-01'),db.today())
  self.assertEqual(a['average'],.5);self.assertEqual(a['actual'],180);self.assertEqual(a['productivity'],18.8)
  db.review_report(rid,{'status':'DEVUELTO','review_note':'Corregir información','version':1},self.chief)
  self.assertEqual(db.personal_analysis('u',db.today(),db.today())['actual'],0)
 def test_render_and_export_private_dates_and_formulas(self):
  db.submit_report({**self.raw,'project':'=HYPERLINK("bad")'},self.user);db.submit_report({**self.raw,'description':'DATOS DEL OTRO DISEÑADOR'},self.other)
  client=self.app().test_client()
  with patch.object(db,'sync_project_activities'),patch.object(routes,'_my_assignment_options',return_value=[]):
   response=client.get('/reportes/?fecha=2026-10-02&analizar=1');self.assertEqual(response.status_code,200);self.assertIn(b'personalChart',response.data)
  result=client.get('/reportes/mis-reportes/excel?desde=2026-10-02&hasta=2026-10-02');self.assertEqual(result.status_code,200)
  w=load_workbook(io.BytesIO(result.data));self.assertEqual(w['Registro'].max_row,2);self.assertEqual(w['Registro']['I2'].value,3)
  self.assertEqual(w['Registro']['A2'].data_type,'s');self.assertTrue(w['Registro']['A2'].value.startswith("'="));self.assertEqual(len(w['Resumen']._charts),1)
  self.assertEqual(client.get('/reportes/mis-reportes/excel?desde=2026-10-03').status_code,400)
 def test_activity_without_observation_and_no_minimum(self):
  raw={k:v for k,v in self.raw.items() if k!='description'}
  raw.update(activity='Diseñar tolva',start='06:00',end='15:00')
  task=db.plan_personal_task(raw,self.user)
  rid=db.submit_report({**raw,'pending_id':task},self.user)
  self.assertEqual(db.report(rid)['description'],'Diseñar tolva')
  with db.connection() as c:
   config=json.loads(c.execute('SELECT config FROM profile_history WHERE user_id=?',('u',)).fetchone()[0]);config['minimum_reports']=20
   c.execute('UPDATE reports SET schedule=? WHERE id=?',(json.dumps(config),rid))
  summary=db.day_summary('u',db.today())
  self.assertEqual(summary['status'],'CUMPLIDO');self.assertEqual(summary['minimum'],0)
  client=self.app().test_client()
  with patch.object(db,'sync_project_activities'),patch.object(routes,'_my_assignment_options',return_value=[]):
   page=client.get('/reportes/?fecha=2026-10-02').get_data(as_text=True)
  self.assertNotIn('name="description"',page);self.assertIn('name="activity"',page)
  self.assertNotIn('Reportes del',page);self.assertNotIn('Avance de mis asignaciones',page);self.assertNotIn('mínimo diario',page)
 def test_assigned_task_list_completion_and_access(self):
  raw={'user_id':'u','title':'Diseñar tolva','start_date':'2026-10-02','due_date':'2026-10-05','planned_hours':8}
  task=db.save_assignment(raw,self.chief)
  source=db.save_assignment({**raw,'title':'ACTIVIDAD AUTOMÁTICA DEL PROYECTO'},self.chief)
  other=db.save_assignment({**raw,'user_id':'other','title':'PRIVADO OTRO DISEÑADOR'},self.chief)
  with db.connection() as c:c.execute("UPDATE assignments SET source_id='source' WHERE id=?",(source,))
  client=self.app().test_client()
  with patch.object(db,'sync_project_activities'):
   response=client.get('/reportes/asignaciones');self.assertEqual(response.status_code,200)
   page=response.get_data(as_text=True);self.assertIn('Diseñar tolva',page);self.assertNotIn('ACTIVIDAD AUTOMÁTICA',page);self.assertNotIn('PRIVADO OTRO',page)
  self.assertEqual(client.post('/reportes/asignaciones/'+other+'/completar',data={'day':'2026-10-02','start':'06:00','end':'09:00'}).status_code,403)
  self.assertEqual(client.post('/reportes/asignaciones/'+source+'/completar',data={'day':'2026-10-02','start':'06:00','end':'09:00'}).status_code,403)
  self.assertEqual(client.post('/reportes/asignaciones/'+task+'/completar',data={'day':'2026-10-02','start':'06:00','end':'09:00'}).status_code,302)
  self.assertEqual(next(t for t in db.assignments(['u']) if t['id']==task)['progress'],100)
  client.post('/reportes/asignaciones/'+task+'/completar',data={'day':'2026-10-02','start':'09:00','end':'10:00'})
  self.assertEqual(len(db.list_reports(['u'],db.today(),db.today())),1)
 def test_chief_assigns_own_team_and_selector_removed(self):
  chief={**self.chief,'name':'Jefe','role':auth.ROLE_CHIEF,'is_chief':True}
  self.user=chief
  designer={'id':'u','name':'Diseñador','is_active':True,'role':auth.ROLE_COLLABORATOR,'profile':db.get_profile('u')}
  client=self.app().test_client()
  with patch.object(routes,'_configured_users',return_value=[designer]),patch.object(db,'sync_project_activities'),patch.object(routes,'get_user_by_id',return_value=designer):
   result=client.post('/reportes/asignaciones/guardar',data={'user_id':'u','title':'Preparar planos','start_date':'2026-10-02','due_date':'2026-10-05','planned_hours':'8'})
   self.assertEqual(result.status_code,302)
   page=client.get('/reportes/asignaciones').get_data(as_text=True);self.assertIn('Asignar una tarea',page);self.assertIn('Preparar planos',page)
  self.user={'id':'u','name':'Diseñador','role':auth.ROLE_COLLABORATOR}
  with patch.object(db,'sync_project_activities'),patch.object(routes,'_my_assignment_options',return_value=[]):
   page=client.get('/reportes/?fecha=2026-10-02').get_data(as_text=True)
   self.assertNotIn('Asignación del jefe',page);self.assertNotIn('<select name="assignment_id"',page)
 def test_chief_cannot_report_and_dashboard_tasks(self):
  designer=dict(self.user)
  pending=db.plan_personal_task(self.raw,designer)
  completed=db.submit_report({**self.raw,'start':'09:00','end':'11:00','activity':'TAREA COMPLETADA DEL DISEÑADOR'},designer)
  assignment=db.save_assignment({'user_id':'u','title':'TAREA ASIGNADA POR JEFE','start_date':'2026-10-02','due_date':'2026-10-02','planned_hours':8},self.chief)
  self.user={**self.chief,'name':'Jefe','role':auth.ROLE_CHIEF,'is_chief':True}
  with self.assertRaises(PermissionError):db.submit_report(self.raw,self.user)
  with self.assertRaises(PermissionError):db.plan_personal_task(self.raw,self.user)
  client=self.app().test_client()
  for endpoint in ['/reportes/enviar','/reportes/mis-tareas/programar','/reportes/mis-tareas/'+pending+'/cancelar','/reportes/asignaciones/'+assignment+'/completar']:
   self.assertEqual(client.post(endpoint,data=self.raw).status_code,403)
  self.assertEqual(client.get('/reportes/mis-reportes/excel').status_code,403)
  response=client.get('/reportes/');self.assertEqual(response.status_code,302);self.assertIn('/reportes/tablero',response.location)
  designer.update(profile=db.get_profile('u'),is_active=True)
  with patch.object(routes,'_configured_users',return_value=[designer]),patch.object(db,'sync_project_activities'):
   page=client.get('/reportes/tablero?desde=2026-10-02&hasta=2026-10-02');self.assertEqual(page.status_code,200)
   text=page.get_data(as_text=True)
   self.assertIn('Tareas programadas por los diseñadores',text);self.assertIn('TAREA ASIGNADA POR JEFE',text);self.assertIn('TAREA COMPLETADA DEL DISEÑADOR',text)
   self.assertIn('ATRASADA',text);self.assertNotIn('Reportado / esperado',text)
   self.assertEqual(client.get('/reportes/tablero?usuario=other').status_code,403)
   export=client.get('/reportes/exportar?desde=2026-10-02&hasta=2026-10-02').get_data(as_text=True)
   self.assertIn('Centro de trabajo',export);self.assertIn('TAREA COMPLETADA DEL DISEÑADOR',export)
 def test_chief_navigation_no_personal_report_link(self):
  spec=importlib.util.spec_from_file_location('real_permissions',root/'permisos.py');module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
  app=self.app()
  with app.test_request_context('/'):
   g.current_user={'id':'chief','role':auth.ROLE_CHIEF,'is_chief':True}
   groups=module.navigation();report_group=next(x for x in groups if x['module']=='reportes')
   endpoints=[x[0] for x in report_group['links']]
   self.assertNotIn('reportes.index',endpoints);self.assertEqual(endpoints,['reportes.dashboard','reportes.tasks'])
 def test_multipoint_all_intermediate_records_and_weekly_tables(self):
  for start,end in [('06:00','07:00'),('09:00','10:00'),('10:00','12:00')]:
   db.submit_report({**self.raw,'start':start,'end':end,'activity':'Tarea '+start},self.user)
  data=db.personal_analysis('u',db.parse_day('2026-10-01'),db.today())
  self.assertEqual(len(data['series']),3)
  self.assertEqual([p['time'] for p in data['series']],['07:00','10:00','12:00'])
  self.assertEqual([p['productivity'] for p in data['series']],[100,50,66.7])
  self.assertEqual(data['weeks'][0]['actual'],240)
  self.assertEqual(data['weeks'][0]['productivity'],25)
  client=self.app().test_client()
  with patch.object(db,'sync_project_activities'),patch.object(routes,'_my_assignment_options',return_value=[]):
   page=client.get('/reportes/?fecha=2026-10-02&analizar=1').get_data(as_text=True)
   self.assertNotIn('Tabla por día',page);self.assertNotIn('Tabla por semana',page)
   self.assertIn('chartValues=[100.0, 50.0, 66.7]',page)
  response=client.get('/reportes/mis-reportes/excel?desde=2026-10-01&hasta=2026-10-02')
  workbook=load_workbook(io.BytesIO(response.data))
  self.assertEqual(workbook['Tendencia'].max_row,4)
  self.assertEqual(workbook['Resumen']._charts[0].series[0].val.numRef.f,"'Tendencia'!$E$2:$E$4")
  self.assertEqual(workbook['Semanal']['E2'].value,16);self.assertEqual(workbook['Semanal']['F2'].value,4)
  with patch.object(db,'now',return_value=datetime(2026,10,6,16,tzinfo=db.BOGOTA)):
   weeks=db.personal_analysis('u',db.parse_day('2026-10-01'),db.today())['weeks']
   self.assertEqual(len(weeks),2);self.assertEqual(weeks[1]['start'],'2026-10-05');self.assertEqual(weeks[1]['count'],0)
 def test_hours_decimal_screenshot_elapsed_and_remaining(self):
  config={'start':'06:00','end':'15:00','breaks':'12:00-12:30','weekdays':[0,1,2,3,4],'modules':[]}
  db.save_profile('u','chief','',config,self.chief,db.get_profile('u')['version'])
  with patch.object(db,'now',return_value=datetime(2026,10,2,11,46,tzinfo=db.BOGOTA)):
   db.submit_report({**self.raw,'start':'06:00','end':'11:38'},self.user)
   summary=db.day_summary('u',db.today())
   self.assertEqual(summary['expected'],510);self.assertEqual(summary['actual'],338)
   self.assertEqual(summary['elapsed'],346);self.assertEqual(summary['unregistered_elapsed'],8)
   self.assertEqual(summary['remaining_work'],164)
   self.assertEqual(db.duration_label(summary['actual']),'5 h 38 min')
   self.assertEqual(db.duration_label(summary['expected']),'8 h 30 min')
   client=self.app().test_client()
   with patch.object(db,'sync_project_activities'),patch.object(routes,'_my_assignment_options',return_value=[]):
    page=client.get('/reportes/?fecha=2026-10-02').get_data(as_text=True)
    self.assertIn('5 h 38 min',page);self.assertIn('0 h 08 min',page);self.assertIn('Falta por trabajar: 2 h 44 min',page)
  with patch.object(db,'now',return_value=datetime(2026,10,2,12,15,tzinfo=db.BOGOTA)):
   self.assertEqual(db.day_summary('u',db.today())['elapsed'],360)
  with patch.object(db,'now',return_value=datetime(2026,10,2,16,tzinfo=db.BOGOTA)):
   summary=db.day_summary('u',db.today());self.assertEqual(summary['remaining_work'],0);self.assertEqual(summary['unregistered_elapsed'],172)
 def test_before_entry_future_and_multiple_breaks(self):
  config={'start':'06:00','end':'15:00','breaks':'08:00-08:15\n12:00-12:30','weekdays':[0,1,2,3,4],'modules':[]}
  db.save_profile('u','chief','',config,self.chief,db.get_profile('u')['version'])
  with patch.object(db,'now',return_value=datetime(2026,10,2,5,30,tzinfo=db.BOGOTA)):
   summary=db.day_summary('u',db.today());self.assertEqual(summary['elapsed'],0);self.assertEqual(summary['unregistered_elapsed'],0);self.assertEqual(summary['remaining_work'],495)
  with patch.object(db,'now',return_value=datetime(2026,10,2,8,10,tzinfo=db.BOGOTA)):
   self.assertEqual(db.day_summary('u',db.today())['elapsed'],120)
  self.assertEqual(db.day_summary('u',db.parse_day('2026-10-05'))['elapsed'],0)
  self.assertEqual(db.net_minutes(6*60,15*60,db.validate_config(config)),495)
 def test_migration_and_assignment_unchanged(self):
  db.ensure_database();db.ensure_database()
  task=db.save_assignment({'user_id':'u','title':'Preparar diseño','start_date':'2026-10-02','due_date':'2026-10-05','planned_hours':8},self.chief)
  pending=db.plan_personal_task({**self.raw,'assignment_id':task},self.user)
  with self.assertRaises(ValueError):db.submit_report({**self.raw,'pending_id':pending},self.user)
  db.submit_report({**self.raw,'pending_id':pending,'assignment_id':task,'progress':30},self.user)
  self.assertEqual(db.assignment(task)['status'],'ACTIVA')
if __name__=='__main__':unittest.main(verbosity=2)
