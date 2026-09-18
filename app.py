from __future__ import annotations

from collections import defaultdict, deque
from datetime import date, datetime, timedelta
from functools import wraps
from io import BytesIO
from pathlib import Path
import hmac
import os
import re
import secrets
import shutil
import threading
import time
import webbrowser
from urllib.parse import urlsplit

from flask import Flask, render_template, request, redirect, url_for, flash, send_file, abort, jsonify, session, g
from openpyxl import load_workbook

from storage import (
    DATA_DIR, MASTER_FILE, CONFIG_FILE, HOLIDAYS_FILE,
    ensure_dirs, ensure_simple_workbooks, save_upload, read_rows, write_rows, clear_workspace,
    workspace_dirty, read_meta, write_meta, set_meta, history_rows, archive_workspace,
    recover_history, delete_history, backup_master, safe_float, clean_code, normalize_text,
    create_autosave, autosave_rows, recover_autosave, delete_autosave,
    get_work_dir, set_user_scope, migrate_legacy_data_to_current_user
)
from auth_store import (
    AUDIT_FILE, ROLE_ADMIN, ROLE_COLLABORATOR, STATUS_ACTIVE, STATUS_DISABLED,
    ensure_security_storage, get_app_secret, has_users, list_users, get_user_by_id,
    create_user, create_initial_admin, authenticate, reset_password, change_role,
    set_user_status, unlock_user,
    audit_event, audit_rows
)
from engine import (
    MASTER_REQUIRED, load_master, invalidate_master_cache, validate_master_file, append_master_code,
    parse_b1, parse_lge_excel, build_details, consolidate, make_lm, movement_proposal, rq_proposal,
    spare_rows, value_rows, parse_code_qty, compare_actual, classification
)
from planos import planos_bp
from biblioteca import biblioteca_bp
from proyectos_diseno.routes import bp as proyectos_diseno_bp
from capacitaciones import bp as capacitaciones_bp, ensure_database as ensure_capacitaciones_database, ensure_worker_started as ensure_capacitaciones_worker
from oportunidades_proyecto import bp as oportunidades_bp, ensure_database as ensure_oportunidades_database
from asistente_oga import bp as asistente_oga_bp

APP_NAME = "SISTEMA TECNICO"
APP_REV = "0.17.5.3"

ensure_security_storage()
app = Flask(__name__)
app.secret_key = get_app_secret()
app.config.update(
    # Capacitaciones permite MP4 grandes; el límite tradicional de 96 MB se
    # conserva por endpoint para el resto de módulos.
    MAX_CONTENT_LENGTH=max(96 * 1024 * 1024, int(os.environ.get("OGA_MAX_VIDEO_MB", "2048")) * 1024 * 1024 + 2 * 1024 * 1024),
    SESSION_COOKIE_NAME="oga_herramientas_session",
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=os.environ.get("OGA_HTTPS", "0") == "1",
    PERMANENT_SESSION_LIFETIME=timedelta(hours=10),
)
app.register_blueprint(planos_bp)
app.register_blueprint(biblioteca_bp)
app.register_blueprint(proyectos_diseno_bp, url_prefix="/proyectos-diseno/app")
app.register_blueprint(capacitaciones_bp)
app.register_blueprint(oportunidades_bp)
app.register_blueprint(asistente_oga_bp)

LOGIN_WINDOW_SECONDS = 15 * 60
LOGIN_IP_LIMIT = 25
IDLE_TIMEOUT_SECONDS = 60 * 60
_IP_LOGIN_ATTEMPTS = defaultdict(deque)
_IP_LOCK = threading.RLock()


def client_ip():
    return request.remote_addr or "DESCONOCIDA"


def csrf_token():
    token = session.get("csrf_token")
    if not token:
        token = secrets.token_urlsafe(32)
        session["csrf_token"] = token
    return token


def _ip_is_limited(ip):
    now = time.time()
    with _IP_LOCK:
        attempts = _IP_LOGIN_ATTEMPTS[ip]
        while attempts and now - attempts[0] > LOGIN_WINDOW_SECONDS:
            attempts.popleft()
        return len(attempts) >= LOGIN_IP_LIMIT


def _record_ip_failure(ip):
    with _IP_LOCK:
        _IP_LOGIN_ATTEMPTS[ip].append(time.time())


def _clear_ip_failures(ip):
    with _IP_LOCK:
        _IP_LOGIN_ATTEMPTS.pop(ip, None)


def admin_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        user = getattr(g, "current_user", None)
        if not user or not user.get("is_admin"):
            abort(403)
        return view(*args, **kwargs)
    return wrapped


@app.before_request
def enforce_security():
    g.current_user = None
    g.audit_logged = False
    set_user_scope(None)

    user_id = session.get("user_id")
    if user_id:
        user = get_user_by_id(user_id)
        version_ok = user and int(session.get("session_version") or 0) == int(user.get("session_version") or 0)
        if user and user.get("is_active") and version_ok:
            last_activity = float(session.get("last_activity") or 0)
            if last_activity and time.time() - last_activity > IDLE_TIMEOUT_SECONDS:
                session.clear()
                if request.endpoint not in {"login", "initial_setup", "static"}:
                    flash("La sesión se cerró por inactividad.", "error")
                    return redirect(url_for("login"))
            else:
                g.current_user = user
                set_user_scope(user["id"])
                ensure_dirs()
                ensure_simple_workbooks()
                session["last_activity"] = time.time()
        else:
            session.clear()

    users_exist = has_users()
    if not users_exist:
        if request.endpoint not in {"initial_setup", "static"}:
            return redirect(url_for("initial_setup"))
    elif request.endpoint == "initial_setup":
        return redirect(url_for("login"))

    public_endpoints = {"login", "initial_setup", "static"}
    if users_exist and request.endpoint not in public_endpoints and g.current_user is None:
        next_path = request.full_path.rstrip("?") if request.method == "GET" else ""
        return redirect(url_for("login", next=next_path))

    if request.endpoint in {"login", "initial_setup", "admin_user_create", "admin_user_password"}:
        if (request.content_length or 0) > 16 * 1024:
            abort(413)

    if request.endpoint != "capacitaciones.upload" and (request.content_length or 0) > 96 * 1024 * 1024:
        abort(413)

    if request.method in {"POST", "PUT", "PATCH", "DELETE"}:
        supplied = request.form.get("_csrf_token", "") or request.headers.get("X-CSRF-Token", "")
        expected = session.get("csrf_token", "")
        if not supplied or not expected or not hmac.compare_digest(str(supplied), str(expected)):
            if request.endpoint and request.endpoint.startswith("planos."):
                return jsonify({"ok": False, "error": "La sesión de seguridad venció. Recargue la página."}), 400
            abort(400, description="La sesión de seguridad venció. Recargue la página.")


AUDIT_ENDPOINTS = {
    "biblioteca.create_equipo": ("CREAR EQUIPO", "BIBLIOTECA"),
    "biblioteca.update_equipo": ("EDITAR EQUIPO", "BIBLIOTECA"),
    "biblioteca.delete_equipo": ("ELIMINAR EQUIPO", "BIBLIOTECA"),
    "biblioteca.create_caracteristica": ("CREAR CARACTERISTICA", "BIBLIOTECA"),
    "biblioteca.update_caracteristica": ("EDITAR CARACTERISTICA", "BIBLIOTECA"),
    "biblioteca.delete_caracteristica": ("ELIMINAR CARACTERISTICA", "BIBLIOTECA"),
    "biblioteca.create_referencia": ("CREAR REFERENCIA", "BIBLIOTECA"),
    "biblioteca.update_referencia": ("EDITAR REFERENCIA", "BIBLIOTECA"),
    "biblioteca.delete_referencia": ("ELIMINAR REFERENCIA", "BIBLIOTECA"),
    "biblioteca.costos": ("CARGAR COSTOS", "BIBLIOTECA"),
    "biblioteca.download_file": ("DESCARGAR DOCUMENTO", "BIBLIOTECA"),
    "backup_create": ("CREAR COPIA", "COPIAS"),
    "backup_recover": ("RECUPERAR COPIA", "COPIAS"),
    "backup_delete": ("ELIMINAR COPIA", "COPIAS"),
    "upload_inventor": ("IMPORTAR EXCEL INVENTOR", "CARGAR LISTA"),
    "upload_b1": ("IMPORTAR EXCEL B1", "CARGAR LISTA"),
    "upload_lge": ("IMPORTAR LISTADO DE EQUIPOS", "CARGAR LISTA"),
    "upload_equipment_inventor": ("IMPORTAR EXCEL DE EQUIPO", "CARGAR LISTA"),
    "remove_lge_row": ("QUITAR EQUIPO", "CARGAR LISTA"),
    "process_route": ("PROCESAR LISTA", "RQ"),
    "rq_date": ("CAMBIAR FECHA BASE", "RQ"),
    "externos_add": ("AGREGAR EXTERNO", "RQ"),
    "externos_remove": ("QUITAR EXTERNO", "RQ"),
    "load_mov": ("IMPORTAR MOV", "RQ"),
    "load_rq": ("IMPORTAR RQ", "RQ"),
    "factory_new": ("CREAR CÓDIGO", "LISTA MAESTRA"),
    "factory_reload": ("RECARGAR MAESTRA", "LISTA MAESTRA"),
    "factory_replace": ("REEMPLAZAR MAESTRA", "LISTA MAESTRA"),
    "factory_download": ("DESCARGAR EXCEL", "LISTA MAESTRA"),
    "calendario_config": ("CAMBIAR CONFIGURACIÓN", "CALENDARIO"),
    "holiday_add": ("AGREGAR FESTIVO", "CALENDARIO"),
    "holiday_delete": ("ELIMINAR FESTIVO", "CALENDARIO"),
    "archive": ("ARCHIVAR TRABAJO", "HISTÓRICO"),
    "history_recover": ("RECUPERAR TRABAJO", "HISTÓRICO"),
    "history_delete": ("ELIMINAR TRABAJO", "HISTÓRICO"),
    "new_list": ("NUEVA LISTA", "TRABAJO"),
    "export_file": ("EXPORTAR EXCEL", "EXPORTACIONES"),
    "export_rq_category": ("EXPORTAR CATEGORÍA RQ", "EXPORTACIONES"),
    "export_all": ("EXPORTAR TODO", "EXPORTACIONES"),
    "planos.upload": ("IMPORTAR PDF", "PLANOS"),
    "planos.review": ("REVISAR PDF", "PLANOS"),
    "planos.clear": ("LIMPIAR REVISIÓN", "PLANOS"),
    "planos.download_current": ("DESCARGAR EXCEL", "PLANOS"),
    "planos.download_master": ("DESCARGAR HISTÓRICO", "PLANOS"),
    "planos.download_report": ("DESCARGAR REPORTE", "PLANOS"),
    "planos.delete_report": ("ELIMINAR REPORTE", "PLANOS"),
    "planos.dist_upload_pdf": ("IMPORTAR PDF DE DISTRIBUCIÓN", "PLANOS"),
    "planos.dist_upload_lm": ("IMPORTAR EXCEL LM", "PLANOS"),
    "planos.dist_upload_dxf": ("IMPORTAR EXCEL DXF", "PLANOS"),
    "planos.dist_review": ("REVISAR DISTRIBUCIÓN", "PLANOS"),
    "planos.dist_clear": ("LIMPIAR DISTRIBUCIÓN", "PLANOS"),
}


@app.after_request
def apply_security_and_audit(response):
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "SAMEORIGIN"
    response.headers["Referrer-Policy"] = "same-origin"
    response.headers["Permissions-Policy"] = "camera=(), microphone=(), geolocation=()"
    response.headers["Content-Security-Policy"] = (
        "default-src 'self'; script-src 'self' 'unsafe-inline'; "
        "style-src 'self' 'unsafe-inline'; img-src 'self' data: blob:; "
        "font-src 'self' data:; connect-src 'self'; frame-src 'self' blob:; "
        "media-src 'self' blob:; worker-src 'self' blob:; "
        "object-src 'none'; base-uri 'self'; form-action 'self'; frame-ancestors 'self'"
    )
    if request.endpoint and request.endpoint.startswith("capacitaciones.") and request.endpoint in {
        "capacitaciones.media_mp4", "capacitaciones.poster", "capacitaciones.hls"
    }:
        response.headers.setdefault("Cache-Control", "private, max-age=3600")
    else:
        response.headers["Cache-Control"] = "no-store" if getattr(g, "current_user", None) else "no-cache"
    if app.config.get("SESSION_COOKIE_SECURE"):
        response.headers["Strict-Transport-Security"] = "max-age=31536000; includeSubDomains"

    endpoint = request.endpoint or ""
    event = AUDIT_ENDPOINTS.get(endpoint)
    if endpoint == "biblioteca.costos" and request.method != "POST":
        event = None
    if event and getattr(g, "current_user", None) and not getattr(g, "audit_logged", False):
        try:
            file_names = [Path(item.filename).name for item in request.files.values() if item and item.filename]
            params = [str(value) for value in (request.view_args or {}).values()]
            detail_parts = file_names + params
            audit_event(
                event[0], event[1], " · ".join(detail_parts),
                "OK" if response.status_code < 400 else f"ERROR {response.status_code}",
                g.current_user, client_ip(),
            )
        except Exception:
            app.logger.exception("No fue posible registrar la auditoría de %s", endpoint)
    return response


@app.template_filter("peso")
def peso(v):
    if v in (None, "", "VALOR NO DISPONIBLE"):
        return v or ""
    try:
        return "$ {:,.0f}".format(float(v)).replace(",", ".")
    except Exception:
        return str(v)


@app.template_filter("num")
def num(v):
    try:
        f = float(v)
        if abs(f - round(f)) < 1e-9:
            return str(int(round(f)))
        return ("{:.4f}".format(f)).rstrip("0").rstrip(".")
    except Exception:
        return str(v or "")


@app.template_filter("filesize")
def filesize(v):
    try:
        size = float(v or 0)
    except Exception:
        return "0 KB"
    if size >= 1024 * 1024:
        return f"{size / (1024 * 1024):.1f} MB"
    return f"{max(size / 1024, 0):.0f} KB"


def init_app_data():
    set_user_scope(None)
    ensure_dirs(include_user=False)
    ensure_simple_workbooks()
    ensure_security_storage()
    ensure_capacitaciones_database()
    ensure_capacitaciones_worker()
    ensure_oportunidades_database()
    if not MASTER_FILE.exists():
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Hoja1"
        for i, h in enumerate(MASTER_REQUIRED, 1):
            ws.cell(1, i).value = h
        ws.cell(1, 16).value = "REGLAS"
        # Las reglas son fijas en el programa. En el Excel maestro solo se indica el numero de regla.
        wb.save(MASTER_FILE)

    # Los trabajos ya no se limpian al iniciar: cada usuario conserva su
    # propio espacio y puede reanudarlo después de cerrar el navegador.


def ctx(**kwargs):
    user = getattr(g, "current_user", None)
    meta = read_meta() if user else {}
    return dict(
        app_name=APP_NAME,
        app_rev=APP_REV,
        dirty=workspace_dirty() if user else False,
        meta=meta,
        current_user=user,
        csrf_token=csrf_token(),
        **kwargs,
    )


def checkpoint(reason, force=False):
    """Guarda el trabajo sin interrumpir la acción principal si falla la copia."""
    try:
        return create_autosave(reason, force=force)
    except Exception:
        app.logger.exception("No fue posible crear el punto de recuperación: %s", reason)
        return None


@app.context_processor
def inject_template_context():
    """Comparte el contexto general con las vistas del Blueprint de Planos."""
    return ctx()


def current_paths():
    work_dir = get_work_dir()
    return {
        "general": work_dir / "INVENTOR_GENERAL.xlsx",
        "b1_raw": work_dir / "B1_ORIGINAL.xls",
        "b1_filtered": work_dir / "B1_FILTRADO.xlsx",
        "lge_source": work_dir / "LGE_ORIGEN.xlsx",
        "lge": work_dir / "LGE.xlsx",
        "details": work_dir / "DETALLE_INVENTOR.xlsx",
        "lm": work_dir / "LM.xlsx",
        "mov": work_dir / "MOV_PROPUESTO.xlsx",
        "rq": work_dir / "RQ_PROPUESTA.xlsx",
        "rq_export": work_dir / "RQ_PARA_FACTORY.xlsx",
        "spares": work_dir / "REPUESTOS.xlsx",
        "value": work_dir / "VALOR.xlsx",
        "eqvalue": work_dir / "VALOR_EQUIPOS.xlsx",
        "trace": work_dir / "TRAZABILIDAD.xlsx",
        "externos_manual": work_dir / "EXTERNOS_MANUAL.xlsx",
    }


def lge_rows():
    return read_rows(current_paths()["lge"])


def save_lge(rows):
    write_rows(current_paths()["lge"], ["CODIGO", "CANTIDAD", "DESCRIPCION", "PART NUMBER", "ARCHIVO"], rows, "LGE")


def externos_manual_codes():
    """Códigos agregados a mano a EXTERNOS. Esta es la ÚNICA forma de que
    un código quede en esa clasificación: no hay regla automática."""
    rows = read_rows(current_paths()["externos_manual"])
    return {clean_code(r.get("CODIGO")) for r in rows if clean_code(r.get("CODIGO"))}


def save_externos_manual(codes):
    rows = [{"CODIGO": c} for c in sorted(codes)]
    write_rows(current_paths()["externos_manual"], ["CODIGO"], rows, "EXTERNOS_MANUAL")


def workspace_sources():
    work_dir = get_work_dir()
    p = current_paths()
    rows = lge_rows()
    if rows:
        sources = []
        missing = []
        for r in rows:
            fn = str(r.get("ARCHIVO") or "").strip()
            if not fn:
                missing.append(f"{r.get('CODIGO')} - {r.get('DESCRIPCION')}")
                continue
            fp = work_dir / "INVENTOR_EQUIPOS" / fn
            if not fp.exists():
                missing.append(f"{r.get('CODIGO')} - {r.get('DESCRIPCION')}")
                continue
            sources.append({"path": fp, "equipment": f"{r.get('CODIGO')} - {r.get('DESCRIPCION')}", "multiplier": safe_float(r.get("CANTIDAD"), 1)})
        return sources, missing
    if p["general"].exists():
        return [{"path": p["general"], "equipment": "GENERAL", "multiplier": 1}], []
    alt = read_meta().get("inventor_general_path")
    if alt and (work_dir / alt).exists():
        return [{"path": work_dir / alt, "equipment": "GENERAL", "multiplier": 1}], []
    return [], ["LISTA DE INVENTOR"]


def base_date_from_meta():
    raw = read_meta().get("fecha_base_rq")
    if raw:
        try:
            return datetime.strptime(raw, "%Y-%m-%d").date()
        except Exception:
            pass
    return date.today()


RQ_CATEGORY_ORDER = ("C.TEMPRANA", "EXTERNO", "PRODUCCION", "ENSAMBLE", "SIN CLASIFICAR")
RQ_CATEGORY_LABELS = {
    "C.TEMPRANA": "COMPRAS TEMPRANAS",
    "EXTERNO": "EXTERNOS",
    "PRODUCCION": "PRODUCCIÓN",
    "ENSAMBLE": "ENSAMBLE",
    "SIN CLASIFICAR": "SIN CLASIFICACIÓN",
}


def rq_category_key(value):
    category = normalize_text(value)
    if category in {"", "SIN CLASIFICACION", "SIN CLASIFICAR"}:
        return "SIN CLASIFICAR"
    return category


def rq_category_slug(category):
    slug = normalize_text(category).replace(".", "-")
    slug = re.sub(r"[^A-Z0-9]+", "-", slug).strip("-")
    return slug.lower() or "sin-clasificar"


def group_rq_rows(rows):
    """Agrupa las filas de la RQ por clasificación base (C.TEMPRANA,
    PRODUCCION, ENSAMBLE o sin clasificar). EXTERNOS ya no se calcula ni se
    cruza aquí — es una lista aparte, totalmente manual (ver
    externos_manual_rows)."""
    grouped = {}
    for row in rows:
        category = rq_category_key(row.get("CLASIFICACION"))
        item = dict(row)
        item["CLASIFICACION"] = category
        grouped.setdefault(category, []).append(item)

    order = {category: index for index, category in enumerate(RQ_CATEGORY_ORDER)}
    categories = sorted(grouped, key=lambda category: (
        1 if category == "SIN CLASIFICAR" else 0,
        order.get(category, 90),
        category,
    ))
    return [{
        "category": category,
        "label": RQ_CATEGORY_LABELS.get(category, category),
        "slug": rq_category_slug(category),
        "rows": grouped[category],
        "is_unclassified": category == "SIN CLASIFICAR",
    } for category in categories]


def externos_manual_rows():
    """Filas para la tabla EXTERNOS: solo lo que se agregó a mano con el
    botón. No depende de la RQ procesada ni le quita nada a las otras
    tablas — es una lista independiente de seguimiento."""
    _, master_map, _ = load_master()
    lm_map = {clean_code(r.get("CODIGO")): r for r in read_rows(current_paths()["lm"])}
    rows = []
    for code in sorted(externos_manual_codes()):
        m = master_map.get(code, {})
        lm = lm_map.get(code, {})
        rows.append({
            "CODIGO": code,
            "NOMBRE": m.get("NOMBRE", ""),
            "UD": m.get("U", ""),
            "PENDIENTE": lm.get("PENDIENTE", ""),
        })
    return rows


def write_rq_export(path, rows, include_category_sheets=True, sheet_title="RQ"):
    """Genera la RQ para uso humano con cada nota debajo del codigo al que pertenece.
    Incluye una hoja general y una hoja adicional por cada clasificacion encontrada."""
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    headers = ["CLASIFICACION", "CODIGO", "NOMBRE", "UD", "CANTIDAD", "FECHA REQUERIDA"]
    blue = "0B64B4"
    thin = Side(style="thin", color="D9E2F3")

    def fill_sheet(ws, sheet_rows):
        ws.append(headers)
        for c in ws[1]:
            c.fill = PatternFill("solid", fgColor=blue)
            c.font = Font(color="FFFFFF", bold=True)
            c.alignment = Alignment(horizontal="center", vertical="center")
        for r in sheet_rows:
            ws.append([r.get(h, "") for h in headers])
            item_row = ws.max_row
            for c in ws[item_row]:
                c.border = Border(bottom=thin)
                c.alignment = Alignment(vertical="top", wrap_text=True)
            if rq_category_key(r.get("CLASIFICACION")) == "SIN CLASIFICAR":
                for c in ws[item_row]:
                    c.fill = PatternFill("solid", fgColor="FCE8E8")
                    c.font = Font(color="9C0006", bold=c.column in {1, 2})
            notes = [x.strip() for x in str(r.get("NOTA") or "").split("\n") if x.strip() and x.strip().upper() != "NO"]
            for note in notes:
                ws.append(["", "NOTA:", note, "", "", ""])
                nr = ws.max_row
                ws.cell(nr, 2).font = Font(bold=True, italic=True, color="5B6570")
                ws.cell(nr, 3).font = Font(italic=True, color="5B6570")
                ws.cell(nr, 3).alignment = Alignment(wrap_text=True)
                ws.merge_cells(start_row=nr, start_column=3, end_row=nr, end_column=6)
                if "VERIFICAR TERMINADO" in note.upper():
                    for cc in range(2, 7):
                        ws.cell(nr, cc).fill = PatternFill("solid", fgColor="FFF2CC")
        widths = [18, 16, 54, 10, 14, 18]
        for i, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:F{max(1, ws.max_row)}"

    wb = Workbook()
    ws_general = wb.active
    ws_general.title = sheet_title[:31]
    fill_sheet(ws_general, rows)

    if include_category_sheets:
        for group in group_rq_rows(rows):
            ws_group = wb.create_sheet(group["label"][:31])
            fill_sheet(ws_group, group["rows"])
            if group["is_unclassified"]:
                ws_group.sheet_properties.tabColor = "C7333E"

    if isinstance(path, (str, Path)):
        Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def process_workspace():
    p = current_paths()
    sources, missing = workspace_sources()
    if missing:
        raise ValueError("Falta cargar la lista de Inventor para: " + "; ".join(missing[:12]))
    details = build_details(sources)
    write_rows(p["details"], [
        "EQUIPO", "ITEM", "FILENAME", "PART NUMBER", "DESCRIPTION", "STOCK NUMBER", "ITEM QTY",
        "REGLA", "UNIDAD", "CANTIDAD CALCULADA", "LONGITUD MM", "AREA M2", "ESTADO CODIGO", "NOTA EXTRA"
    ], details, "DETALLE", status_col="ESTADO CODIGO", warning_col="NOTA EXTRA")

    cons = consolidate(details)
    lm = make_lm(cons)
    write_rows(p["lm"], ["CODIGO", "NOMBRE", "UD", "CANTIDAD REQUERIDA", "GESTIONADO", "PENDIENTE", "REGLA", "NOTA", "ESTADO", "ADVERTENCIA"], lm, "LM", status_col="ESTADO", warning_col="ADVERTENCIA")

    mov = movement_proposal(lm)
    write_rows(p["mov"], ["CODIGO", "NOMBRE", "UD", "PENDIENTE", "EXISTENCIA B1", "MOV PROPUESTO"], mov, "MOV_PROPUESTO")

    rq = rq_proposal(lm, base_date_from_meta())
    write_rows(p["rq"], ["CLASIFICACION", "CODIGO", "NOMBRE", "UD", "CANTIDAD", "FECHA BASE", "DIAS", "FECHA REQUERIDA", "NOTA"], rq, "RQ_PROPUESTA", warning_col="NOTA")
    write_rq_export(p["rq_export"], rq)

    reps = spare_rows(details)
    write_rows(p["spares"], ["CODIGO", "NOMBRE", "CANTIDAD", "UD", "CLASIFICACION", "FRECUENCIA", "NOTA", "OBSERVACION"], reps, "REPUESTOS")

    vals, total, missing_price, eqvals = value_rows(lm, details)
    write_rows(p["value"], ["CODIGO", "NOMBRE", "UD", "CANTIDAD", "PRECIO", "SUBTOTAL"], vals, "VALOR", currency_headers=["PRECIO", "SUBTOTAL"])
    write_rows(p["eqvalue"], ["EQUIPO", "VALOR", "SIN VALOR"], eqvals, "VALOR_EQUIPOS", currency_headers=["VALOR"])

    meta = read_meta()
    meta.update({
        "procesado": True,
        "ultima_procesada": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "valor_total": total,
        "referencias_sin_valor": missing_price,
    })
    write_meta(meta)
    return len(details), len(lm)


@app.route("/")
def home():
    return redirect(url_for("inicio"))


@app.route("/inicio")
def inicio():
    rows_master, _, _ = load_master()
    return render_template("inicio.html", **ctx(
        master_count=len(rows_master),
        history_count=len(history_rows()),
    ))


@app.route("/proyectos-diseno")
def proyectos_diseno_page():
    params = {"view": request.args.get("view", "dashboard")}
    project_id = str(request.args.get("project_id", "")).strip()
    if project_id:
        params["project_id"] = project_id
    return redirect(url_for("proyectos_diseno.index", **params))


@app.route("/configuracion-inicial", methods=["GET", "POST"])
def initial_setup():
    if has_users():
        return redirect(url_for("login"))
    if request.method == "POST":
        try:
            user = create_initial_admin(
                request.form.get("username"),
                request.form.get("name"),
                request.form.get("password"),
            )
            set_user_scope(user["id"])
            ensure_dirs()
            migrated = migrate_legacy_data_to_current_user()
            ensure_simple_workbooks()
            session.clear()
            session.permanent = True
            session["user_id"] = user["id"]
            session["session_version"] = user["session_version"]
            session["login_at"] = time.time()
            session["last_activity"] = time.time()
            session["csrf_token"] = secrets.token_urlsafe(32)
            g.current_user = user
            audit_event(
                "CREAR ADMINISTRADOR INICIAL", "SEGURIDAD",
                "Datos V12 recuperados" if migrated else "Instalación nueva",
                "OK", user, client_ip(),
            )
            flash("Administrador creado correctamente.", "ok")
            return redirect(url_for("inicio"))
        except Exception as exc:
            flash(str(exc), "error")
    return render_template("initial_setup.html", **ctx())


@app.route("/iniciar-sesion", methods=["GET", "POST"])
def login():
    if g.current_user:
        return redirect(url_for("inicio"))
    if request.method == "POST":
        username = request.form.get("username", "").strip().lower()
        ip = client_ip()
        attempted_user = {"username": username or "VACÍO", "name": "", "role": ""}
        if _ip_is_limited(ip):
            audit_event("INGRESO BLOQUEADO", "SEGURIDAD", "Límite temporal por equipo", "BLOQUEADO", attempted_user, ip)
            flash("Demasiados intentos. Espere 15 minutos antes de volver a intentar.", "error")
            return render_template("login.html", **ctx(), username=username), 429

        user, status = authenticate(username, request.form.get("password", ""))
        if status == "ok" and user:
            _clear_ip_failures(ip)
            session.clear()
            session.permanent = True
            session["user_id"] = user["id"]
            session["session_version"] = user["session_version"]
            session["login_at"] = time.time()
            session["last_activity"] = time.time()
            session["csrf_token"] = secrets.token_urlsafe(32)
            g.current_user = user
            set_user_scope(user["id"])
            ensure_dirs()
            ensure_simple_workbooks()
            audit_event("INICIAR SESIÓN", "SEGURIDAD", "Acceso autorizado", "OK", user, ip)
            next_path = request.args.get("next", "")
            parsed_next = urlsplit(next_path)
            if (
                not next_path.startswith("/")
                or next_path.startswith("//")
                or parsed_next.scheme
                or parsed_next.netloc
            ):
                next_path = url_for("inicio")
            return redirect(next_path)

        _record_ip_failure(ip)
        result = "BLOQUEADO" if status == "locked" else "DENEGADO"
        audit_event("INTENTO DE INGRESO", "SEGURIDAD", "Credenciales no válidas o cuenta inactiva", result, attempted_user, ip)
        flash("Usuario o contraseña incorrectos, o acceso temporalmente bloqueado.", "error")
    return render_template(
        "login.html",
        **ctx(),
        username=request.form.get("username", "").strip().lower(),
    )


@app.post("/cerrar-sesion")
def logout():
    user = g.current_user
    if user:
        audit_event("CERRAR SESIÓN", "SEGURIDAD", "Cierre voluntario", "OK", user, client_ip())
        g.audit_logged = True
    session.clear()
    set_user_scope(None)
    return redirect(url_for("login"))


@app.route("/administracion/usuarios")
@admin_required
def admin_users():
    query = request.args.get("q", "").strip()
    return render_template(
        "admin_users.html",
        **ctx(
            users=list_users(),
            audit=audit_rows(500, query),
            audit_query=query,
            admin_page=True,
        ),
    )


@app.post("/administracion/usuarios/crear")
@admin_required
def admin_user_create():
    try:
        target = create_user(
            request.form.get("username"),
            request.form.get("name"),
            request.form.get("password"),
            request.form.get("role", ROLE_COLLABORATOR),
            g.current_user["username"],
        )
        audit_event("CREAR USUARIO", "ADMINISTRACIÓN", f"{target['username']} · {target['role']}", "OK", g.current_user, client_ip())
        g.audit_logged = True
        flash(f"Usuario {target['username']} creado correctamente.", "ok")
    except Exception as exc:
        audit_event("CREAR USUARIO", "ADMINISTRACIÓN", "Operación rechazada", "ERROR", g.current_user, client_ip())
        g.audit_logged = True
        flash(str(exc), "error")
    return redirect(url_for("admin_users"))


@app.post("/administracion/usuarios/<user_id>/rol")
@admin_required
def admin_user_role(user_id):
    if str(user_id) == str(g.current_user["id"]):
        flash("No puede cambiar su propio rol.", "error")
        return redirect(url_for("admin_users"))
    try:
        target = change_role(user_id, request.form.get("role"))
        audit_event("CAMBIAR ROL", "ADMINISTRACIÓN", f"{target['username']} → {target['role']}", "OK", g.current_user, client_ip())
        g.audit_logged = True
        flash(f"Rol de {target['username']} actualizado.", "ok")
    except Exception as exc:
        audit_event("CAMBIAR ROL", "ADMINISTRACIÓN", str(user_id), "ERROR", g.current_user, client_ip())
        g.audit_logged = True
        flash(str(exc), "error")
    return redirect(url_for("admin_users"))


@app.post("/administracion/usuarios/<user_id>/estado")
@admin_required
def admin_user_status(user_id):
    if str(user_id) == str(g.current_user["id"]):
        flash("No puede desactivar su propia cuenta.", "error")
        return redirect(url_for("admin_users"))
    try:
        target = set_user_status(user_id, request.form.get("status"))
        audit_event("CAMBIAR ESTADO", "ADMINISTRACIÓN", f"{target['username']} → {target['status']}", "OK", g.current_user, client_ip())
        g.audit_logged = True
        flash(f"Estado de {target['username']} actualizado.", "ok")
    except Exception as exc:
        audit_event("CAMBIAR ESTADO", "ADMINISTRACIÓN", str(user_id), "ERROR", g.current_user, client_ip())
        g.audit_logged = True
        flash(str(exc), "error")
    return redirect(url_for("admin_users"))


@app.post("/administracion/usuarios/<user_id>/desbloquear")
@admin_required
def admin_user_unlock(user_id):
    try:
        target = unlock_user(user_id)
        audit_event("DESBLOQUEAR USUARIO", "ADMINISTRACIÓN", target["username"], "OK", g.current_user, client_ip())
        g.audit_logged = True
        flash(f"Usuario {target['username']} desbloqueado.", "ok")
    except Exception as exc:
        audit_event("DESBLOQUEAR USUARIO", "ADMINISTRACIÓN", str(user_id), "ERROR", g.current_user, client_ip())
        g.audit_logged = True
        flash(str(exc), "error")
    return redirect(url_for("admin_users"))


@app.post("/administracion/usuarios/<user_id>/password")
@admin_required
def admin_user_password(user_id):
    try:
        target = reset_password(user_id, request.form.get("password"))
        if str(target["id"]) == str(g.current_user["id"]):
            session["session_version"] = target["session_version"]
        audit_event("RESTABLECER CONTRASEÑA", "ADMINISTRACIÓN", target["username"], "OK", g.current_user, client_ip())
        g.audit_logged = True
        flash(f"Contraseña de {target['username']} actualizada. Entréguesela de forma privada.", "ok")
    except Exception as exc:
        audit_event("RESTABLECER CONTRASEÑA", "ADMINISTRACIÓN", str(user_id), "ERROR", g.current_user, client_ip())
        g.audit_logged = True
        flash(str(exc), "error")
    return redirect(url_for("admin_users"))


@app.route("/administracion/auditoria/exportar")
@admin_required
def admin_audit_export():
    audit_event("EXPORTAR AUDITORÍA", "ADMINISTRACIÓN", "AUDITORIA.xlsx", "OK", g.current_user, client_ip())
    g.audit_logged = True
    return send_file(AUDIT_FILE, as_attachment=True, download_name="AUDITORIA_OGA.xlsx")


@app.route("/copias-seguridad")
def backups():
    return render_template("backups.html", **ctx(copies=autosave_rows()))


@app.post("/copias-seguridad/crear")
def backup_create():
    copy = checkpoint("Copia manual", force=True)
    if copy:
        flash("Copia de seguridad creada correctamente.", "ok")
    else:
        flash("No hay un trabajo actual para guardar.", "error")
    return redirect(url_for("backups"))


@app.post("/copias-seguridad/<autosave_id>/recuperar")
def backup_recover(autosave_id):
    try:
        if workspace_dirty():
            checkpoint("Antes de restaurar otra copia", force=True)
        target = recover_autosave(autosave_id)
        checkpoint("Trabajo restaurado", force=True)
        flash(f"Copia recuperada: {target.get('fecha')} — {target.get('motivo')}", "ok")
        return redirect(url_for("cargar"))
    except Exception as e:
        flash(str(e), "error")
        return redirect(url_for("backups"))


@app.post("/copias-seguridad/<autosave_id>/eliminar")
def backup_delete(autosave_id):
    if delete_autosave(autosave_id):
        flash("Copia de seguridad eliminada.", "ok")
    else:
        flash("La copia de seguridad no existe.", "error")
    return redirect(url_for("backups"))


@app.post("/autosave/checkpoint")
def autosave_checkpoint():
    copy = checkpoint("Autoguardado periódico")
    return jsonify({"ok": True, "saved": bool(copy), "fecha": copy.get("fecha") if copy else None})


@app.route("/cargar")
def cargar():
    p = current_paths()
    return render_template("cargar.html", **ctx(
        inventor_loaded=p["general"].exists(),
        b1=read_rows(p["b1_filtered"]),
        lge=lge_rows(),
    ))


@app.post("/cargar/inventor")
def upload_inventor():
    f = request.files.get("file")
    if not f or not f.filename:
        flash("Seleccione una lista de Inventor.", "error")
        return redirect(url_for("cargar"))
    ext = Path(f.filename).suffix.lower()
    if ext not in {".xlsx", ".xls"}:
        flash("La lista de Inventor debe ser Excel (.xlsx o .xls).", "error")
        return redirect(url_for("cargar"))
    target = get_work_dir() / ("INVENTOR_GENERAL" + ext)
    save_upload(f, target)
    # Normalize current general path naming if xls used.
    if target != current_paths()["general"]:
        set_meta("inventor_general_path", target.name)
    checkpoint("Lista de Inventor cargada")
    flash("Lista de Inventor cargada.", "ok")
    return redirect(url_for("cargar"))


@app.post("/cargar/b1")
def upload_b1():
    f = request.files.get("file")
    if not f or not f.filename:
        flash("Seleccione el archivo B1.", "error")
        return redirect(url_for("cargar"))
    ext = Path(f.filename).suffix.lower()
    if ext not in {".xls", ".xlsx"}:
        flash("B1 debe ser un archivo Excel .xls o .xlsx.", "error")
        return redirect(url_for("cargar"))
    raw = get_work_dir() / ("B1_ORIGINAL" + ext)
    save_upload(f, raw)
    try:
        rows = parse_b1(raw)
    except Exception as e:
        raw.unlink(missing_ok=True)
        flash(str(e), "error")
        return redirect(url_for("cargar"))
    write_rows(current_paths()["b1_filtered"], ["CODIGO", "NOMBRE", "UD", "EXISTENCIA"], rows, "B1")
    set_meta("b1_original", raw.name)
    checkpoint("Inventario B1 cargado")
    flash(f"B1 cargado: {len(rows)} referencias. Solo se conservaron CODIGO, NOMBRE, UD y EXISTENCIA.", "ok")
    return redirect(url_for("cargar"))


@app.post("/cargar/lge")
def upload_lge():
    f = request.files.get("file")
    if not f or not f.filename:
        flash("Seleccione el listado de equipos exportado desde Inventor.", "error")
        return redirect(url_for("cargar"))
    ext = Path(f.filename).suffix.lower()
    if ext not in {".xlsx", ".xls"}:
        flash("El listado de equipos debe ser Excel (.xlsx o .xls).", "error")
        return redirect(url_for("cargar"))
    target = get_work_dir() / ("LGE_ORIGEN" + ext)
    save_upload(f, target)
    set_meta("lge_original", target.name)
    try:
        rows = parse_lge_excel(target)
    except Exception as e:
        target.unlink(missing_ok=True)
        flash(str(e), "error")
        return redirect(url_for("cargar"))
    if not rows:
        flash("No se detectaron equipos en el Excel de Inventor.", "error")
        return redirect(url_for("cargar"))
    save_lge(rows)
    checkpoint("Listado de equipos cargado")
    flash(f"Listado de equipos cargado: {len(rows)} equipos/componentes detectados.", "ok")
    return redirect(url_for("cargar"))


@app.post("/lge/<code>/upload")
def upload_equipment_inventor(code):
    rows = lge_rows()
    idx = request.form.get("row", type=int)
    f = request.files.get("file")
    if idx is None or idx < 0 or idx >= len(rows) or not f or not f.filename:
        flash("No fue posible asociar la lista al equipo.", "error")
        return redirect(url_for("cargar"))
    ext = Path(f.filename).suffix.lower()
    if ext not in {".xlsx", ".xls"}:
        flash("La lista de Inventor debe ser Excel.", "error")
        return redirect(url_for("cargar"))
    folder = get_work_dir() / "INVENTOR_EQUIPOS"
    folder.mkdir(parents=True, exist_ok=True)
    fn = f"{idx+1:03d}_{re.sub(r'[^A-Za-z0-9_-]', '_', code)}{ext}"
    save_upload(f, folder / fn)
    rows[idx]["ARCHIVO"] = fn
    save_lge(rows)
    checkpoint(f"Lista cargada para {code}")
    flash(f"Lista cargada para {code}.", "ok")
    return redirect(url_for("cargar"))


@app.post("/lge/remove")
def remove_lge_row():
    rows = lge_rows()
    idx = request.form.get("row", type=int)
    if idx is not None and 0 <= idx < len(rows):
        checkpoint("Antes de eliminar un equipo", force=True)
        fn = str(rows[idx].get("ARCHIVO") or "")
        if fn:
            (get_work_dir() / "INVENTOR_EQUIPOS" / fn).unlink(missing_ok=True)
        removed = rows.pop(idx)
        save_lge(rows)
        checkpoint(f"Equipo {removed.get('CODIGO')} eliminado")
        flash(f"Equipo {removed.get('CODIGO')} eliminado del proceso.", "ok")
    return redirect(url_for("cargar"))


@app.post("/process")
def process_route():
    try:
        dcount, lcount = process_workspace()
        checkpoint("Proceso RQ terminado")
        flash(f"Proceso terminado: {dcount} filas de Inventor y {lcount} referencias consolidadas.", "ok")
        return redirect(url_for("rq"))
    except Exception as e:
        flash(str(e), "error")
        return redirect(url_for("cargar"))


@app.route("/rq")
def rq():
    p = current_paths()
    rqrows = read_rows(p["rq"])
    rq_all = []
    for row in rqrows:
        item = dict(row)
        cat = rq_category_key(row.get("CLASIFICACION"))
        item["CLASIFICACION_LABEL"] = RQ_CATEGORY_LABELS.get(cat, cat)
        rq_all.append(item)
    return render_template("rq.html", **ctx(
        mov=read_rows(p["mov"]),
        rqrows=rqrows,
        rq_all=rq_all,
        rq_groups=group_rq_rows(rqrows),
        externos=externos_manual_rows(),
        trace=read_rows(p["trace"]),
        fecha_base=base_date_from_meta().isoformat(),
    ))


@app.post("/rq/fecha")
def rq_date():
    raw = request.form.get("fecha_base", "")
    try:
        datetime.strptime(raw, "%Y-%m-%d")
        set_meta("fecha_base_rq", raw)
        if current_paths()["lm"].exists():
            process_workspace()
        checkpoint("Fecha base de RQ actualizada")
        flash("Fecha base de RQ actualizada.", "ok")
    except Exception as e:
        flash(f"Fecha no válida: {e}", "error")
    return redirect(url_for("rq"))


@app.post("/rq/externos/agregar")
def externos_add():
    code = clean_code(request.form.get("codigo", ""))
    if not code:
        flash("Escriba un código válido.", "error")
        return redirect(url_for("rq"))
    _, master_map, _ = load_master()
    if code not in master_map:
        flash(f"El código {code} no existe en la Lista Maestra Factory. Verifíquelo antes de agregarlo.", "error")
        return redirect(url_for("rq"))
    codes = externos_manual_codes()
    if code in codes:
        flash(f"El código {code} ya está en EXTERNOS.", "error")
        return redirect(url_for("rq"))
    codes.add(code)
    save_externos_manual(codes)
    checkpoint(f"Código {code} agregado a EXTERNOS")
    flash(f"Código {code} agregado a EXTERNOS.", "ok")
    return redirect(url_for("rq"))


@app.post("/rq/externos/quitar")
def externos_remove():
    code = clean_code(request.form.get("codigo", ""))
    codes = externos_manual_codes()
    if code in codes:
        codes.discard(code)
        save_externos_manual(codes)
        checkpoint(f"Código {code} quitado de EXTERNOS")
        flash(f"Código {code} quitado de EXTERNOS.", "ok")
    return redirect(url_for("rq"))


def load_actual(kind):
    number = request.form.get("number", "").strip()
    f = request.files.get("file")
    if not number or not f or not f.filename:
        raise ValueError("Escriba el número y seleccione el Excel realizado.")
    ext = Path(f.filename).suffix.lower()
    if ext not in {".xls", ".xlsx"}:
        raise ValueError("El archivo realizado debe ser Excel.")
    folder = get_work_dir() / ("MOV_REAL" if kind == "MOV" else "RQ_REAL")
    folder.mkdir(parents=True, exist_ok=True)
    target = folder / f"{kind}_{re.sub(r'[^A-Za-z0-9_-]', '_', number)}{ext}"
    save_upload(f, target)
    actual = parse_code_qty(target)

    if kind == "MOV":
        expected = {clean_code(r.get("CODIGO")): safe_float(r.get("MOV PROPUESTO")) for r in read_rows(current_paths()["mov"])}
    else:
        expected = {clean_code(r.get("CODIGO")): safe_float(r.get("CANTIDAD")) for r in read_rows(current_paths()["rq"])}
    comparison = compare_actual(actual, expected)
    comp_map = {r["CODIGO"]: r for r in comparison}
    trace = read_rows(current_paths()["trace"])
    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    for a in actual:
        code = clean_code(a["CODIGO"])
        c = comp_map.get(code, {})
        trace.append({
            "TIPO": kind, "NUMERO": number, "CODIGO": code, "CANTIDAD": safe_float(a["CANTIDAD"]),
            "ESTADO": c.get("ESTADO", ""), "OBSERVACION": f"Esperado: {c.get('ESPERADO', 0):g}", "FECHA": now
        })
    write_rows(current_paths()["trace"], ["TIPO", "NUMERO", "CODIGO", "CANTIDAD", "ESTADO", "OBSERVACION", "FECHA"], trace, "TRAZABILIDAD", status_col="ESTADO")
    if current_paths()["details"].exists():
        # Reprocess to update LM and remaining RQ after confirmed actual.
        process_workspace()
    return comparison


@app.post("/rq/load-mov")
def load_mov():
    try:
        comp = load_actual("MOV")
        checkpoint("Movimiento real cargado")
        counts = {k: sum(1 for r in comp if r["ESTADO"] == k) for k in ["CORRECTO", "FALTANTE", "SOBRANTE"]}
        flash(f"Movimiento cargado. Correctos: {counts['CORRECTO']}, faltantes: {counts['FALTANTE']}, sobrantes: {counts['SOBRANTE']}.", "ok")
    except Exception as e:
        flash(str(e), "error")
    return redirect(url_for("rq"))


@app.post("/rq/load-rq")
def load_rq():
    try:
        comp = load_actual("RQ")
        checkpoint("RQ real cargada")
        counts = {k: sum(1 for r in comp if r["ESTADO"] == k) for k in ["CORRECTO", "FALTANTE", "SOBRANTE"]}
        flash(f"RQ cargada. Correctos: {counts['CORRECTO']}, faltantes: {counts['FALTANTE']}, sobrantes: {counts['SOBRANTE']}.", "ok")
    except Exception as e:
        flash(str(e), "error")
    return redirect(url_for("rq"))



@app.route("/lm")
def lm():
    rows = read_rows(current_paths()["lm"])
    q = request.args.get("q", "").strip()
    if q:
        nq = normalize_text(q)
        rows = [
            r for r in rows
            if nq in normalize_text(r.get("CODIGO")) or nq in normalize_text(r.get("NOMBRE"))
        ]
    return render_template("lm.html", **ctx(rows=rows, busqueda=q))

@app.route("/repuestos")
def repuestos():
    return render_template("repuestos.html", **ctx(rows=read_rows(current_paths()["spares"])))


@app.route("/valor")
def valor():
    rows = read_rows(current_paths()["value"])
    q = request.args.get("q", "").strip()
    if q:
        nq = normalize_text(q)
        rows = [
            r for r in rows
            if nq in normalize_text(r.get("CODIGO")) or nq in normalize_text(r.get("NOMBRE"))
        ]
    return render_template("valor.html", **ctx(
        rows=rows,
        equipos=read_rows(current_paths()["eqvalue"]),
        total=read_meta().get("valor_total", 0),
        missing=read_meta().get("referencias_sin_valor", 0),
        busqueda=q,
    ))


@app.route("/factory")
def factory():
    q = request.args.get("q", "").strip()
    page = max(request.args.get("page", 1, type=int), 1)
    per_page = 75
    rows, _, rules = load_master()
    if q:
        nq = normalize_text(q)
        rows = [r for r in rows if nq in normalize_text(r.get("CODIGO")) or nq in normalize_text(r.get("NOMBRE"))]
    total = len(rows)
    start = (page - 1) * per_page
    shown = rows[start:start + per_page]
    pages = max((total + per_page - 1) // per_page, 1)
    return render_template("factory.html", **ctx(rows=shown, rules=rules, q=q, page=page, pages=pages, total=total))


@app.post("/factory/new")
@admin_required
def factory_new():
    code = request.form.get("CODIGO", "").strip()
    _, mp, _ = load_master()
    if not re.fullmatch(r"\d+", code):
        flash("El código Factory debe contener únicamente números.", "error")
        return redirect(url_for("factory"))
    if code in mp:
        flash("Ese código ya existe. Los códigos existentes solo se editan directamente en el Excel maestro.", "error")
        return redirect(url_for("factory"))
    vals = {h: request.form.get(h, "").strip() for h in MASTER_REQUIRED}
    if vals.get("NOTA") == "SI" and not vals.get("DESCRIPCION DE LA NOTA"):
        flash("Debe escribir la descripción de la nota cuando NOTA = SI.", "error")
        return redirect(url_for("factory"))
    vals["CODIGO"] = code
    vals["ULT.COSTO"] = safe_float(vals.get("ULT.COSTO"), 0) if vals.get("ULT.COSTO") else ""
    append_master_code(vals)
    flash("Nuevo código agregado como una nueva fila del Excel maestro.", "ok")
    return redirect(url_for("factory", q=code))


@app.post("/factory/reload")
@admin_required
def factory_reload():
    invalidate_master_cache()
    rows, _, _ = load_master(force=True)
    flash(f"Lista Maestra actualizada desde Excel: {len(rows)} códigos cargados.", "ok")
    return redirect(url_for("factory"))


@app.post("/factory/replace")
@admin_required
def factory_replace():
    f = request.files.get("file")
    if not f or not f.filename:
        flash("Seleccione la nueva Lista Maestra Factory.", "error")
        return redirect(url_for("factory"))
    if Path(f.filename).suffix.lower() != ".xlsx":
        flash("La Lista Maestra debe ser .xlsx.", "error")
        return redirect(url_for("factory"))
    temp = DATA_DIR / "_NUEVA_MAESTRA_TEMP.xlsx"
    save_upload(f, temp)
    try:
        missing = validate_master_file(temp)
        if missing:
            temp.unlink(missing_ok=True)
            flash("La nueva Lista Maestra no tiene estas columnas: " + ", ".join(missing), "error")
            return redirect(url_for("factory"))
        backup_master()
        shutil.move(str(temp), str(MASTER_FILE))
        invalidate_master_cache()
        rows, _, _ = load_master(force=True)
        flash(f"Nueva Lista Maestra cargada. {len(rows)} códigos disponibles.", "ok")
    except Exception as e:
        temp.unlink(missing_ok=True)
        flash(str(e), "error")
    return redirect(url_for("factory"))


@app.route("/factory/download")
def factory_download():
    return send_file(MASTER_FILE, as_attachment=True, download_name="LISTA_MAESTRA_FACTORY.xlsx")


@app.route("/calendario")
def calendario():
    cfg = {normalize_text(r.get("CLAVE")): r.get("VALOR") for r in read_rows(CONFIG_FILE)}
    return render_template("calendario.html", **ctx(config=cfg, holidays=read_rows(HOLIDAYS_FILE)))


@app.post("/calendario/config")
@admin_required
def calendario_config():
    rows = [
        {"CLAVE": "DIAS_C.TEMPRANA", "VALOR": max(0, request.form.get("temprana", 30, type=int))},
        {"CLAVE": "DIAS_PRODUCCION", "VALOR": max(0, request.form.get("produccion", 8, type=int))},
        {"CLAVE": "DIAS_ENSAMBLE", "VALOR": max(0, request.form.get("ensamble", 8, type=int))},
    ]
    write_rows(CONFIG_FILE, ["CLAVE", "VALOR"], rows, "CONFIGURACION")
    flash("Tiempos de compra actualizados.", "ok")
    return redirect(url_for("calendario"))


@app.post("/calendario/holiday/add")
@admin_required
def holiday_add():
    raw = request.form.get("fecha", "")
    desc = request.form.get("descripcion", "").strip()
    try:
        datetime.strptime(raw, "%Y-%m-%d")
    except Exception:
        flash("Fecha no válida.", "error")
        return redirect(url_for("calendario"))
    rows = read_rows(HOLIDAYS_FILE)
    rows.append({"FECHA": raw, "DESCRIPCION": desc})
    write_rows(HOLIDAYS_FILE, ["FECHA", "DESCRIPCION"], rows, "FESTIVOS")
    flash("Día no laboral agregado.", "ok")
    return redirect(url_for("calendario"))


@app.post("/calendario/holiday/delete")
@admin_required
def holiday_delete():
    idx = request.form.get("row", type=int)
    rows = read_rows(HOLIDAYS_FILE)
    if idx is not None and 0 <= idx < len(rows):
        rows.pop(idx)
        write_rows(HOLIDAYS_FILE, ["FECHA", "DESCRIPCION"], rows, "FESTIVOS")
    return redirect(url_for("calendario"))


@app.route("/historico")
def historico():
    return render_template("historico.html", **ctx(rows=history_rows()))


@app.post("/archive")
def archive():
    name = request.form.get("name", "").strip()
    if not name:
        flash("Escriba NOMBRE / N.º PROYECTO para archivar.", "error")
        return redirect(request.referrer or url_for("cargar"))
    try:
        archive_workspace(name)
        flash("Trabajo archivado. El área de trabajo quedó limpia.", "ok")
        return redirect(url_for("historico"))
    except Exception as e:
        flash(str(e), "error")
        return redirect(request.referrer or url_for("cargar"))


@app.post("/historico/<history_id>/recover")
def history_recover(history_id):
    action = request.form.get("current_action", "discard")
    try:
        if workspace_dirty():
            if action == "archive":
                name = request.form.get("archive_name", "").strip()
                if not name:
                    raise ValueError("Debe escribir un nombre para archivar el trabajo actual.")
                archive_workspace(name)
            elif action == "discard":
                checkpoint("Antes de descartar para recuperar histórico", force=True)
                clear_workspace()
            else:
                return redirect(url_for("historico"))
        target = recover_history(history_id)
        checkpoint("Histórico recuperado", force=True)
        flash(f"Histórico recuperado: {target.get('NOMBRE')}", "ok")
        return redirect(url_for("cargar"))
    except Exception as e:
        flash(str(e), "error")
        return redirect(url_for("historico"))


@app.post("/historico/<history_id>/delete")
def history_delete(history_id):
    delete_history(history_id)
    flash("Histórico eliminado.", "ok")
    return redirect(url_for("historico"))


@app.post("/new-list")
def new_list():
    saved = checkpoint("Antes de iniciar una nueva lista", force=True)
    clear_workspace()
    if saved:
        flash("Nueva lista iniciada. El trabajo anterior quedó disponible en COPIAS DE SEGURIDAD.", "ok")
    else:
        flash("Nueva lista iniciada. La Lista Maestra Factory se mantiene.", "ok")
    return redirect(url_for("cargar"))


@app.route("/export/<kind>")
def export_file(kind):
    p = current_paths()
    if kind == "rq" and p["rq"].exists():
        grouped_rows = [
            row
            for group in group_rq_rows(read_rows(p["rq"]))
            for row in group["rows"]
        ]
        write_rq_export(p["rq_export"], grouped_rows)
    mapping = {
        "b1": p["b1_filtered"], "lm": p["lm"], "mov": p["mov"], "rq": p["rq_export"],
        "repuestos": p["spares"], "valor": p["value"], "detalle": p["details"], "traza": p["trace"],
    }
    target = mapping.get(kind)
    if not target or not target.exists():
        abort(404)
    return send_file(target, as_attachment=True, download_name=target.name)


@app.route("/export/rq/categoria/<slug>")
def export_rq_category(slug):
    rows = read_rows(current_paths()["rq"])
    group = next((group for group in group_rq_rows(rows) if group["slug"] == slug), None)
    if not group:
        abort(404)
    output = BytesIO()
    write_rq_export(
        output,
        group["rows"],
        include_category_sheets=False,
        sheet_title=group["label"],
    )
    output.seek(0)
    filename = f"RQ_{group['slug'].upper().replace('-', '_')}.xlsx"
    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route("/export/all")
def export_all():
    out = get_work_dir() / "EXPORTACION_COMPLETA.xlsx"
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font
    wb = Workbook()
    wb.remove(wb.active)
    p = current_paths()
    sources = [
        ("LM", p["lm"]), ("MOV", p["mov"]), ("RQ", p["rq"]), ("REPUESTOS", p["spares"]),
        ("VALOR", p["value"]), ("VALOR_EQUIPOS", p["eqvalue"]), ("TRAZABILIDAD", p["trace"]), ("B1", p["b1_filtered"])
    ]
    for name, fp in sources:
        if not fp.exists():
            continue
        rows = read_rows(fp)
        ws = wb.create_sheet(name[:31])
        if rows:
            headers = list(rows[0].keys())
            ws.append(headers)
            for r in rows:
                ws.append([r.get(h, "") for h in headers])
            for c in ws[1]:
                c.fill = PatternFill("solid", fgColor="0B64B4")
                c.font = Font(color="FFFFFF", bold=True)
            ws.freeze_panes = "A2"
    if not wb.sheetnames:
        ws = wb.create_sheet("SIN_DATOS")
        ws["A1"] = "No hay datos procesados."
    wb.save(out)
    return send_file(out, as_attachment=True, download_name="GENERADOR_RQ_EXPORTACION_COMPLETA.xlsx")


if __name__ == "__main__":
    init_app_data()
    from waitress import serve

    host = os.environ.get("OGA_HOST", "127.0.0.1")
    port = int(os.environ.get("OGA_PORT", "5005"))
    threads = max(4, min(int(os.environ.get("OGA_THREADS", "24")), 48))

    def open_browser():
        webbrowser.open_new(f"http://127.0.0.1:{port}")

    if host in {"127.0.0.1", "localhost"} and os.environ.get("OGA_NO_BROWSER", "0") != "1":
        threading.Timer(1.2, open_browser).start()
    print(f"OGA {APP_REV} disponible en http://{host}:{port}")
    serve(app, host=host, port=port, threads=threads)
else:
    init_app_data()
