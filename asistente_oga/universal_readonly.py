from __future__ import annotations

from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import quote
import json
import re
import sqlite3
import threading

from openpyxl import load_workbook

from .readonly import (
    count_opportunities,
    database_ready,
    find_exact,
    normalize,
    search_opportunities,
    stats as opportunity_stats,
)

BASE_DIR = Path(__file__).resolve().parent.parent
DATA_DIR = BASE_DIR / "data"
PROJECT_DB_DIR = BASE_DIR / "proyectos_diseno" / "BASES_DE_DATOS"
LIBRARY_DIR = DATA_DIR / "BIBLIOTECA"
MASTER_FILE = DATA_DIR / "LISTA_MAESTRA_FACTORY.xlsx"
CAP_DB = DATA_DIR / "capacitaciones" / "capacitaciones.db"
SECURITY_DIR = DATA_DIR / "SEGURIDAD"

# Cache de lectura por mtime: evita abrir los mismos XLSX en cada pregunta.
# Nunca escribe al archivo; si otra parte de OGA actualiza un Excel, el cambio
# de mtime invalida automaticamente la copia en memoria.
_EXCEL_CACHE: dict[tuple[str, str, int], tuple[int, int, tuple[dict[str, Any], ...]]] = {}
_EXCEL_CACHE_LOCK = threading.RLock()

COMMON_STOP = {
    "que", "cual", "cuales", "como", "donde", "cuando", "quien", "quienes",
    "hay", "tiene", "tienen", "tenemos", "dame", "dime", "muestra", "muestrame",
    "busca", "buscar", "encuentra", "trae", "traeme", "consulta", "consultar", "revisa",
    "de", "del", "la", "las", "el", "los", "un", "una", "unos", "unas", "en", "con",
    "por", "para", "a", "al", "y", "o", "me", "mi", "mis", "su", "sus", "oga",
    "informacion", "info", "datos", "dato", "registro", "registros", "actual", "actuales",
    "favor", "puedes", "podrias", "quiero", "necesito", "ver", "sobre", "todos", "todas",
    "cuantos", "cuántos", "cuantas", "cuántas", "cantidad", "total",
    "estan", "están", "son", "sea", "sean", "estaba", "estaban",
}

MODULE_WORDS = {
    "diseno": {"diseno", "diseño", "disenador", "diseñador", "disenadores", "diseñadores", "pmp", "etapa", "avance", "bodega", "proyecto de diseno", "proyectos de diseno", "diseno mecanico", "diseño mecanico", "proyectos por empezar", "proyecto por empezar", "adicional", "adicionales", "trabajos adicionales"},
    "lista_maestra": {"lista maestra", "factory", "codigo factory", "costo", "costos", "repuesto", "repuestos"},
    "rq": {"rq", "requisicion", "requisiciones", "mov", "movimiento", "compras tempranas", "rq propuesta", "rq para factory"},
    "planos": {"plano", "planos", "revisor", "revisor de planos", "revisor planos", "planos revisados", "revision de planos", "revisión de planos", "reporte de planos", "reportes de planos", "faltan planos", "falta plano", "planos faltan", "plano faltante", "planos faltantes", "dxf", "distribucion de planos"},
    "biblioteca": {"biblioteca", "referencia", "referencias", "ficha tecnica", "ficha técnica", "modelo 3d", "biblioteca de equipos"},
    "capacitaciones": {"capacitacion", "capacitaciones", "video", "videos", "tutorial", "tutoriales"},
    "auditoria": {"auditoria", "auditoría", "registro de actividades", "actividad de usuario", "quien hizo", "quién hizo"},
    "usuarios": {"usuarios", "usuario del sistema", "usuarios registrados", "usuarios del sistema", "colaboradores registrados"},
    "calendario": {"calendario", "festivo", "festivos", "feriado", "feriados", "dias produccion", "días producción", "dias ensamble", "días ensamble"},
    "backups": {"copia de seguridad", "copias de seguridad", "backup", "backups", "autoguardado", "autoguardados"},
}


def _excel_rows(path: Path, sheet_name: str | None = None, max_rows: int = 30000) -> list[dict[str, Any]]:
    if not path.exists() or not path.is_file():
        return []
    try:
        stat = path.stat()
        cache_key = (str(path.resolve()), str(sheet_name or ""), int(max_rows))
        with _EXCEL_CACHE_LOCK:
            cached = _EXCEL_CACHE.get(cache_key)
            if cached and cached[0] == stat.st_mtime_ns and cached[1] == stat.st_size:
                return [dict(row) for row in cached[2]]
    except OSError:
        return []
    try:
        wb = load_workbook(path, data_only=True, read_only=True)
    except Exception:
        return []
    try:
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
        iterator = ws.iter_rows(values_only=True)
        try:
            headers = [str(value or "").strip() for value in next(iterator)]
        except StopIteration:
            return []
        rows: list[dict[str, Any]] = []
        for values in iterator:
            if len(rows) >= max_rows:
                break
            if not any(value not in (None, "") for value in values):
                continue
            rows.append({headers[i]: values[i] if i < len(values) else "" for i in range(len(headers))})
        try:
            stat = path.stat()
            with _EXCEL_CACHE_LOCK:
                _EXCEL_CACHE[cache_key] = (stat.st_mtime_ns, stat.st_size, tuple(dict(row) for row in rows))
        except OSError:
            pass
        return rows
    finally:
        wb.close()


def _words(text: str, extra_stop: Iterable[str] = ()) -> list[str]:
    stop = COMMON_STOP | {normalize(x) for x in extra_stop}
    parts = re.findall(r"[a-z0-9áéíóúñü._/-]+", normalize(text))
    out: list[str] = []
    for part in parts:
        if len(part) < 2 or part in stop:
            continue
        if part not in out:
            out.append(part)
    return out[:12]


def _contains_all(row: dict[str, Any], terms: list[str], fields: Iterable[str] | None = None) -> bool:
    if not terms:
        return True
    keys = list(fields) if fields else list(row.keys())
    haystack = normalize(" ".join(str(row.get(key) or "") for key in keys))
    return all(term in haystack for term in terms)


def _is_count(text: str) -> bool:
    n = normalize(text)
    return any(word in n.split() for word in ("cuantos", "cuantas", "cantidad", "total"))


def _current_user_dir(context: dict[str, Any]) -> Path | None:
    user_id = str(context.get("_user_id") or "").strip()
    if not user_id or not re.fullmatch(r"[A-Za-z0-9_-]+", user_id):
        return None
    return DATA_DIR / "USUARIOS" / user_id


def _read_json(path: Path) -> dict[str, Any]:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
        return value if isinstance(value, dict) else {}
    except Exception:
        return {}


# ---------------------------------------------------------------------------
# PROYECTOS DE DISENO
# ---------------------------------------------------------------------------

def _design_data() -> dict[str, Any]:
    projects = _excel_rows(PROJECT_DB_DIR / "Proyectos.xlsx")
    designers = _excel_rows(PROJECT_DB_DIR / "Disenadores.xlsx")
    progress = _excel_rows(PROJECT_DB_DIR / "Avance_Proyectos.xlsx")
    equipment = _excel_rows(PROJECT_DB_DIR / "Equipos_Proyectos.xlsx")
    periods = _excel_rows(PROJECT_DB_DIR / "Periodos_Etapas.xlsx")
    revisions = _excel_rows(PROJECT_DB_DIR / "Revisiones_Proyectos.xlsx")
    upcoming = _excel_rows(PROJECT_DB_DIR / "Proyectos_Por_Empezar.xlsx")
    additionals = _excel_rows(PROJECT_DB_DIR / "Adicionales.xlsx")

    designer_map = {str(row.get("id") or ""): str(row.get("nombre") or "") for row in designers}
    progress_by: dict[str, list[dict[str, Any]]] = defaultdict(list)
    equipment_by: dict[str, list[dict[str, Any]]] = defaultdict(list)
    periods_by: dict[str, list[dict[str, Any]]] = defaultdict(list)
    revisions_by: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in progress:
        progress_by[str(row.get("project_id") or "")].append(row)
    for row in equipment:
        equipment_by[str(row.get("project_id") or "")].append(row)
    for row in periods:
        periods_by[str(row.get("project_id") or "")].append(row)
    for row in revisions:
        revisions_by[str(row.get("project_id") or "")].append(row)

    enriched = []
    today = date.today()
    for project in projects:
        pid = str(project.get("id") or "")
        rows = progress_by.get(pid, [])
        total = sum(float(r.get("porcentaje") or 0) for r in rows)
        completed = sum(float(r.get("porcentaje") or 0) for r in rows if bool(r.get("cumplida")))
        avance = round((completed / total * 100) if total else 0, 1)
        pending_stages = []
        for stage in (1, 2, 3):
            sr = [r for r in rows if int(float(r.get("etapa") or 0)) == stage]
            if sr and not all(bool(r.get("cumplida")) for r in sr):
                pending_stages.append(stage)
        current_stage = pending_stages[0] if pending_stages else None
        late = False
        if str(project.get("estado") or "").casefold() != "finalizado" and current_stage:
            ends = []
            for period in periods_by.get(pid, []):
                try:
                    stage = int(float(period.get("etapa") or 0))
                except Exception:
                    continue
                if stage != current_stage:
                    continue
                try:
                    ends.append(date.fromisoformat(str(period.get("fin") or "")[:10]))
                except Exception:
                    pass
            late = bool(ends and max(ends) < today)
        row = dict(project)
        row.update({
            "designer_name": designer_map.get(str(project.get("designer_id") or ""), "Sin diseñador"),
            "avance": avance,
            "etapa_actual": current_stage,
            "retrasado": late,
            "progress_rows": rows,
            "equipment_rows": equipment_by.get(pid, []),
            "period_rows": periods_by.get(pid, []),
            "revision_rows": revisions_by.get(pid, []),
        })
        enriched.append(row)
    return {
        "projects": enriched,
        "designers": designers,
        "upcoming": upcoming,
        "additionals": additionals,
    }


def _design_card(row: dict[str, Any], detailed: bool = False) -> dict[str, Any]:
    status = str(row.get("estado") or "Activo")
    stage = row.get("etapa_actual")
    items = [
        f"Diseñador: {row.get('designer_name') or 'Sin diseñador'}",
        f"Tipo: {row.get('tipo_proyecto') or '-'} · Tamaño: {row.get('tamano_nombre') or '-'}",
        f"PMP: {row.get('pmp_nombre') or '-'} · Bodega: {row.get('bodega') or '-'}",
    ]
    if stage:
        items.append(f"Etapa actual: 0{stage}")
    if row.get("retrasado"):
        items.append("Alerta: fecha programada de la etapa actual vencida")
    if detailed:
        pending = [r for r in row.get("progress_rows") or [] if not bool(r.get("cumplida"))]
        if pending:
            items.append("Pendientes: " + "; ".join(str(r.get("actividad") or "") for r in pending[:5]))
        eq = row.get("equipment_rows") or []
        if eq:
            items.append("Equipos: " + "; ".join(f"{r.get('codigo') or '-'} x{r.get('cantidad') or '-'}" for r in eq[:8]))
    return {
        "type": "design_project",
        "title": f"Proyecto {row.get('numero') or '-'}",
        "subtitle": str(row.get("cliente") or "Sin cliente"),
        "status": status,
        "meta": f"Avance {row.get('avance') or 0:g}%" if isinstance(row.get("avance"), (int, float)) else "",
        "items": items,
        "url": f"/proyectos-diseno?view=projects&project_id={quote(str(row.get('id') or ''))}",
        "action_label": "Abrir proyecto",
        "entity_key": str(row.get("id") or ""),
    }


def _design_answer(message: str) -> dict[str, Any]:
    data = _design_data()
    projects = data["projects"]
    n = normalize(message)

    # Consultas sobre recursos del modulo, no solo sobre la tabla Proyectos.
    if any(x in n for x in ("disenadores", "diseñadores", "equipo de diseno", "equipo de diseño", "carga de diseno", "carga de diseño")) and "proyecto " not in n:
        active_projects = [p for p in projects if normalize(p.get("estado")) != "finalizado"]
        load = Counter(str(p.get("designer_id") or "") for p in active_projects)
        cards = []
        for d in data["designers"]:
            if not bool(d.get("activo", True)):
                continue
            did = str(d.get("id") or "")
            cards.append({
                "type": "designer",
                "title": str(d.get("nombre") or "Diseñador"),
                "subtitle": "Diseño Mecánico",
                "status": "Activo",
                "meta": f"{load.get(did, 0)} proyecto(s) activo(s)",
                "url": "/proyectos-diseno?view=designers",
                "action_label": "Abrir Diseñadores",
                "entity_key": did,
            })
        cards.sort(key=lambda c: (-int(str(c["meta"]).split()[0]), normalize(c["title"])))
        return _result(f"Hay {len(cards)} diseñador(es) activo(s). La carga indicada corresponde a proyectos no finalizados.", cards, "PROYECTOS_DISENO", [str(c.get("entity_key") or "") for c in cards])

    if any(x in n for x in ("por empezar", "proximos proyectos", "próximos proyectos", "proyectos proximos", "proyectos próximos")):
        designer_map = {str(d.get("id") or ""): str(d.get("nombre") or "") for d in data["designers"]}
        rows = data["upcoming"]
        cards = [{
            "type": "upcoming_project",
            "title": f"Proyecto {r.get('numero') or '-'}",
            "subtitle": str(r.get("cliente") or "Sin cliente"),
            "meta": f"Inicio {r.get('fecha_inicio') or '-'} · {r.get('dias_referencia') or '-'} días",
            "items": [f"Diseñador: {designer_map.get(str(r.get('designer_id') or ''), 'Sin diseñador') }"],
            "url": "/proyectos-diseno?view=projects",
            "action_label": "Abrir Proyectos",
            "entity_key": str(r.get("id") or ""),
        } for r in rows[:20]]
        return _result(f"Hay {len(rows)} proyecto(s) registrados por empezar.", cards, "PROYECTOS_DISENO", [str(r.get("id") or "") for r in rows[:20]])

    if any(x in n for x in ("adicional", "adicionales", "trabajos adicionales", "ot adicional")):
        rows = data["additionals"]
        terms = _words(message, extra_stop={"adicional", "adicionales", "trabajo", "trabajos", "ot"})
        if terms:
            rows = [r for r in rows if _contains_all(r, terms)]
        cards = [{
            "type": "design_additional",
            "title": str(r.get("titulo") or "Adicional"),
            "subtitle": str(r.get("proyecto") or ""),
            "meta": f"Solicita: {r.get('solicita') or '-'} · {r.get('dias_duracion') or '-'} día(s)",
            "items": [str(r.get("descripcion") or "")],
            "url": "/proyectos-diseno?view=additionals",
            "action_label": "Abrir Adicionales",
            "entity_key": str(r.get("id") or ""),
        } for r in rows[:20]]
        return _result(f"Encontre {len(rows)} adicional(es) en Diseño Mecánico.", cards, "PROYECTOS_DISENO", [str(r.get("id") or "") for r in rows[:20]])

    match = re.search(r"\bproyecto\s*[:#-]?\s*([a-z0-9._/-]+)", n)
    if match:
        number = match.group(1)
        exact = [r for r in projects if normalize(r.get("numero")) == number]
        if exact:
            return _result(
                f"Encontre {len(exact)} proyecto(s) de Diseño Mecánico con número {number}.",
                [_design_card(r, True) for r in exact[:12]],
                "PROYECTOS_DISENO",
                [str(r.get("id") or "") for r in exact[:20]],
            )

    rows = list(projects)
    if any(x in n for x in ("finalizado", "finalizados", "terminado", "terminados")):
        rows = [r for r in rows if normalize(r.get("estado")) == "finalizado"]
    elif any(x in n for x in ("activo", "activos", "en curso", "trabajando")):
        rows = [r for r in rows if normalize(r.get("estado")) != "finalizado"]
    if any(x in n for x in ("retrasado", "retrasados", "atrasado", "atrasados", "vencido", "vencidos")):
        rows = [r for r in rows if r.get("retrasado")]

    designer_names = [str(d.get("nombre") or "") for d in data["designers"] if d.get("nombre")]
    selected_designer = next((name for name in designer_names if normalize(name) in n), "")
    if not selected_designer:
        # Permitir "proyectos de Naranjo" o "proyectos de Brayan" sin exigir
        # el nombre completo, siempre que la coincidencia sea unica.
        possible = []
        for name in designer_names:
            parts = [p for p in normalize(name).split() if len(p) >= 4]
            if any(re.search(rf"\b{re.escape(part)}\b", n) for part in parts):
                possible.append(name)
        if len(possible) == 1:
            selected_designer = possible[0]
    if selected_designer:
        rows = [r for r in rows if normalize(r.get("designer_name")) == normalize(selected_designer)]

    terms = _words(message, extra_stop={
        "proyecto", "proyectos", "diseno", "diseño", "mecanico", "mecánico", "disenador", "diseñador",
        "disenadores", "diseñadores", "activo", "activos", "finalizado", "finalizados", "retrasado", "retrasados",
        "avance", "etapa", "pmp", "bodega", "mayor", "menor", "mas", "más", "menos", "avanzado", "avanzados", "porcentaje",
    })
    if selected_designer:
        designer_parts = set(_words(selected_designer))
        terms = [t for t in terms if t not in designer_parts]
    if terms:
        rows = [r for r in rows if _contains_all(r, terms, fields=("numero", "cliente", "tipo_proyecto", "pmp_nombre", "bodega", "tamano_nombre", "designer_name", "notas_historico", "estatus_historico"))]

    if any(x in n for x in ("mayor avance", "mas avanzado", "más avanzado", "mayor porcentaje")):
        rows.sort(key=lambda r: float(r.get("avance") or 0), reverse=True)
    elif any(x in n for x in ("menor avance", "menos avanzado", "menor porcentaje")):
        rows.sort(key=lambda r: float(r.get("avance") or 0))
    else:
        rows.sort(key=lambda r: (normalize(r.get("estado")) == "finalizado", str(r.get("numero") or "")))
    total = len(rows)
    if _is_count(message):
        answer = f"Hay {total} proyecto(s) de Diseño Mecánico que cumplen la consulta."
    elif any(x in n for x in ("avance", "como va", "cómo va")) and len(rows) == 1:
        answer = f"El Proyecto {rows[0].get('numero') or '-'} registra {rows[0].get('avance') or 0:g}% de avance."
    else:
        answer = f"Encontre {total} proyecto(s) en Diseño Mecánico. Te muestro hasta 12."
    return _result(answer, [_design_card(r, True) for r in rows[:12]], "PROYECTOS_DISENO", [str(r.get("id") or "") for r in rows[:20]])


# ---------------------------------------------------------------------------
# LISTA MAESTRA / FACTORY
# ---------------------------------------------------------------------------

def _master_answer(message: str) -> dict[str, Any]:
    rows = _excel_rows(MASTER_FILE, max_rows=20000)
    n = normalize(message)
    wants_spares = any(x in n for x in ("repuesto", "repuestos"))
    wants_without_cost = any(x in n for x in ("sin costo", "sin precio", "sin valor"))
    if wants_spares:
        rows = [r for r in rows if normalize(r.get("REPUESTO ") or r.get("REPUESTO")) in {"si", "sí", "true", "1"}]
    if wants_without_cost:
        def _missing_cost(r):
            raw = r.get("ULT.COSTO")
            if raw in (None, ""):
                return True
            try:
                return float(raw) <= 0
            except Exception:
                return True
        rows = [r for r in rows if _missing_cost(r)]
    code_match = re.search(r"\b(\d{6,})\b", n)
    if code_match:
        code = code_match.group(1)
        exact = [r for r in rows if normalize(r.get("CODIGO")) == code]
        if exact:
            rows = exact
        else:
            rows = [r for r in rows if code in normalize(r.get("CODIGO"))]
    else:
        terms = _words(message, extra_stop={"lista", "maestra", "factory", "codigo", "código", "codigos", "códigos", "costo", "costos", "precio", "repuesto", "repuestos", "referencia", "referencias"})
        if terms:
            rows = [r for r in rows if _contains_all(r, terms)]
    # Si una frase cotidiana trae dos palabras y la combinacion exacta no existe,
    # recuperar coincidencias parciales ordenadas por relevancia en vez de fallar.
    if not rows and not code_match and not (wants_spares or wants_without_cost):
        loose_terms = _words(message, extra_stop={"lista", "maestra", "factory", "codigo", "código", "codigos", "códigos", "costo", "costos", "precio", "repuesto", "repuestos", "referencia", "referencias"})
        if loose_terms:
            all_rows = _excel_rows(MASTER_FILE, max_rows=20000)
            scored = []
            for row in all_rows:
                hay = normalize(" ".join(str(v or "") for v in row.values()))
                score = sum(1 for term in loose_terms if term in hay)
                if score:
                    scored.append((score, row))
            scored.sort(key=lambda item: item[0], reverse=True)
            rows = [row for _, row in scored]
    total = len(rows)
    cards = []
    for r in rows[:15]:
        cost = r.get("ULT.COSTO")
        try:
            cost_text = f"$ {float(cost):,.0f}".replace(",", ".")
        except Exception:
            cost_text = str(cost or "Sin costo")
        items = [f"Unidad: {r.get('U') or '-'}", f"Tipo: {r.get('TIPO ') or r.get('TIPO') or '-'}", f"Repuesto: {r.get('REPUESTO ') or '-'}"]
        note = str(r.get("NOTA ") or "").strip()
        desc_note = str(r.get("DESCRIPCION DE LA NOTA ") or "").strip()
        if note and note.upper() not in {"NO", "N/A"}:
            items.append(f"Nota: {desc_note or note}")
        cards.append({
            "type": "master_code",
            "title": f"{r.get('CODIGO') or '-'} · {r.get('NOMBRE') or '-'}",
            "subtitle": "Lista Maestra Factory",
            "meta": cost_text,
            "items": items,
            "url": "/factory",
            "action_label": "Abrir Lista Maestra",
            "entity_key": str(r.get("CODIGO") or ""),
        })
    if _is_count(message):
        answer = f"La Lista Maestra tiene {total} código(s) que cumplen la consulta."
    else:
        answer = f"Encontre {total} código(s) en Lista Maestra Factory. Te muestro hasta 15." if total else "No encontre códigos en Lista Maestra Factory con esa consulta."
    return _result(answer, cards, "LISTA_MAESTRA", [str(r.get("CODIGO") or "") for r in rows[:20]])


# ---------------------------------------------------------------------------
# BIBLIOTECA
# ---------------------------------------------------------------------------

def _library_answer(message: str) -> dict[str, Any]:
    equipments = _excel_rows(LIBRARY_DIR / "equipos.xlsx", max_rows=20000)
    refs = _excel_rows(LIBRARY_DIR / "referencias.xlsx", max_rows=50000)
    costs = _excel_rows(LIBRARY_DIR / "costos.xlsx", max_rows=50000)
    cost_map = {normalize(r.get("referencia")): r.get("costo") for r in costs}
    equip_map = {normalize(r.get("codigo")): str(r.get("nombre") or "") for r in equipments}
    n = normalize(message)
    terms = _words(message, extra_stop={"biblioteca", "equipo", "equipos", "referencia", "referencias", "plano", "planos", "pdf", "ficha", "tecnica", "técnica", "modelo", "3d", "lista", "materiales"})

    # Buscar primero referencias cuando la consulta parece una referencia concreta.
    ref_rows = []
    for row in refs:
        hay = normalize(" ".join(str(row.get(k) or "") for k in ("referencia_codigo", "equipo_codigo", "equipo_nombre", "selecciones_json")))
        if terms and all(t in hay for t in terms):
            ref_rows.append(row)
    if not terms and any(word in n for word in ("referencias", "referencia")):
        ref_rows = refs[:15]

    want_docs = any(word in n for word in ("pdf", "plano", "planos", "ficha", "modelo 3d", "archivo", "documento"))
    if not terms and want_docs:
        if "ficha" in n:
            ref_rows = [r for r in refs if str(r.get("ficha_tecnica") or "").strip()]
        elif "modelo 3d" in n:
            ref_rows = [r for r in refs if str(r.get("modelo3d") or "").strip()]
        elif "lista" in n and "material" in n:
            ref_rows = [r for r in refs if str(r.get("lista_materiales") or "").strip()]
        else:
            ref_rows = [r for r in refs if str(r.get("planos") or "").strip()]
    cards = []
    for row in ref_rows[:15]:
        ref = str(row.get("referencia_codigo") or "")
        eq_code = str(row.get("equipo_codigo") or "")
        items = []
        docs = (("planos", row.get("planos"), "PDF"), ("modelo3d", row.get("modelo3d"), "ZIP"), ("lista_materiales", row.get("lista_materiales"), "XLSM"), ("ficha_tecnica", row.get("ficha_tecnica"), "DOCX"))
        for kind, filename, label in docs:
            if str(filename or "").strip():
                items.append(f"{label}: {filename}")
        cost = cost_map.get(normalize(ref), "")
        try:
            meta = f"$ {float(cost):,.0f}".replace(",", ".") if str(cost).strip() else "Sin costo registrado"
        except Exception:
            meta = str(cost or "Sin costo registrado")
        card = {
            "type": "library_reference",
            "title": ref or "Referencia",
            "subtitle": str(row.get("equipo_nombre") or equip_map.get(normalize(eq_code)) or eq_code),
            "meta": meta,
            "items": items[:6],
            "url": f"/biblioteca/equipos?equipo={quote(eq_code)}",
            "action_label": "Abrir Biblioteca",
            "entity_key": ref,
        }
        # Si pide un documento y hay plano PDF, ofrecerlo de forma directa.
        if want_docs and str(row.get("planos") or "").strip():
            filename = str(row.get("planos"))
            card["view_url"] = f"/biblioteca/uploads/planos/{quote(filename)}"
            card["view_label"] = "Abrir plano PDF"
            card["download_url"] = f"/biblioteca/download/planos/{quote(filename)}"
            card["download_label"] = "Descargar plano"
            card["kind"] = "PDF"
        cards.append(card)

    if not cards and terms:
        scored = []
        for row in refs:
            hay = normalize(" ".join(str(row.get(k) or "") for k in ("referencia_codigo", "equipo_codigo", "equipo_nombre", "selecciones_json")))
            score = sum(1 for term in terms if term in hay)
            if score:
                scored.append((score, row))
        scored.sort(key=lambda item: item[0], reverse=True)
        ref_rows = [row for _, row in scored]
        for row in ref_rows[:15]:
            ref = str(row.get("referencia_codigo") or "")
            eq_code = str(row.get("equipo_codigo") or "")
            cards.append({
                "type": "library_reference",
                "title": ref or "Referencia",
                "subtitle": str(row.get("equipo_nombre") or equip_map.get(normalize(eq_code)) or eq_code),
                "meta": "Coincidencia parcial en Biblioteca",
                "items": [str(row.get("selecciones_json") or "")[:260]],
                "url": f"/biblioteca/equipos?equipo={quote(eq_code)}",
                "action_label": "Abrir Biblioteca",
                "entity_key": ref,
            })
    if cards:
        return _result(f"Encontre {len(ref_rows)} referencia(s) en Biblioteca. Te muestro hasta 15.", cards, "BIBLIOTECA", [str(r.get("referencia_codigo") or "") for r in ref_rows[:20]])

    eq_rows = equipments
    if terms:
        eq_rows = [r for r in equipments if _contains_all(r, terms, fields=("codigo", "nombre"))]
    cards = [{
        "type": "library_equipment",
        "title": f"{r.get('codigo') or '-'} · {r.get('nombre') or '-'}",
        "subtitle": "Equipo de Biblioteca",
        "meta": f"Item {r.get('item') or '-'}",
        "url": f"/biblioteca/equipos?equipo={quote(str(r.get('codigo') or ''))}",
        "action_label": "Abrir Biblioteca",
        "entity_key": str(r.get("codigo") or ""),
    } for r in eq_rows[:15]]
    return _result(
        f"Encontre {len(eq_rows)} equipo(s) en Biblioteca. Te muestro hasta 15." if eq_rows else "No encontre equipos o referencias en Biblioteca con esa consulta.",
        cards, "BIBLIOTECA", [str(r.get("codigo") or "") for r in eq_rows[:20]],
    )


# ---------------------------------------------------------------------------
# CAPACITACIONES
# ---------------------------------------------------------------------------

def _cap_rows() -> list[dict[str, Any]]:
    if not CAP_DB.exists():
        return []
    try:
        uri = CAP_DB.resolve().as_uri() + "?mode=ro"
        conn = sqlite3.connect(uri, uri=True, timeout=3)
        conn.row_factory = sqlite3.Row
        conn.execute("PRAGMA query_only=ON")
        try:
            return [dict(r) for r in conn.execute("SELECT * FROM videos WHERE eliminado_en IS NULL ORDER BY creado_en DESC LIMIT 2000").fetchall()]
        finally:
            conn.close()
    except sqlite3.Error:
        return []


def _training_answer(message: str) -> dict[str, Any]:
    rows = _cap_rows()
    n = normalize(message)
    status_map = {"listo": "LISTO", "listos": "LISTO", "pendiente": "PENDIENTE", "pendientes": "PENDIENTE", "procesando": "PROCESANDO", "error": "ERROR", "errores": "ERROR"}
    target_status = next((value for word, value in status_map.items() if word in n.split()), "")
    if target_status:
        rows = [r for r in rows if str(r.get("estado") or "").upper() == target_status]
    terms = _words(message, extra_stop={"capacitacion", "capacitaciones", "video", "videos", "tutorial", "tutoriales", "listo", "listos", "pendiente", "pendientes", "procesando", "error", "errores"})
    if terms:
        rows = [r for r in rows if _contains_all(r, terms, fields=("nombre", "descripcion", "subido_por_usuario", "subido_por_nombre", "estado"))]
    cards = [{
        "type": "training_video",
        "title": str(r.get("nombre") or "Video"),
        "subtitle": str(r.get("descripcion") or "")[:220],
        "status": str(r.get("estado") or ""),
        "meta": f"Subido por {r.get('subido_por_nombre') or r.get('subido_por_usuario') or '-'} · {str(r.get('creado_en') or '')[:10]}",
        "url": f"/capacitaciones/{quote(str(r.get('id') or ''))}",
        "action_label": "Abrir capacitación",
        "entity_key": str(r.get("id") or ""),
    } for r in rows[:12]]
    answer = f"Hay {len(rows)} video(s) de capacitación que cumplen la consulta." if _is_count(message) else (f"Encontre {len(rows)} video(s) de capacitación. Te muestro hasta 12." if rows else "No encontre videos de capacitación con esa consulta.")
    return _result(answer, cards, "CAPACITACIONES", [str(r.get("id") or "") for r in rows[:20]])


# ---------------------------------------------------------------------------
# RQ / TRABAJO DEL USUARIO
# ---------------------------------------------------------------------------

def _rq_answer(message: str, context: dict[str, Any]) -> dict[str, Any]:
    user_dir = _current_user_dir(context)
    if user_dir is None:
        return _result("No pude identificar el espacio de usuario para consultar RQ.", [], "RQ")
    work = user_dir / "TRABAJO_ACTUAL"
    history_file = user_dir / "HISTORICO.xlsx"
    meta = _read_json(work / "META.json") if work.exists() else {}
    history = _excel_rows(history_file, max_rows=100)
    n = normalize(message)

    if any(word in n for word in ("historico", "histórico", "archivado", "archivados", "anteriores")):
        cards = [{
            "type": "rq_history",
            "title": str(r.get("NOMBRE") or "Trabajo RQ"),
            "subtitle": str(r.get("FECHA") or ""),
            "meta": str(r.get("ARCHIVO") or ""),
            "url": "/historico",
            "action_label": "Abrir Histórico",
            "entity_key": str(r.get("ID") or ""),
        } for r in history[:10]]
        return _result(f"Tu histórico de RQ contiene {len(history)} trabajo(s) archivados.", cards, "RQ", [str(r.get("ID") or "") for r in history[:20]])

    files = sorted([p.name for p in work.iterdir() if p.is_file()]) if work.exists() else []

    # Consultas directas sobre los Excel generados en el trabajo actual.
    workbook_candidates = []
    if any(x in n for x in ("repuesto", "repuestos")):
        workbook_candidates = ["REPUESTOS.xlsx"]
    elif any(x in n for x in ("valor", "costo total", "costos")):
        workbook_candidates = ["VALOR.xlsx", "VALOR_EQUIPOS.xlsx"]
    elif re.search(r"\bmov\b|movimiento", n):
        workbook_candidates = ["MOV_PROPUESTO.xlsx"]
    elif any(x in n for x in ("rq propuesta", "rq para factory", "lineas rq", "líneas rq")):
        workbook_candidates = ["RQ_PROPUESTA.xlsx", "RQ_PARA_FACTORY.xlsx"]
    if workbook_candidates and work.exists():
        table_cards = []
        total_rows = 0
        for filename in workbook_candidates:
            path = work / filename
            rows = _excel_rows(path, max_rows=5000)
            if not rows:
                continue
            total_rows += len(rows)
            for row in rows[:10]:
                pairs = [f"{key}: {value}" for key, value in row.items() if value not in (None, "")][:8]
                title = str(row.get("CODIGO") or row.get("Código") or row.get("ITEM") or row.get("NOMBRE") or filename)
                table_cards.append({"type": "rq_row", "title": title, "subtitle": filename, "items": pairs, "url": "/rq", "action_label": "Abrir RQ"})
        if table_cards:
            return _result(f"Encontre {total_rows} registro(s) en los archivos del RQ actual.", table_cards[:15], "RQ")

    items = []
    if meta:
        items.append(f"Procesado: {'Sí' if meta.get('procesado') else 'No'}")
        if meta.get("ultima_procesada"):
            items.append(f"Último proceso: {meta.get('ultima_procesada')}")
        if meta.get("valor_total") not in (None, ""):
            try:
                items.append(f"Valor total: $ {float(meta.get('valor_total')):,.0f}".replace(",", "."))
            except Exception:
                items.append(f"Valor total: {meta.get('valor_total')}")
        if meta.get("referencias_sin_valor") not in (None, ""):
            items.append(f"Referencias sin valor: {meta.get('referencias_sin_valor')}")
    if files:
        items.append("Archivos: " + ", ".join(files[:12]))
    card = {
        "type": "rq_current",
        "title": "RQ · Trabajo actual",
        "subtitle": "Generador de RQ",
        "status": "Procesado" if meta.get("procesado") else ("Con archivos" if files else "Sin trabajo actual"),
        "items": items,
        "url": "/rq",
        "action_label": "Abrir RQ",
    }
    answer = "Este es el estado de tu trabajo RQ actual." if (meta or files) else "No hay un trabajo RQ activo en tu espacio de usuario."
    return _result(answer, [card], "RQ")


# ---------------------------------------------------------------------------
# REVISOR DE PLANOS
# ---------------------------------------------------------------------------

def _planos_answer(message: str, context: dict[str, Any]) -> dict[str, Any]:
    user_dir = _current_user_dir(context)
    if user_dir is None:
        return _result("No pude identificar el espacio de usuario para consultar el Revisor de Planos.", [], "PLANOS")
    root = user_dir / "PLANOS"
    state = _read_json(root / "estado.json")
    reports = root / "REPORTES"
    report_files = []
    if reports.exists():
        report_files = sorted(
            [p for p in reports.glob("*.xlsx") if p.name != "lista_planos_revisados.xlsx"],
            key=lambda p: p.stat().st_mtime,
            reverse=True,
        )
    n = normalize(message)

    result = state.get("last_result") if isinstance(state.get("last_result"), dict) else {}
    items = result.get("items") if isinstance(result.get("items"), list) else []
    missing = [r for r in items if "FALTA PLANO" in str(r.get("observacion") or "").upper()]
    if any(phrase in n for phrase in ("faltan planos", "falta plano", "planos faltan", "plano faltante", "planos faltantes", "faltantes", "errores del plano", "errores planos")):
        cards = [{
            "type": "drawing_missing",
            "title": str(r.get("codigo") or "Código sin plano"),
            "subtitle": str(r.get("descripcion") or ""),
            "meta": f"Cantidad {r.get('cantidad') or '-'} · Tipo {r.get('tipo') or '-'}",
            "status": "FALTA PLANO",
        } for r in missing[:20]]
        return _result(f"El último plano revisado tiene {len(missing)} ítem(s) marcados como FALTA PLANO.", cards, "PLANOS")

    if any(word in n for word in ("reporte", "reportes", "excel", "historico", "histórico")):
        cards = [{
            "type": "drawing_report",
            "title": p.name,
            "subtitle": datetime.fromtimestamp(p.stat().st_mtime).strftime("%Y-%m-%d %H:%M"),
            "meta": f"{p.stat().st_size / 1024:.0f} KB",
            "view_url": f"/planos/download/report/{quote(p.name)}",
            "view_label": "Descargar reporte",
            "entity_key": p.name,
        } for p in report_files[:15]]
        return _result(f"Encontre {len(report_files)} reporte(s) del Revisor de Planos.", cards, "PLANOS", [p.name for p in report_files[:20]])

    pdf_name = str(state.get("pdf_name") or result.get("pdf_name") or "")
    card_items = []
    if pdf_name:
        card_items.append(f"PDF actual: {pdf_name}")
    if items:
        card_items.append(f"Ítems revisados: {len(items)}")
        card_items.append(f"Faltan planos: {len(missing)}")
    if result.get("pages"):
        card_items.append(f"Páginas: {result.get('pages')}")
    dist = state.get("distribution") if isinstance(state.get("distribution"), dict) else {}
    if dist:
        if dist.get("pdf_name"):
            card_items.append(f"Distribución PDF: {dist.get('pdf_name')}")
        if dist.get("lm_name"):
            card_items.append(f"LM: {dist.get('lm_name')}")
        if dist.get("dxf_name"):
            card_items.append(f"DXF: {dist.get('dxf_name')}")
    return _result(
        "Este es el estado disponible del Revisor de Planos.",
        [{"type": "planos_status", "title": "Revisor de Planos", "subtitle": pdf_name or "Sin PDF actual", "items": card_items, "url": "/planos/", "action_label": "Abrir Revisor"}],
        "PLANOS",
    )


# ---------------------------------------------------------------------------
# AUDITORIA Y USUARIOS (respeta permisos de la app)
# ---------------------------------------------------------------------------

def _audit_answer(message: str, context: dict[str, Any]) -> dict[str, Any]:
    if not bool(context.get("_is_admin")):
        return _result("La auditoría del sistema solo puede ser consultada por administradores.", [], "AUDITORIA")
    rows = _excel_rows(SECURITY_DIR / "AUDITORIA.xlsx", sheet_name="AUDITORIA", max_rows=5000)
    terms = _words(message, extra_stop={"auditoria", "auditoría", "actividad", "actividades", "registro", "quien", "quién", "hizo", "usuario", "usuarios"})
    if terms:
        rows = [r for r in rows if _contains_all(r, terms, fields=("FECHA", "USUARIO", "NOMBRE", "ROL", "ACCION", "MODULO", "DETALLE", "RESULTADO"))]
    rows = list(reversed(rows))
    cards = [{
        "type": "audit",
        "title": f"{r.get('ACCION') or '-'} · {r.get('MODULO') or '-'}",
        "subtitle": f"{r.get('NOMBRE') or r.get('USUARIO') or '-'} · {r.get('FECHA') or '-'}",
        "status": str(r.get("RESULTADO") or ""),
        "items": [str(r.get("DETALLE") or "")[:300]],
    } for r in rows[:20]]
    return _result(f"Encontre {len(rows)} actividad(es) en auditoría que cumplen la consulta.", cards, "AUDITORIA")


def _users_answer(context: dict[str, Any]) -> dict[str, Any]:
    if not bool(context.get("_is_admin")):
        return _result("La lista de usuarios del sistema solo puede ser consultada por administradores.", [], "USUARIOS")
    rows = _excel_rows(SECURITY_DIR / "USUARIOS.xlsx", sheet_name="USUARIOS", max_rows=1000)
    # Nunca exponer hashes, bloqueos internos ni secretos.
    cards = [{
        "type": "user",
        "title": str(r.get("NOMBRE") or r.get("USUARIO") or "Usuario"),
        "subtitle": str(r.get("USUARIO") or ""),
        "status": str(r.get("ESTADO") or ""),
        "meta": str(r.get("ROL") or ""),
    } for r in rows]
    return _result(f"Hay {len(rows)} usuario(s) registrados en Sistema OGA.", cards, "USUARIOS")


# ---------------------------------------------------------------------------
# CALENDARIO / FESTIVOS Y COPIAS (SOLO LECTURA)
# ---------------------------------------------------------------------------

def _calendar_answer(message: str) -> dict[str, Any]:
    holidays = _excel_rows(DATA_DIR / "FESTIVOS.xlsx", max_rows=5000)
    config = _excel_rows(DATA_DIR / "CONFIGURACION.xlsx", max_rows=500)
    today = date.today()
    upcoming = []
    for row in holidays:
        raw = row.get("FECHA") or row.get("fecha")
        try:
            day = raw.date() if isinstance(raw, datetime) else date.fromisoformat(str(raw)[:10])
        except Exception:
            continue
        if day >= today:
            copy = dict(row); copy["_date"] = day
            upcoming.append(copy)
    upcoming.sort(key=lambda r: r["_date"])
    cards = []
    if any(x in normalize(message) for x in ("configuracion", "configuración", "dias produccion", "días producción", "dias ensamble", "días ensamble")):
        cards = [{"type": "config", "title": str(r.get("CLAVE") or "Configuración"), "meta": str(r.get("VALOR") or "")} for r in config]
        return _result(f"Encontre {len(config)} parámetro(s) de calendario/configuración.", cards, "CALENDARIO")
    for row in upcoming[:15]:
        cards.append({"type": "holiday", "title": str(row.get("NOMBRE") or row.get("nombre") or "Festivo"), "subtitle": row["_date"].isoformat(), "url": "/calendario", "action_label": "Abrir Calendario"})
    return _result(f"Hay {len(holidays)} festivo(s) registrados; {len(upcoming)} son desde hoy en adelante.", cards, "CALENDARIO")


def _backups_answer(context: dict[str, Any]) -> dict[str, Any]:
    user_dir = _current_user_dir(context)
    cards = []
    if user_dir:
        auto = user_dir / "AUTOGUARDADO"
        if auto.exists():
            for p in sorted(auto.glob("*.zip"), key=lambda x: x.stat().st_mtime, reverse=True)[:10]:
                cards.append({"type": "backup", "title": p.name, "subtitle": "Autoguardado personal", "meta": f"{p.stat().st_size/1024/1024:.1f} MB"})
    master = DATA_DIR / "BACKUP_MAESTRA"
    if master.exists():
        for p in sorted(master.glob("*.xlsx"), key=lambda x: x.stat().st_mtime, reverse=True)[:10]:
            cards.append({"type": "backup", "title": p.name, "subtitle": "Copia Lista Maestra", "meta": f"{p.stat().st_size/1024/1024:.1f} MB"})
    return _result(f"Encontre {len(cards)} copia(s) recientes disponibles para consulta.", cards[:15], "BACKUPS")


# ---------------------------------------------------------------------------
# RESUMEN GENERAL Y ENRUTAMIENTO
# ---------------------------------------------------------------------------

def _module_detect(message: str, context: dict[str, Any]) -> str | None:
    n = normalize(message)
    path = normalize(context.get("path") or "")
    endpoint = normalize(context.get("endpoint") or "")

    # Si la conversacion ya esta dentro de un modulo, una pregunta corta como
    # "y cuales tienen PDF?" debe conservar ese contexto. Las menciones
    # explicitas de otro modulo que aparecen mas abajo siguen teniendo prioridad
    # cuando son inequívocas.
    last_module = normalize(context.get("assistant_module") or "")
    context_modules = {
        "proyectos_diseno": "diseno",
        "lista_maestra": "lista_maestra",
        "rq": "rq",
        "planos": "planos",
        "biblioteca": "biblioteca",
        "capacitaciones": "capacitaciones",
        "auditoria": "auditoria",
        "usuarios": "usuarios",
        "calendario": "calendario",
        "backups": "backups",
    }

    # Prioridad para evitar ambiguedades: "repuestos de mi RQ" pertenece a RQ,
    # no a Lista Maestra solo por contener la palabra repuestos.
    priority = ("auditoria", "usuarios", "rq", "diseno", "planos", "capacitaciones", "biblioteca", "lista_maestra", "calendario", "backups")
    for module in priority:
        phrases = MODULE_WORDS.get(module, set())
        if any(normalize(phrase) in n for phrase in phrases):
            return module

    if last_module in context_modules:
        # Frases de continuidad; evita arrastrar el contexto cuando el usuario
        # hace una pregunta totalmente nueva y explicita.
        if (
            n.startswith(("y ", "tambien ", "también ", "de esos", "de esas", "cuales", "cuáles", "que ", "qué ", "dame", "muestra", "muestrame", "muéstrame"))
            or len(_words(message)) <= 5
        ):
            return context_modules[last_module]

    # El contexto de la pagina resuelve preguntas cortas como "que hay aqui".
    if "proyectos-diseno" in path or "proyectos_diseno" in endpoint:
        if any(w in n for w in ("proyecto", "proyectos", "avance", "etapa", "cliente", "disenador", "diseñador", "retrasado")):
            return "diseno"
    if "/rq" in path or endpoint.endswith("rq"):
        return "rq"
    if "/planos" in path or endpoint.startswith("planos"):
        return "planos"
    if "/biblioteca" in path or endpoint.startswith("biblioteca"):
        return "biblioteca"
    if "/capacitaciones" in path or endpoint.startswith("capacitaciones"):
        return "capacitaciones"
    if "/factory" in path or endpoint.endswith("factory"):
        return "lista_maestra"
    if "/calendario" in path or endpoint.endswith("calendario"):
        return "calendario"
    if "/admin" in path and "usuario" in n:
        return "usuarios"

    return None



def _global_search(message: str, context: dict[str, Any]) -> dict[str, Any] | None:
    """Busca texto libre en varias fuentes sin escribir ni modificar nada.

    Se usa solo cuando la pregunta no declara un modulo concreto. El objetivo es
    que consultas humanas como "busca bomba de vacio" no queden encerradas en
    Oportunidades: se revisan las fuentes principales y se indica de donde sale
    cada resultado.
    """
    terms = _words(message, extra_stop={
        "proyecto", "proyectos", "sistema", "sistemas", "oga", "modulo", "módulo",
        "resultado", "resultados", "general", "global", "todo", "toda", "todos", "todas",
    })
    if not terms:
        return None

    cards: list[dict[str, Any]] = []
    counts: list[str] = []

    # Oportunidades: el indice SQLite ya contiene criterios, subsistemas,
    # entradas, salidas, equipos y archivos asociados.
    if database_ready():
        tokens = [(None, term) for term in terms]
        opp_rows = search_opportunities(tokens, 5)
        if opp_rows:
            opp_count = count_opportunities(tokens)
            counts.append(f"{opp_count} en Oportunidades")
            for row in opp_rows[:5]:
                cards.append({
                    "type": "opportunity",
                    "title": f"OD · Proyecto {row.get('proyecto') or '-'}",
                    "subtitle": f"Oportunidades · {row.get('cliente') or 'Sin cliente'}",
                    "status": str(row.get("estado") or ""),
                    "meta": f"Radicado {row.get('radicado') or '-'} · ID OD {row.get('id')}",
                    "url": f"/oportunidades-proyecto/?od={row.get('id')}",
                    "action_label": "Abrir OD",
                    "entity_key": str(row.get("id") or ""),
                })

    # Proyectos Diseño, incluyendo codigos de equipos asignados.
    ddata = _design_data()
    dmatches = []
    for row in ddata["projects"]:
        eq_text = " ".join(
            f"{eq.get('codigo') or ''} {eq.get('nombre') or ''} {eq.get('descripcion') or ''}"
            for eq in (row.get("equipment_rows") or [])
        )
        hay = normalize(" ".join([
            str(row.get("numero") or ""), str(row.get("cliente") or ""),
            str(row.get("tipo_proyecto") or ""), str(row.get("pmp_nombre") or ""),
            str(row.get("bodega") or ""), str(row.get("tamano_nombre") or ""),
            str(row.get("designer_name") or ""), str(row.get("notas_historico") or ""), eq_text,
        ]))
        if all(term in hay for term in terms):
            dmatches.append(row)
    if dmatches:
        counts.append(f"{len(dmatches)} en Proyectos Diseño")
        cards.extend(_design_card(row, True) for row in dmatches[:4])

    # Lista Maestra / Factory.
    master = _excel_rows(MASTER_FILE, max_rows=20000)
    mmatches = [row for row in master if _contains_all(row, terms)]
    if mmatches:
        counts.append(f"{len(mmatches)} en Lista Maestra")
        for row in mmatches[:4]:
            cost = row.get("ULT.COSTO")
            try:
                cost_text = f"$ {float(cost):,.0f}".replace(",", ".")
            except Exception:
                cost_text = str(cost or "Sin costo")
            cards.append({
                "type": "master_code",
                "title": f"{row.get('CODIGO') or '-'} · {row.get('NOMBRE') or '-'}",
                "subtitle": "Lista Maestra Factory",
                "meta": cost_text,
                "url": "/factory",
                "action_label": "Abrir Lista Maestra",
                "entity_key": str(row.get("CODIGO") or ""),
            })

    # Biblioteca de equipos y referencias.
    refs = _excel_rows(LIBRARY_DIR / "referencias.xlsx", max_rows=50000)
    rmatches = []
    for row in refs:
        hay = normalize(" ".join(str(row.get(k) or "") for k in ("referencia_codigo", "equipo_codigo", "equipo_nombre", "selecciones_json")))
        if all(term in hay for term in terms):
            rmatches.append(row)
    if rmatches:
        counts.append(f"{len(rmatches)} en Biblioteca")
        for row in rmatches[:4]:
            ref = str(row.get("referencia_codigo") or "")
            eq_code = str(row.get("equipo_codigo") or "")
            cards.append({
                "type": "library_reference",
                "title": ref or "Referencia",
                "subtitle": f"Biblioteca · {row.get('equipo_nombre') or eq_code}",
                "meta": f"Equipo {eq_code or '-'}",
                "url": f"/biblioteca/equipos?equipo={quote(eq_code)}",
                "action_label": "Abrir Biblioteca",
                "entity_key": ref,
            })

    # Capacitaciones.
    videos = _cap_rows()
    vmatches = [row for row in videos if _contains_all(row, terms, fields=("nombre", "descripcion", "subido_por_usuario", "subido_por_nombre", "estado"))]
    if vmatches:
        counts.append(f"{len(vmatches)} en Capacitaciones")
        for row in vmatches[:3]:
            cards.append({
                "type": "training_video",
                "title": str(row.get("nombre") or "Video"),
                "subtitle": "Capacitaciones · " + str(row.get("descripcion") or "")[:160],
                "status": str(row.get("estado") or ""),
                "url": f"/capacitaciones/{quote(str(row.get('id') or ''))}",
                "action_label": "Abrir capacitación",
                "entity_key": str(row.get("id") or ""),
            })

    # Trabajo RQ actual del usuario: revisar solo archivos ya generados.
    user_dir = _current_user_dir(context)
    if user_dir:
        work = user_dir / "TRABAJO_ACTUAL"
        rq_hits = []
        if work.exists():
            for filename in ("RQ_PROPUESTA.xlsx", "RQ_PARA_FACTORY.xlsx", "MOV_PROPUESTO.xlsx", "REPUESTOS.xlsx", "VALOR.xlsx", "VALOR_EQUIPOS.xlsx"):
                path = work / filename
                for row in _excel_rows(path, max_rows=5000):
                    if _contains_all(row, terms):
                        rq_hits.append((filename, row))
                        if len(rq_hits) >= 5:
                            break
                if len(rq_hits) >= 5:
                    break
        if rq_hits:
            counts.append(f"{len(rq_hits)}+ en tu RQ actual" if len(rq_hits) == 5 else f"{len(rq_hits)} en tu RQ actual")
            for filename, row in rq_hits[:3]:
                pairs = [f"{k}: {v}" for k, v in row.items() if v not in (None, "")][:6]
                cards.append({
                    "type": "rq_row",
                    "title": str(row.get("CODIGO") or row.get("Código") or row.get("ITEM") or row.get("NOMBRE") or filename),
                    "subtitle": f"RQ · {filename}",
                    "items": pairs,
                    "url": "/rq",
                    "action_label": "Abrir RQ",
                })

        # Estado del Revisor de Planos actual.
        state = _read_json(user_dir / "PLANOS" / "estado.json")
        result = state.get("last_result") if isinstance(state.get("last_result"), dict) else {}
        pitems = result.get("items") if isinstance(result.get("items"), list) else []
        phits = []
        for row in pitems:
            hay = normalize(" ".join(str(v or "") for v in row.values()))
            if all(term in hay for term in terms):
                phits.append(row)
        if phits:
            counts.append(f"{len(phits)} en tu revisión de planos")
            for row in phits[:3]:
                cards.append({
                    "type": "drawing_item",
                    "title": str(row.get("codigo") or "Ítem de plano"),
                    "subtitle": "Revisor de Planos · " + str(row.get("descripcion") or ""),
                    "status": str(row.get("observacion") or ""),
                    "url": "/planos/",
                    "action_label": "Abrir Revisor",
                })

    if not cards:
        return _result(
            "Busque esa expresion en Oportunidades, Proyectos Diseño, Lista Maestra, Biblioteca, Capacitaciones y tu espacio de trabajo, pero no encontre coincidencias.",
            [], "SISTEMA_OGA",
        )

    summary = "; ".join(counts)
    return _result(
        f"Hice una busqueda global de {' '.join(terms)}. Coincidencias: {summary}. Te muestro una seleccion indicando el modulo de origen.",
        cards[:20], "SISTEMA_OGA",
        [str(card.get("entity_key") or "") for card in cards if card.get("entity_key")][:20],
    )


def _project_overview(message: str) -> dict[str, Any]:
    data = _design_data()
    design_projects = data["projects"]
    design_active = [r for r in design_projects if normalize(r.get("estado")) != "finalizado"]
    cards = [_design_card(r) for r in design_active[:6]]

    opp_total = count_opportunities([]) if database_ready() else 0
    if database_ready():
        for row in search_opportunities([], 6):
            cards.append({
                "type": "opportunity",
                "title": f"OD · Proyecto {row.get('proyecto') or '-'}",
                "subtitle": str(row.get("cliente") or ""),
                "status": str(row.get("estado") or ""),
                "meta": f"Radicado {row.get('radicado') or '-'} · ID OD {row.get('id')}",
                "url": f"/oportunidades-proyecto/?od={row.get('id')}",
                "action_label": "Abrir OD",
            })
    answer = (
        f"En Sistema OGA encuentro {len(design_projects)} proyecto(s) en Diseño Mecánico "
        f"({len(design_active)} activos) y {opp_total} oportunidad(es) de proyecto. "
        "Te muestro una muestra de ambos módulos; si me dices 'proyectos de diseño' u 'oportunidades' puedo filtrar uno solo."
    )
    return _result(answer, cards, "SISTEMA_OGA")


def _system_summary(context: dict[str, Any]) -> dict[str, Any]:
    design = _design_data()
    dprojects = design["projects"]
    opp = opportunity_stats() if database_ready() else {"total": 0, "subsystems": 0}
    master_count = len(_excel_rows(MASTER_FILE, max_rows=30000))
    library_eq = len(_excel_rows(LIBRARY_DIR / "equipos.xlsx", max_rows=30000))
    library_refs = len(_excel_rows(LIBRARY_DIR / "referencias.xlsx", max_rows=50000))
    videos = _cap_rows()
    user_dir = _current_user_dir(context)
    reports = 0
    rq_history = 0
    if user_dir:
        reports_dir = user_dir / "PLANOS" / "REPORTES"
        if reports_dir.exists():
            reports = len([p for p in reports_dir.glob("*.xlsx") if p.name != "lista_planos_revisados.xlsx"])
        rq_history = len(_excel_rows(user_dir / "HISTORICO.xlsx", max_rows=100))
    cards = [
        {"title": "Oportunidades de Proyecto", "meta": f"{opp.get('total', 0)} OD · {opp.get('subsystems', 0)} subsistemas", "url": "/oportunidades-proyecto/", "action_label": "Abrir Oportunidades"},
        {"title": "Proyectos Diseño Mecánico", "meta": f"{len(dprojects)} proyectos · {sum(1 for r in dprojects if normalize(r.get('estado')) != 'finalizado')} activos", "url": "/proyectos-diseno", "action_label": "Abrir Proyectos"},
        {"title": "Lista Maestra Factory", "meta": f"{master_count} códigos", "url": "/factory", "action_label": "Abrir Lista Maestra"},
        {"title": "Biblioteca", "meta": f"{library_eq} equipos · {library_refs} referencias", "url": "/biblioteca/", "action_label": "Abrir Biblioteca"},
        {"title": "Capacitaciones", "meta": f"{len(videos)} videos", "url": "/capacitaciones/", "action_label": "Abrir Capacitaciones"},
        {"title": "Revisor de Planos", "meta": f"{reports} reportes en tu espacio", "url": "/planos/", "action_label": "Abrir Revisor"},
        {"title": "RQ", "meta": f"{rq_history} trabajos archivados en tu espacio", "url": "/rq", "action_label": "Abrir RQ"},
    ]
    return _result(
        "El modo local ya puede consultar los principales módulos de Sistema OGA. Las consultas son estrictamente de solo lectura.",
        cards,
        "SISTEMA_OGA",
    )


def _result(answer: str, cards: list[dict[str, Any]], module: str, entity_keys: list[str] | None = None) -> dict[str, Any]:
    return {
        "ok": True,
        "mode": "READ_ONLY",
        "answer": answer,
        "cards": cards,
        "sources": [],
        "context_update": {"module": module, "entity_keys": list(entity_keys or [])[:20]},
    }


def universal_manifest() -> tuple[dict[str, str], ...]:
    return (
        {"name": "consultar_proyectos_diseno", "mode": "READ_ONLY"},
        {"name": "consultar_lista_maestra", "mode": "READ_ONLY"},
        {"name": "consultar_rq", "mode": "READ_ONLY"},
        {"name": "consultar_revisor_planos", "mode": "READ_ONLY"},
        {"name": "consultar_biblioteca", "mode": "READ_ONLY"},
        {"name": "consultar_capacitaciones", "mode": "READ_ONLY"},
        {"name": "consultar_auditoria", "mode": "READ_ONLY_ADMIN"},
        {"name": "consultar_calendario", "mode": "READ_ONLY"},
        {"name": "consultar_copias", "mode": "READ_ONLY"},
        {"name": "resumen_sistema", "mode": "READ_ONLY"},
    )


def answer_universal(message: str, context: dict[str, Any] | None = None) -> dict[str, Any] | None:
    context = context or {}
    n = normalize(message)

    # Ayuda y consultas globales deben mostrar el alcance real del asistente.
    if any(phrase in n for phrase in ("que puedes hacer", "como funcionas", "para que sirves", "que puedes consultar", "ayuda asistente", "modulos puedes")):
        return _system_summary(context)
    if any(phrase in n for phrase in ("resumen del sistema", "resumen oga", "que hay en oga", "que informacion hay", "que modulos hay")):
        return _system_summary(context)

    module = _module_detect(message, context)
    if module == "diseno":
        return _design_answer(message)
    if module == "lista_maestra":
        return _master_answer(message)
    if module == "rq":
        return _rq_answer(message, context)
    if module == "planos":
        return _planos_answer(message, context)
    if module == "biblioteca":
        return _library_answer(message)
    if module == "capacitaciones":
        return _training_answer(message)
    if module == "auditoria":
        return _audit_answer(message, context)
    if module == "usuarios":
        return _users_answer(context)
    if module == "calendario":
        return _calendar_answer(message)
    if module == "backups":
        return _backups_answer(context)

    # Una consulta completamente generica de proyectos ya no debe asumir OD.
    generic_project_phrases = {
        "que proyectos hay", "que proyectos tenemos", "cuales proyectos hay",
        "cuales son los proyectos", "muestrame los proyectos", "muestra los proyectos",
        "lista los proyectos", "listar proyectos", "proyectos",
    }
    if n.strip(" ?¿!¡.") in generic_project_phrases:
        return _project_overview(message)

    # Numero de proyecto sin aclarar modulo: buscar en Diseño y dejar que el
    # motor de Oportunidades también pueda responder si no aparece aquí.
    match = re.search(r"\bproyecto\s*[:#-]?\s*([a-z0-9._/-]+)", n)
    if match and not any(x in n for x in ("radicado", "oportunidad", "od")):
        number = match.group(1)
        design = [r for r in _design_data()["projects"] if normalize(r.get("numero")) == number]
        opp = find_exact("proyecto", number) if database_ready() else []
        if design and opp:
            cards = [_design_card(r, True) for r in design]
            for row in opp[:8]:
                cards.append({"type": "opportunity", "title": f"OD · Proyecto {row.get('proyecto') or '-'}", "subtitle": str(row.get("cliente") or ""), "status": str(row.get("estado") or ""), "meta": f"Radicado {row.get('radicado') or '-'} · ID OD {row.get('id')}", "url": f"/oportunidades-proyecto/?od={row.get('id')}", "action_label": "Abrir OD"})
            return _result(f"El número {number} aparece tanto en Proyectos Diseño como en Oportunidades. Te muestro ambas fuentes para no asumir una incorrecta.", cards, "SISTEMA_OGA")
        if design:
            return _result(f"Encontre el Proyecto {number} en Diseño Mecánico.", [_design_card(r, True) for r in design], "PROYECTOS_DISENO", [str(r.get("id") or "") for r in design])
        if opp:
            # El motor especializado de Oportunidades devuelve ID, adjuntos y
            # mantiene la referencia para preguntas siguientes como "sus PDF".
            return None

    # Las preguntas que nombran explícitamente Oportunidades deben conservar
    # el parser especializado (criterios, adjuntos, similares y contexto OD).
    if re.search(r"\b(oportunidad|oportunidades|od|radicado|subsistema|subsistemas|oferta|ofertas)\b", n):
        return None

    # En la pantalla global del asistente, texto libre sin modulo se busca en
    # todas las fuentes principales. Dentro de Oportunidades se conserva el
    # parser especializado de OD para no perder su contexto conversacional.
    path = normalize(context.get("path") or "")
    last_module = normalize(context.get("assistant_module") or "")
    if "/oportunidades" not in path and last_module not in {"oportunidades", "oportunidades_proyecto"}:
        global_result = _global_search(message, context)
        if global_result is not None:
            return global_result

    return None
