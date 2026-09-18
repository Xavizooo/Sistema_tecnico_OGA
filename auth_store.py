"""Usuarios y auditoría almacenados en Excel para la instalación local de OGA.

El archivo nunca guarda contraseñas legibles. Solo contiene hashes Argon2id y
permanece fuera de ``static`` para que no pueda descargarse desde el navegador.
"""

from __future__ import annotations

from datetime import datetime, timedelta
from pathlib import Path
import hashlib
import os
import re
import secrets
import threading
import uuid

from argon2 import PasswordHasher
from argon2.exceptions import InvalidHashError, VerificationError, VerifyMismatchError
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


BASE_DIR = Path(__file__).resolve().parent
SECURITY_DIR = BASE_DIR / "data" / "SEGURIDAD"
USERS_FILE = SECURITY_DIR / "USUARIOS.xlsx"
AUDIT_FILE = SECURITY_DIR / "AUDITORIA.xlsx"
SECRET_FILE = SECURITY_DIR / ".app_secret"

ROLE_ADMIN = "ADMINISTRADOR"
ROLE_COLLABORATOR = "COLABORADOR"
VALID_ROLES = {ROLE_ADMIN, ROLE_COLLABORATOR}
STATUS_ACTIVE = "ACTIVO"
STATUS_DISABLED = "DESACTIVADO"

USER_HEADERS = [
    "ID", "USUARIO", "NOMBRE", "PASSWORD_HASH", "ROL", "ESTADO",
    "INTENTOS_FALLIDOS", "VENTANA_FALLOS_DESDE", "BLOQUEADO_HASTA",
    "SESSION_VERSION", "ULTIMO_INGRESO", "CREADO_EN", "ACTUALIZADO_EN", "CREADO_POR",
]
AUDIT_HEADERS = [
    "FECHA", "USUARIO", "NOMBRE", "ROL", "ACCION", "MODULO",
    "DETALLE", "RESULTADO", "IP",
]

_LOCK = threading.RLock()
_PASSWORD_HASHER = PasswordHasher(
    time_cost=2,
    memory_cost=19456,
    parallelism=1,
    hash_len=32,
    salt_len=16,
)
_DUMMY_HASH = _PASSWORD_HASHER.hash(secrets.token_urlsafe(32))
_COMMON_PASSWORDS = {
    "123456789012", "administrador", "administrator", "contrasena123",
    "contraseña123", "password1234", "qwerty123456", "oga123456789",
}


def _now():
    return datetime.now()


def _timestamp(value=None):
    return (value or _now()).strftime("%Y-%m-%d %H:%M:%S")


def _parse_timestamp(value):
    if not value:
        return None
    try:
        return datetime.strptime(str(value), "%Y-%m-%d %H:%M:%S")
    except (TypeError, ValueError):
        return None


def normalize_username(value):
    return str(value or "").strip().lower()


def validate_username(value):
    username = normalize_username(value)
    if not re.fullmatch(r"[a-z0-9][a-z0-9._@-]{2,49}", username):
        raise ValueError("El usuario debe tener entre 3 y 50 caracteres y usar letras, números, punto, guion o @.")
    return username


def validate_password(password, username=""):
    password = str(password or "")
    if len(password) < 12:
        raise ValueError("La contraseña debe tener mínimo 12 caracteres.")
    if len(password) > 128:
        raise ValueError("La contraseña no puede superar 128 caracteres.")
    lowered = password.casefold()
    if lowered in _COMMON_PASSWORDS:
        raise ValueError("La contraseña elegida es demasiado común.")
    normalized_user = normalize_username(username).split("@", 1)[0]
    if len(normalized_user) >= 4 and normalized_user in lowered:
        raise ValueError("La contraseña no debe contener el nombre de usuario.")
    return password


def _safe_excel_text(value, limit):
    text = str(value or "").replace("\x00", "").replace("\r", " ").replace("\n", " ").strip()[:limit]
    if text.startswith(("=", "+", "-", "@")):
        text = "'" + text
    return text


def ensure_security_storage():
    SECURITY_DIR.mkdir(parents=True, exist_ok=True)
    try:
        os.chmod(SECURITY_DIR, 0o700)
    except OSError:
        pass
    if not SECRET_FILE.exists():
        temporary = SECRET_FILE.with_suffix(".tmp")
        temporary.write_text(secrets.token_urlsafe(64), encoding="utf-8")
        os.replace(temporary, SECRET_FILE)
    _restrict_file(SECRET_FILE)
    if not USERS_FILE.exists():
        _write_user_rows([])
    if not AUDIT_FILE.exists():
        _write_audit_rows([])


def get_app_secret():
    ensure_security_storage()
    return SECRET_FILE.read_text(encoding="utf-8").strip()


def _protection_password():
    if not SECRET_FILE.exists():
        SECURITY_DIR.mkdir(parents=True, exist_ok=True)
        SECRET_FILE.write_text(secrets.token_urlsafe(64), encoding="utf-8")
        _restrict_file(SECRET_FILE)
    secret = SECRET_FILE.read_text(encoding="utf-8").strip()
    return hashlib.sha256((secret + "|OGA-EXCEL|").encode("utf-8")).hexdigest()[:24]


def _restrict_file(path):
    try:
        os.chmod(path, 0o600)
    except OSError:
        pass


def _save_atomic(workbook, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.stem}.{uuid.uuid4().hex}.tmp{path.suffix}")
    try:
        workbook.save(temporary)
        os.replace(temporary, path)
        _restrict_file(path)
    finally:
        workbook.close()
        temporary.unlink(missing_ok=True)


def _style_header(sheet):
    fill = PatternFill("solid", fgColor="1E5DA8")
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, cell in enumerate(sheet[1], start=1):
        width = 18
        if cell.value in {"PASSWORD_HASH", "DETALLE"}:
            width = 52
        elif cell.value in {"FECHA", "BLOQUEADO_HASTA", "ULTIMO_INGRESO", "CREADO_EN", "ACTUALIZADO_EN"}:
            width = 21
        sheet.column_dimensions[get_column_letter(index)].width = width


def _protect_workbook(workbook, protected_sheet):
    password = _protection_password()
    protected_sheet.protection.sheet = True
    protected_sheet.protection.set_password(password)
    workbook.security.lockStructure = True
    workbook.security.set_workbook_password(password)


def _write_user_rows(rows):
    workbook = Workbook()
    info = workbook.active
    info.title = "INFORMACION"
    info["A1"] = "ARCHIVO ADMINISTRADO POR LA APLICACIÓN OGA"
    info["A2"] = "No editar manualmente. Las contraseñas se almacenan únicamente como hashes Argon2id."
    info["A1"].font = Font(bold=True, color="1E5DA8", size=13)
    info.column_dimensions["A"].width = 95
    sheet = workbook.create_sheet("USUARIOS")
    sheet.append(USER_HEADERS)
    for row in rows:
        sheet.append([row.get(header, "") for header in USER_HEADERS])
    _style_header(sheet)
    sheet.column_dimensions["D"].hidden = True
    _protect_workbook(workbook, sheet)
    sheet.sheet_state = "veryHidden"
    _save_atomic(workbook, USERS_FILE)


def _load_user_rows():
    ensure_security_storage()
    workbook = load_workbook(USERS_FILE, data_only=True, read_only=True)
    try:
        sheet = workbook["USUARIOS"]
        values = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(values)]
        rows = []
        for values_row in values:
            if not any(value not in (None, "") for value in values_row):
                continue
            row = {headers[index]: values_row[index] if index < len(values_row) else "" for index in range(len(headers))}
            row["USUARIO"] = normalize_username(row.get("USUARIO"))
            row["INTENTOS_FALLIDOS"] = int(row.get("INTENTOS_FALLIDOS") or 0)
            row["SESSION_VERSION"] = int(row.get("SESSION_VERSION") or 1)
            rows.append(row)
        return rows
    finally:
        workbook.close()


def _public_user(row):
    if not row:
        return None
    blocked_until = _parse_timestamp(row.get("BLOQUEADO_HASTA"))
    return {
        "id": str(row.get("ID") or ""),
        "username": normalize_username(row.get("USUARIO")),
        "name": str(row.get("NOMBRE") or "").strip(),
        "role": str(row.get("ROL") or ROLE_COLLABORATOR),
        "status": str(row.get("ESTADO") or STATUS_DISABLED),
        "last_login": str(row.get("ULTIMO_INGRESO") or ""),
        "created_at": str(row.get("CREADO_EN") or ""),
        "created_by": str(row.get("CREADO_POR") or ""),
        "locked_until": _timestamp(blocked_until) if blocked_until and blocked_until > _now() else "",
        "session_version": int(row.get("SESSION_VERSION") or 1),
        "is_admin": str(row.get("ROL")) == ROLE_ADMIN,
        "is_active": str(row.get("ESTADO")) == STATUS_ACTIVE,
    }


def has_users():
    with _LOCK:
        return bool(_load_user_rows())


def list_users():
    with _LOCK:
        rows = _load_user_rows()
        return [_public_user(row) for row in sorted(rows, key=lambda item: (str(item.get("ESTADO")) != STATUS_ACTIVE, normalize_username(item.get("USUARIO"))))]


def get_user_by_id(user_id):
    with _LOCK:
        target = str(user_id or "")
        return next((_public_user(row) for row in _load_user_rows() if str(row.get("ID")) == target), None)


def get_user_by_username(username):
    with _LOCK:
        target = normalize_username(username)
        return next((_public_user(row) for row in _load_user_rows() if normalize_username(row.get("USUARIO")) == target), None)


def create_user(username, name, password, role, created_by):
    with _LOCK:
        username = validate_username(username)
        name = str(name or "").strip()
        if len(name) < 3:
            raise ValueError("Escriba el nombre completo del colaborador.")
        if len(name) > 100 or any(ord(character) < 32 for character in name):
            raise ValueError("El nombre no puede superar 100 caracteres ni contener caracteres de control.")
        if name.lstrip().startswith(("=", "+", "-", "@")):
            raise ValueError("El nombre contiene un inicio no permitido.")
        role = str(role or ROLE_COLLABORATOR).upper()
        if role not in VALID_ROLES:
            raise ValueError("Rol no válido.")
        password = validate_password(password, username)
        rows = _load_user_rows()
        if any(normalize_username(row.get("USUARIO")) == username for row in rows):
            raise ValueError("Ese nombre de usuario ya existe.")
        now = _timestamp()
        row = {
            "ID": uuid.uuid4().hex,
            "USUARIO": username,
            "NOMBRE": name,
            "PASSWORD_HASH": _PASSWORD_HASHER.hash(password),
            "ROL": role,
            "ESTADO": STATUS_ACTIVE,
            "INTENTOS_FALLIDOS": 0,
            "VENTANA_FALLOS_DESDE": "",
            "BLOQUEADO_HASTA": "",
            "SESSION_VERSION": 1,
            "ULTIMO_INGRESO": "",
            "CREADO_EN": now,
            "ACTUALIZADO_EN": now,
            "CREADO_POR": str(created_by or "SISTEMA"),
        }
        rows.append(row)
        _write_user_rows(rows)
        return _public_user(row)


def create_initial_admin(username, name, password):
    """Crea de forma atómica el administrador permanente del primer inicio."""
    with _LOCK:
        if _load_user_rows():
            raise ValueError("La configuración inicial ya fue completada.")
        return create_user(username, name, password, ROLE_ADMIN, "CONFIGURACIÓN INICIAL")


def authenticate(username, password):
    """Valida credenciales y devuelve (usuario, estado).

    Estados: ok, invalid, disabled o locked. La interfaz usa un mensaje
    genérico para no revelar si un usuario existe.
    """
    with _LOCK:
        username = normalize_username(username)
        password = str(password or "")
        if len(username) > 50 or len(password) > 128:
            try:
                _PASSWORD_HASHER.verify(_DUMMY_HASH, password[:128])
            except VerificationError:
                pass
            return None, "invalid"
        rows = _load_user_rows()
        index = next((i for i, row in enumerate(rows) if normalize_username(row.get("USUARIO")) == username), None)
        if index is None:
            try:
                _PASSWORD_HASHER.verify(_DUMMY_HASH, str(password or ""))
            except VerificationError:
                pass
            return None, "invalid"

        row = rows[index]
        if str(row.get("ESTADO")) != STATUS_ACTIVE:
            return None, "disabled"

        now = _now()
        blocked_until = _parse_timestamp(row.get("BLOQUEADO_HASTA"))
        if blocked_until and blocked_until > now:
            return None, "locked"

        try:
            valid = _PASSWORD_HASHER.verify(str(row.get("PASSWORD_HASH") or ""), password)
        except (VerifyMismatchError, InvalidHashError, VerificationError):
            valid = False

        if not valid:
            window_start = _parse_timestamp(row.get("VENTANA_FALLOS_DESDE"))
            if not window_start or now - window_start > timedelta(minutes=15):
                attempts = 1
                window_start = now
            else:
                attempts = int(row.get("INTENTOS_FALLIDOS") or 0) + 1
            row["INTENTOS_FALLIDOS"] = attempts
            row["VENTANA_FALLOS_DESDE"] = _timestamp(window_start)
            if attempts >= 5:
                row["BLOQUEADO_HASTA"] = _timestamp(now + timedelta(minutes=15))
            row["ACTUALIZADO_EN"] = _timestamp(now)
            _write_user_rows(rows)
            return None, "locked" if attempts >= 5 else "invalid"

        row["INTENTOS_FALLIDOS"] = 0
        row["VENTANA_FALLOS_DESDE"] = ""
        row["BLOQUEADO_HASTA"] = ""
        row["ULTIMO_INGRESO"] = _timestamp(now)
        row["ACTUALIZADO_EN"] = _timestamp(now)
        if _PASSWORD_HASHER.check_needs_rehash(str(row.get("PASSWORD_HASH"))):
            row["PASSWORD_HASH"] = _PASSWORD_HASHER.hash(str(password))
        _write_user_rows(rows)
        return _public_user(row), "ok"


def reset_password(user_id, new_password):
    with _LOCK:
        rows = _load_user_rows()
        index = next((i for i, row in enumerate(rows) if str(row.get("ID")) == str(user_id)), None)
        if index is None:
            raise ValueError("Usuario no encontrado.")
        row = rows[index]
        password = validate_password(new_password, row.get("USUARIO"))
        row["PASSWORD_HASH"] = _PASSWORD_HASHER.hash(password)
        row["SESSION_VERSION"] = int(row.get("SESSION_VERSION") or 1) + 1
        row["INTENTOS_FALLIDOS"] = 0
        row["VENTANA_FALLOS_DESDE"] = ""
        row["BLOQUEADO_HASTA"] = ""
        row["ACTUALIZADO_EN"] = _timestamp()
        _write_user_rows(rows)
        return _public_user(row)


def _active_admin_count(rows):
    return sum(1 for row in rows if row.get("ROL") == ROLE_ADMIN and row.get("ESTADO") == STATUS_ACTIVE)


def change_role(user_id, role):
    with _LOCK:
        role = str(role or "").upper()
        if role not in VALID_ROLES:
            raise ValueError("Rol no válido.")
        rows = _load_user_rows()
        index = next((i for i, row in enumerate(rows) if str(row.get("ID")) == str(user_id)), None)
        if index is None:
            raise ValueError("Usuario no encontrado.")
        row = rows[index]
        if row.get("ROL") == ROLE_ADMIN and role != ROLE_ADMIN and row.get("ESTADO") == STATUS_ACTIVE and _active_admin_count(rows) <= 1:
            raise ValueError("No se puede cambiar el rol del último administrador activo.")
        row["ROL"] = role
        row["SESSION_VERSION"] = int(row.get("SESSION_VERSION") or 1) + 1
        row["ACTUALIZADO_EN"] = _timestamp()
        _write_user_rows(rows)
        return _public_user(row)


def set_user_status(user_id, status):
    with _LOCK:
        status = str(status or "").upper()
        if status not in {STATUS_ACTIVE, STATUS_DISABLED}:
            raise ValueError("Estado no válido.")
        rows = _load_user_rows()
        index = next((i for i, row in enumerate(rows) if str(row.get("ID")) == str(user_id)), None)
        if index is None:
            raise ValueError("Usuario no encontrado.")
        row = rows[index]
        if row.get("ROL") == ROLE_ADMIN and row.get("ESTADO") == STATUS_ACTIVE and status != STATUS_ACTIVE and _active_admin_count(rows) <= 1:
            raise ValueError("No se puede desactivar el último administrador activo.")
        row["ESTADO"] = status
        row["SESSION_VERSION"] = int(row.get("SESSION_VERSION") or 1) + 1
        if status == STATUS_ACTIVE:
            row["INTENTOS_FALLIDOS"] = 0
            row["VENTANA_FALLOS_DESDE"] = ""
            row["BLOQUEADO_HASTA"] = ""
        row["ACTUALIZADO_EN"] = _timestamp()
        _write_user_rows(rows)
        return _public_user(row)


def unlock_user(user_id):
    with _LOCK:
        rows = _load_user_rows()
        index = next((i for i, row in enumerate(rows) if str(row.get("ID")) == str(user_id)), None)
        if index is None:
            raise ValueError("Usuario no encontrado.")
        row = rows[index]
        row["INTENTOS_FALLIDOS"] = 0
        row["VENTANA_FALLOS_DESDE"] = ""
        row["BLOQUEADO_HASTA"] = ""
        row["ACTUALIZADO_EN"] = _timestamp()
        _write_user_rows(rows)
        return _public_user(row)


def _write_audit_rows(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "AUDITORIA"
    sheet.append(AUDIT_HEADERS)
    for row in rows:
        sheet.append([row.get(header, "") for header in AUDIT_HEADERS])
    _style_header(sheet)
    _protect_workbook(workbook, sheet)
    _save_atomic(workbook, AUDIT_FILE)


def _load_audit_rows():
    ensure_security_storage()
    workbook = load_workbook(AUDIT_FILE, data_only=True, read_only=True)
    try:
        sheet = workbook["AUDITORIA"]
        values = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(values)]
        rows = []
        for values_row in values:
            if not any(value not in (None, "") for value in values_row):
                continue
            rows.append({headers[index]: values_row[index] if index < len(values_row) else "" for index in range(len(headers))})
        return rows
    finally:
        workbook.close()


def audit_event(action, module, detail="", result="OK", user=None, ip=""):
    with _LOCK:
        user = user or {}
        rows = _load_audit_rows()
        rows.append({
            "FECHA": _timestamp(),
            "USUARIO": _safe_excel_text(user.get("username", "DESCONOCIDO"), 80),
            "NOMBRE": _safe_excel_text(user.get("name", ""), 100),
            "ROL": _safe_excel_text(user.get("role", ""), 40),
            "ACCION": _safe_excel_text(action, 80),
            "MODULO": _safe_excel_text(module, 80),
            "DETALLE": _safe_excel_text(detail, 500),
            "RESULTADO": _safe_excel_text(result, 40),
            "IP": _safe_excel_text(ip, 80),
        })
        _write_audit_rows(rows)


def audit_rows(limit=500, query=""):
    with _LOCK:
        rows = list(reversed(_load_audit_rows()))
        query = str(query or "").strip().casefold()
        if query:
            rows = [row for row in rows if query in " ".join(str(value or "") for value in row.values()).casefold()]
        safe_limit = max(1, min(int(limit or 500), 5000))
        return rows[:safe_limit]
