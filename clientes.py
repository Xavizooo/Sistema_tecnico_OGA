from __future__ import annotations

from datetime import datetime
import json
import os
from pathlib import Path
import re
import threading
import unicodedata
import uuid

from flask import Blueprint, flash, redirect, render_template, request, url_for
from openpyxl import load_workbook
import xlrd

from storage import DATA_DIR


bp = Blueprint("clientes", __name__, url_prefix="/clientes")

MODULE_DIR = Path(__file__).resolve().parent
PAISES_FILE = MODULE_DIR / "data" / "paises_y_10_ciudades_importantes.json"
CLIENTES_DIR = DATA_DIR / "CLIENTES"
CLIENTES_FILE = CLIENTES_DIR / "clientes.json"
_CLIENTES_LOCK = threading.RLock()


def _load_paises() -> dict[str, list[str]]:
    if not PAISES_FILE.exists():
        raise RuntimeError(f"No se encontró el archivo de países: {PAISES_FILE}")

    with PAISES_FILE.open("r", encoding="utf-8") as fh:
        data = json.load(fh)

    data.pop("_meta", None)
    return {
        str(pais): [str(ciudad) for ciudad in ciudades]
        for pais, ciudades in data.items()
        if isinstance(ciudades, list)
    }


def _ensure_clientes_file() -> None:
    CLIENTES_DIR.mkdir(parents=True, exist_ok=True)
    if not CLIENTES_FILE.exists():
        _save_clientes([])


def _load_clientes() -> list[dict]:
    _ensure_clientes_file()
    with _CLIENTES_LOCK:
        try:
            with CLIENTES_FILE.open("r", encoding="utf-8") as fh:
                data = json.load(fh)
            return data if isinstance(data, list) else []
        except (json.JSONDecodeError, OSError):
            return []


def _save_clientes(rows: list[dict]) -> None:
    CLIENTES_DIR.mkdir(parents=True, exist_ok=True)
    temp = CLIENTES_FILE.with_suffix(".tmp")
    with _CLIENTES_LOCK:
        with temp.open("w", encoding="utf-8") as fh:
            json.dump(rows, fh, ensure_ascii=False, indent=2)
        os.replace(temp, CLIENTES_FILE)


def _filtered_clientes(rows: list[dict], query: str) -> list[dict]:
    query = (query or "").strip().casefold()
    if not query:
        return rows

    fields = ("nombre", "pais", "ciudad", "nota", "nit", "tipo_industria")
    return [
        row
        for row in rows
        if any(query in str(row.get(field, "")).casefold() for field in fields)
    ]



_EXCEL_HEADER_ALIASES = {
    "nombre": {
        "cliente razon social",
        "cliente",
        "razon social",
        "nombre",
        "nombre cliente",
    },
    "nit": {
        "tax id nit",
        "tax id",
        "nit",
        "identificacion",
        "identificacion nit",
    },
    "tipo_industria": {
        "tipo de industria",
        "tipo industria",
        "industria",
    },
    "ciudad": {"ciudad", "city"},
    "pais": {"pais", "country"},
    "nota": {"nota", "notas", "observacion", "observaciones"},
}
_EXCEL_REQUIRED_FIELDS = {"nombre", "nit", "ciudad", "pais"}


def _header_key(value) -> str:
    text = str(value or "").strip().casefold()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return " ".join(text.split())


def _excel_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "SI" if value else "NO"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _excel_nit(value) -> str:
    # El NIT puede ser alfanumerico. Se conserva tal como viene del Excel
    # (salvo espacios exteriores) para soportar identificaciones extranjeras
    # y formatos propios de cada cliente.
    return _excel_text(value)


def _map_excel_headers(row) -> dict[str, int]:
    mapped: dict[str, int] = {}
    for index, value in enumerate(row):
        key = _header_key(value)
        if not key:
            continue
        for field, aliases in _EXCEL_HEADER_ALIASES.items():
            if key in aliases and field not in mapped:
                mapped[field] = index
                break
    return mapped


def _parse_excel_rows(rows) -> tuple[list[dict], list[str]] | None:
    iterator = iter(rows)
    header_map: dict[str, int] | None = None

    for row_number, row in enumerate(iterator, start=1):
        if row_number > 25:
            break
        candidate = _map_excel_headers(row)
        if _EXCEL_REQUIRED_FIELDS.issubset(candidate):
            header_map = candidate
            break

    if header_map is None:
        return None

    parsed: list[dict] = []
    issues: list[str] = []
    data_row_number = row_number

    for row in iterator:
        data_row_number += 1
        values = list(row)
        if not any(_excel_text(value) for value in values):
            continue

        def cell(field: str):
            index = header_map.get(field)
            return values[index] if index is not None and index < len(values) else None

        nombre = _excel_text(cell("nombre"))
        nit = _excel_nit(cell("nit"))
        pais = _excel_text(cell("pais"))
        ciudad = _excel_text(cell("ciudad"))
        tipo_industria = _excel_text(cell("tipo_industria"))
        nota = _excel_text(cell("nota"))

        missing = []
        if not nombre:
            missing.append("cliente")
        if not nit:
            missing.append("NIT")
        if not pais:
            missing.append("país")
        if not ciudad:
            missing.append("ciudad")
        if missing:
            issues.append(f"Fila {data_row_number}: faltan {', '.join(missing)}.")
            continue
        if len(nombre) > 160:
            issues.append(f"Fila {data_row_number}: el nombre supera 160 caracteres.")
            continue
        if len(nit) > 30:
            issues.append(f"Fila {data_row_number}: el NIT supera 30 caracteres.")
            continue
        if len(tipo_industria) > 120:
            issues.append(f"Fila {data_row_number}: el tipo de industria supera 120 caracteres.")
            continue

        parsed.append(
            {
                "nombre": nombre,
                "nit": nit,
                "tipo_industria": tipo_industria,
                "ciudad": ciudad,
                "pais": pais,
                "nota": nota,
            }
        )

    return parsed, issues


def _parse_clientes_excel(file_storage) -> tuple[list[dict], list[str]]:
    filename = str(file_storage.filename or "").strip()
    extension = Path(filename).suffix.casefold()
    if extension not in {".xlsx", ".xls"}:
        raise ValueError("El archivo debe ser un Excel .xlsx o .xls.")

    try:
        if extension == ".xlsx":
            file_storage.stream.seek(0)
            workbook = load_workbook(file_storage.stream, read_only=True, data_only=True)
            for sheet in workbook.worksheets:
                result = _parse_excel_rows(sheet.iter_rows(values_only=True))
                if result is not None:
                    return result
        else:
            file_storage.stream.seek(0)
            workbook = xlrd.open_workbook(file_contents=file_storage.stream.read())
            for index in range(workbook.nsheets):
                sheet = workbook.sheet_by_index(index)
                result = _parse_excel_rows(
                    (sheet.row_values(row_index) for row_index in range(sheet.nrows))
                )
                if result is not None:
                    return result
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f"No fue posible leer el Excel: {exc}") from exc

    raise ValueError(
        "No se reconocieron los encabezados. Use al menos: CLIENTE (RAZON SOCIAL), "
        "TAX ID/NIT, CIUDAD y PAIS. TIPO DE INDUSTRIA también se importa cuando está presente."
    )

def _normalize_view(raw: str) -> str:
    return "nuevo" if str(raw or "").strip().lower() == "nuevo" else "lista"


def _render(*, form_data: dict | None = None, status: int = 200):
    view = _normalize_view(request.args.get("view", "lista"))
    paises = _load_paises()
    all_clientes = _load_clientes()
    query = request.args.get("q", "").strip()

    clientes = sorted(
        _filtered_clientes(all_clientes, query),
        key=lambda row: str(row.get("nombre", "")).casefold(),
    )

    response = render_template(
        "clientes.html",
        clientes=clientes,
        total_clientes=len(all_clientes),
        paises=paises,
        query=query,
        form_data=form_data or {},
        clientes_page=True,
        clientes_view=view,
    )
    return response, status


@bp.get("/")
def index():
    return _render()


@bp.post("/nuevo")
def create():
    nombre = request.form.get("nombre", "").strip()
    pais = request.form.get("pais", "").strip()
    ciudad = request.form.get("ciudad", "").strip()
    nota = request.form.get("nota", "").strip()
    nit = request.form.get("nit", "").strip()
    tipo_industria = request.form.get("tipo_industria", "").strip()

    form_data = {
        "nombre": nombre,
        "pais": pais,
        "ciudad": ciudad,
        "nota": nota,
        "nit": nit,
        "tipo_industria": tipo_industria,
    }

    errors: list[str] = []
    paises = _load_paises()

    if not nombre:
        errors.append("El nombre del cliente es obligatorio.")
    elif len(nombre) > 160:
        errors.append("El nombre del cliente no puede superar 160 caracteres.")

    if pais not in paises:
        errors.append("Seleccione un país válido.")
    elif ciudad not in paises[pais]:
        errors.append("Seleccione una ciudad válida para el país indicado.")

    if not nota:
        errors.append("Tienes que poner una nota")
    elif len(nombre) > 160:
            errors.append("El nombre del cliente no puede superar 160 caracteres.")

    if tipo_industria and len(tipo_industria) > 120:
        errors.append("El tipo de industria no puede superar 120 caracteres.")

    if not nit:
        errors.append("El NIT es obligatorio.")
    elif len(nit) > 30:
        errors.append("El NIT no puede superar 30 caracteres.")

    rows = _load_clientes()
    if nit and any(str(row.get("nit", "")) == nit for row in rows):
        errors.append("Ya existe un cliente registrado con ese NIT.")

    if errors:
        for message in errors:
            flash(message, "error")
        # Fuerza la vista de formulario cuando hay error de validación.
        # Se renderiza directamente para conservar los valores ingresados.
        paises = _load_paises()
        all_clientes = _load_clientes()
        response = render_template(
            "clientes.html",
            clientes=sorted(all_clientes, key=lambda row: str(row.get("nombre", "")).casefold()),
            total_clientes=len(all_clientes),
            paises=paises,
            query="",
            form_data=form_data,
            clientes_page=True,
            clientes_view="nuevo",
        )
        return response, 400

    rows.append(
        {
            "id": uuid.uuid4().hex,
            "nombre": nombre,
            "pais": pais,
            "ciudad": ciudad,
            "nota": nota,
            "nit": nit,
            "tipo_industria": tipo_industria,
            "creado": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        }
    )
    _save_clientes(rows)

    flash(f"Cliente {nombre} registrado correctamente.", "ok")
    return redirect(url_for("clientes.index", view="lista"))

@bp.post("/subida-masiva")
def bulk_upload():
    archivo = request.files.get("archivo")
    if archivo is None or not str(archivo.filename or "").strip():
        flash("Seleccione un archivo Excel para la subida masiva.", "error")
        return redirect(url_for("clientes.index", view="lista"))

    try:
        imported_rows, issues = _parse_clientes_excel(archivo)
    except ValueError as exc:
        flash(str(exc), "error")
        return redirect(url_for("clientes.index", view="lista"))

    rows = _load_clientes()
    existing_nits = {
        str(row.get("nit", "")).strip()
        for row in rows
        if str(row.get("nit", "")).strip()
    }
    imported_nits: set[str] = set()
    duplicate_count = 0
    imported_count = 0
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    for item in imported_rows:
        nit = item["nit"]
        if nit in existing_nits or nit in imported_nits:
            duplicate_count += 1
            continue

        rows.append(
            {
                "id": uuid.uuid4().hex,
                "nombre": item["nombre"],
                "pais": item["pais"],
                "ciudad": item["ciudad"],
                "nota": item["nota"],
                "nit": nit,
                "tipo_industria": item["tipo_industria"],
                "creado": timestamp,
            }
        )
        imported_nits.add(nit)
        imported_count += 1

    if imported_count:
        _save_clientes(rows)
        message = f"Subida masiva completada: {imported_count} cliente"
        message += " importado" if imported_count == 1 else "s importados"
        if duplicate_count:
            message += f"; {duplicate_count} duplicado" + (" omitido" if duplicate_count == 1 else "s omitidos")
        if issues:
            message += f"; {len(issues)} fila" + (" omitida" if len(issues) == 1 else "s omitidas") + " por datos inválidos"
        flash(message + ".", "ok")
    else:
        omitted = duplicate_count + len(issues)
        if omitted:
            flash(f"No se importaron clientes. Se omitieron {omitted} filas.", "error")
        else:
            flash("El Excel no contiene filas de clientes para importar.", "error")

    for issue in issues[:5]:
        flash(issue, "error")
    if len(issues) > 5:
        flash(f"Hay {len(issues) - 5} observaciones adicionales no mostradas.", "error")

    return redirect(url_for("clientes.index", view="lista"))

