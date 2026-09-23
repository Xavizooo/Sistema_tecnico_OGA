from __future__ import annotations

import math
import os
import re
import tempfile
import threading
import unicodedata
import uuid
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from flask import Blueprint, flash, jsonify, redirect, render_template, request, url_for
from openpyxl import Workbook, load_workbook

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data" / "INDUSTRIA"
DATA_DIR.mkdir(parents=True, exist_ok=True)
_LOCK = threading.RLock()
_ALLOWED_UPLOADS = {".xlsx", ".xlsm"}

bp = Blueprint("industria", __name__)


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _norm(value: Any) -> str:
    txt = unicodedata.normalize("NFKD", _text(value)).encode("ascii", "ignore").decode("ascii")
    txt = re.sub(r"[^A-Z0-9]+", " ", txt.upper()).strip()
    return re.sub(r"\s+", " ", txt)


@dataclass(frozen=True)
class Field:
    key: str
    label: str
    code: str = ""
    kind: str = "text"  # text | number | textarea
    source_col: int | None = None


@dataclass(frozen=True)
class Section:
    slug: str
    title: str
    sheet: str
    filename: str
    fields: tuple[Field, ...]
    source_start_row: int
    description: str


SECTIONS: dict[str, Section] = {
    "puntos-origen": Section(
        slug="puntos-origen",
        title="Tipos de punto de origen",
        sheet="TIPOS DE PUNTOS DE ORIGEN",
        filename="TIPOS_DE_PUNTOS_DE_ORIGEN.xlsx",
        source_start_row=3,
        description="Catálogo de puntos de origen y destino, cantidades y restricciones de altura.",
        fields=(
            Field("tipo_origen", "Tipo de punto de origen", "PO_01", source_col=1),
            Field("cantidad_origen", "Cantidad de puntos de origen", "PO_02", "number", 2),
            Field("restriccion_origen", "Restricción de altura (origen)", "PO_03", source_col=3),
            Field("tipo_destino", "Tipo de punto de destino", "PO_01", source_col=4),
            Field("cantidad_destino", "Cantidad de puntos de destino", "PO_02", "number", 5),
            Field("restriccion_destino", "Restricción de altura (destino)", "PO_03", source_col=6),
        ),
    ),
    "ciudades": Section(
        slug="ciudades",
        title="Ciudades",
        sheet="CIUDADES",
        filename="CIUDADES.xlsx",
        source_start_row=2,
        description="Datos ambientales y eléctricos de las ciudades utilizadas por los proyectos.",
        fields=(
            Field("pais", "País", source_col=1),
            Field("ciudad", "Ciudad", source_col=2),
            Field("altitud_m", "Altitud (m.s.n.m)", kind="number", source_col=3),
            Field("presion_bar", "Presión ATM (bar)", kind="number", source_col=4),
            Field("temperatura_c", "Temperatura promedio (°C)", kind="number", source_col=5),
            Field("frecuencia_hz", "Frecuencia (Hz)", kind="number", source_col=6),
        ),
    ),
    "materiales": Section(
        slug="materiales",
        title="Tipos de material",
        sheet="TIPOS DE MATERIALES",
        filename="TIPOS_DE_MATERIALES.xlsx",
        source_start_row=3,
        description="Propiedades de producto requeridas para la selección y comparación de soluciones.",
        fields=(
            Field("codigo", "Código", source_col=1, kind="number"),
            Field("nombre", "Nombre común", "MT-02", source_col=2),
            Field("densidad", "Densidad aparente (kg/l)", "MT-03", "number", 3),
            Field("tamano_micras", "Tamaño medio (micras)", "MT-04", "number", 4),
            Field("temperatura_producto", "Temperatura del producto (°C)", "MT-16", "number", 5),
            Field("fluidez", "Fluidez", "MT-05", source_col=6),
            Field("abrasividad", "Abrasividad", "MT-06", source_col=7),
            Field("humedad", "Humedad (%)", "MT-07", "number", 8),
            Field("grasa", "Contenido grasa", "M-08", "number", 9),
            Field("regularidad", "Regularidad partícula", "MT-13", source_col=10),
            Field("absorcion_aire", "Absorción de aire", "MT-14", source_col=11),
            Field("angulo_reposo", "Ángulo de reposo con horizontal", "MT-15", "number", 12),
            Field("higroscopico", "Higroscópico", "MT-09", source_col=13),
            Field("explosividad", "Explosividad", "MT-10", source_col=14),
            Field("corrosivo", "Corrosivo", "MT-11", source_col=15),
            Field("riesgo_biologico", "Riesgo biológico", "MT-12", source_col=16),
        ),
    ),
    "cb-matriz": Section(
        slug="cb-matriz",
        title="CB_MATRIZ",
        sheet="CB_MATRIZ",
        filename="CB_MATRIZ.xlsx",
        source_start_row=5,
        description="Matriz de criterios base y opciones de entrada para la definición del sistema.",
        fields=(
            Field("subsistema", "Subsistema 1", "CB_OD_A", source_col=2),
            Field("voltaje", "Voltaje de potencia (V)", "CB_OD_A1", "number", 3),
            Field("material_contacto", "Material en contacto con producto", "CB_OD_A2", source_col=4),
            Field("material_estructura", "Material estructuras", "CB_OD_A3", source_col=5),
            Field("material_transportar", "Material a transportar", "CB_OD_A4", source_col=6),
            Field("flujo", "Flujo (kg/h) (oferta)", "CB_OD_A5", "number", 7),
            Field("distancia_horizontal", "Distancia horizontal de transporte (m)", "CB_OD_A6", "number", 8),
            Field("distancia_vertical", "Distancia vertical de transporte (m)", "CB_OD_A7", "number", 9),
            Field("curvas_90", "Cantidad de curvas de transporte x 90°", "CB_OD_A8", "number", 10),
            Field("distancia_soplado", "Distancia unidad de soplado / vacío (m)", "CB_OD_A9", "number", 11),
            Field("codos_soplado", "Codos tubería de soplado / vacío", "CB_OD_A10", "number", 12),
            Field("tipologia", "Preferencia tipológica de subsistema", "CB_OD_A11", source_col=13),
            Field("tipologia_descripcion", "Descripción tipología", source_col=14),
            Field("acople", "Preferencia de acoples tub. transporte", "CB_OD_A12", source_col=15),
            Field("tipo_flujo", "Tipo de flujo", "CB_OD_A13", source_col=16),
            Field("pesaje", "Pesaje en proceso OGA", "CB_OD_A14", source_col=17),
            Field("atex", "ATEX", "CB_OD_A15", source_col=18),
            Field("nec", "NEC", "CB_OD_A16", source_col=19),
            Field("exterior", "Exterior", "CB_OD_A17", source_col=20),
            Field("aire_comprimido", "Disponibilidad de aire comprimido", "CB_OD_A18", source_col=21),
        ),
    ),
    "codigo-equipos": Section(
        slug="codigo-equipos",
        title="Código de equipos",
        sheet="EQUIPOS",
        filename="CODIGO_DE_EQUIPOS.xlsx",
        source_start_row=2,
        description="Histórico de proyectos, subsistemas, tipos de equipo, referencias y cantidades para consulta rápida de códigos utilizados.",
        fields=(
            Field("proyecto", "Proyecto", source_col=1),
            Field("subsistema", "Nombre del subsistema", source_col=2),
            Field("tipo_equipo", "Tipo de Equipo", source_col=3),
            Field("referencia", "Referencia", source_col=4),
            Field("cantidad", "Cantidad", kind="number", source_col=5),
        ),
    ),
    "cs": Section(
        slug="cs",
        title="CS criterios de salida",
        sheet="CS-CRITERIOS DE SALIDA",
        filename="CS_CRITERIOS_DE_SALIDA.xlsx",
        source_start_row=5,
        description="Criterios de salida, opciones de equipos y búsqueda numérica por proximidad.",
        fields=(
            Field("codigo_cs", "Código CS", source_col=2),
            Field("criterio", "Criterio de salida", source_col=3),
            Field("nomenclatura", "Nomenclatura", source_col=5),
            Field("descripcion_transporte", "Descripción transporte", source_col=6),
            Field("diametro_tuberia", "Ø tubería de transporte", source_col=7),
            Field("tipo_acople", "Tipo de acople", source_col=8),
            Field("potencia_hp", "Potencia (hp)", kind="number", source_col=9),
            Field("caudal_cfm", "Caudal (CFM)", kind="number", source_col=10),
            Field("diferencial_psi", "Diferencial de presión (PSI)", kind="number", source_col=11),
            Field("tipo_bomba", "Tipo de bomba", source_col=12),
            Field("area_filtracion_m2", "Área de filtración (m²)", kind="number", source_col=13),
            Field("micraje_um", "Micraje área de filtración (µm)", kind="number", source_col=14),
            Field("consumo_aire_cfm", "Consumo de aire (CFM)", kind="number", source_col=15),
            Field("presion_alimentacion_psi", "Presión de alimentación (PSI)", kind="number", source_col=16),
            Field("equipo_slot", "Equipo", source_col=17),
            Field("tag", "TAG", source_col=18),
            Field("descripcion_equipo", "Descripción equipo", source_col=19),
            Field("referencia", "Referencia", source_col=20),
        ),
    ),
}

NUMERIC_CS_FIELDS = {
    "potencia_hp": "Potencia (hp)",
    "caudal_cfm": "Caudal (CFM)",
    "diferencial_psi": "Diferencial de presión (PSI)",
    "area_filtracion_m2": "Área de filtración (m²)",
    "micraje_um": "Micraje área de filtración (µm)",
    "consumo_aire_cfm": "Consumo de aire (CFM)",
    "presion_alimentacion_psi": "Presión de alimentación (PSI)",
}

CS_FIELD_UNITS = {
    "potencia_hp": "hp",
    "caudal_cfm": "CFM",
    "diferencial_psi": "PSI",
    "area_filtracion_m2": "m²",
    "micraje_um": "µm",
    "consumo_aire_cfm": "CFM",
    "presion_alimentacion_psi": "PSI",
}


def _path(section: Section) -> Path:
    return DATA_DIR / section.filename


def _atomic_save(section: Section, rows: list[dict[str, str]]) -> None:
    path = _path(section)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = section.sheet[:31]
    ws.append([f.label for f in section.fields] + ["_OGA_ID"])
    for row in rows:
        ws.append([_text(row.get(f.key)) for f in section.fields] + [_text(row.get("id")) or uuid.uuid4().hex])
    ws.freeze_panes = "A2"
    for idx, field in enumerate(section.fields, 1):
        ws.column_dimensions[ws.cell(1, idx).column_letter].width = min(max(len(field.label) + 3, 15), 42)
    fd, temp_name = tempfile.mkstemp(prefix=f".{path.stem}_", suffix=".xlsx", dir=path.parent)
    os.close(fd)
    temp = Path(temp_name)
    try:
        wb.save(temp)
        check = load_workbook(temp, read_only=True, data_only=False)
        check.close()
        os.replace(temp, path)
    finally:
        temp.unlink(missing_ok=True)


def _ensure(section: Section) -> None:
    if not _path(section).exists():
        _atomic_save(section, [])


def _load(section: Section) -> list[dict[str, str]]:
    with _LOCK:
        _ensure(section)
        wb = load_workbook(_path(section), read_only=True, data_only=True)
        try:
            ws = wb.active
            rows: list[dict[str, str]] = []
            for values in ws.iter_rows(min_row=2, values_only=True):
                if not any(v not in (None, "") for v in values):
                    continue
                row = {field.key: _text(values[i]) if i < len(values) else "" for i, field in enumerate(section.fields)}
                id_idx = len(section.fields)
                row["id"] = _text(values[id_idx]) if id_idx < len(values) else ""
                if not row["id"]:
                    row["id"] = uuid.uuid4().hex
                rows.append(row)
            return rows
        finally:
            wb.close()


def _is_allowed_upload(filename: str) -> bool:
    return Path(Path(filename).name).suffix.lower() in _ALLOWED_UPLOADS


def _parse_source(section: Section, file_storage) -> list[dict[str, str]]:
    if not _is_allowed_upload(file_storage.filename or ""):
        raise ValueError("Solo se permiten archivos Excel .xlsx o .xlsm.")
    file_storage.stream.seek(0)
    wb = load_workbook(file_storage.stream, read_only=True, data_only=True)
    try:
        exact_sheet = section.sheet in wb.sheetnames
        ws = wb[section.sheet] if exact_sheet else wb.active

        if exact_sheet:
            rows: list[dict[str, str]] = []
            for values in ws.iter_rows(min_row=section.source_start_row, values_only=True):
                values = tuple(values)
                row = {}
                for field in section.fields:
                    idx = (field.source_col - 1) if field.source_col else None
                    row[field.key] = _text(values[idx]) if idx is not None and idx < len(values) else ""
                if any(row.values()):
                    row["id"] = uuid.uuid4().hex
                    rows.append(row)
            return rows

        aliases = {_norm(f.label): f.key for f in section.fields}
        mapping: dict[int, str] = {}
        rows: list[dict[str, str]] = []
        header_found = False
        for row_num, values in enumerate(ws.iter_rows(values_only=True), start=1):
            values = tuple(values)
            if not header_found:
                if row_num > 12:
                    break
                candidate: dict[int, str] = {}
                for idx, value in enumerate(values):
                    key = aliases.get(_norm(value))
                    if key:
                        candidate[idx] = key
                if len(candidate) >= max(1, min(3, len(section.fields))):
                    mapping = candidate
                    header_found = True
                continue
            row = {f.key: "" for f in section.fields}
            for idx, key in mapping.items():
                row[key] = _text(values[idx]) if idx < len(values) else ""
            if any(row.values()):
                row["id"] = uuid.uuid4().hex
                rows.append(row)
        if not header_found:
            raise ValueError(f"No encontré la hoja '{section.sheet}' ni encabezados compatibles.")
        return rows
    finally:
        wb.close()


def _row_key(section: Section, row: dict[str, str]) -> tuple[str, ...]:
    if section.slug == "ciudades":
        return (_norm(row.get("pais")), _norm(row.get("ciudad")))
    if section.slug == "materiales":
        return (_norm(row.get("codigo")) or _norm(row.get("nombre")),)
    if section.slug == "puntos-origen":
        return (_norm(row.get("tipo_origen")), _norm(row.get("tipo_destino")))
    if section.slug == "cb-matriz":
        return (_norm(row.get("subsistema")), _norm(row.get("tipologia")), _norm(row.get("voltaje")))
    if section.slug == "codigo-equipos":
        return (
            _norm(row.get("proyecto")),
            _norm(row.get("subsistema")),
            _norm(row.get("tipo_equipo")),
            _norm(row.get("referencia")),
            _norm(row.get("cantidad")),
        )
    if section.slug == "cs":
        return (_norm(row.get("codigo_cs")), _norm(row.get("criterio")), _norm(row.get("tag")), _norm(row.get("referencia")))
    return tuple(_norm(row.get(f.key)) for f in section.fields[:2])


def _merge(section: Section, existing: list[dict[str, str]], incoming: list[dict[str, str]]) -> tuple[list[dict[str, str]], int, int]:
    index = {_row_key(section, r): r for r in existing if any(_row_key(section, r))}
    added = updated = 0
    for src in incoming:
        key = _row_key(section, src)
        target = index.get(key) if any(key) else None
        if target is None:
            existing.append(src)
            if any(key):
                index[key] = src
            added += 1
        else:
            changed = False
            for field in section.fields:
                value = _text(src.get(field.key))
                if value and value != target.get(field.key, ""):
                    target[field.key] = value
                    changed = True
            if changed:
                updated += 1
    return existing, added, updated


def _simple_numeric(value: Any) -> float | None:
    raw = _text(value)
    if not raw:
        return None

    # Las celdas del maestro pueden contener instrucciones (por ejemplo,
    # "EJEMPLO 8,4 hp") que NO son datos del catálogo. Se ignoran para
    # evitar resultados falsos en el search de proximidad.
    normalized = _norm(raw)
    if any(marker in normalized for marker in (
        "NO APLICA",
        "EJEMPLO",
        "NUMERO COMPUESTO",
        "MAXIMO",
    )):
        return None

    # Acepta valores reales con o sin unidad: 8.4, 8,4, 380 CFM, 48 m²,
    # 5 µm, 5 um, 5 micras, etc. Debe ser toda la celda, no una frase.
    txt = re.sub(r"\s+", "", raw)
    m = re.fullmatch(
        r"(-?\d+(?:[\.,]\d+)?)(?:hp|cfm|psi|m2|m²|µm|μm|um|micra|micras)?",
        txt,
        flags=re.IGNORECASE,
    )
    if not m:
        return None
    try:
        number = float(m.group(1).replace(",", "."))
        return number if math.isfinite(number) else None
    except ValueError:
        return None


def _proximity_counts(rows: list[dict[str, str]]) -> dict[str, int]:
    return {
        field: sum(1 for row in rows if _simple_numeric(row.get(field)) is not None)
        for field in NUMERIC_CS_FIELDS
    }


def _proximity(rows: list[dict[str, str]], field: str, target: float) -> dict[str, Any]:
    """Return the numeric window immediately around the value typed by the user.

    The center of the window is always the searched value. If that value exists in
    the catalogue, ``exact`` is true. The neighbours are always *distinct* values:
    the greatest catalogue value strictly below the target and the smallest value
    strictly above it. Example: catalogue [10, 20, 50], target 20 -> 10 | 20 | 50.
    """
    candidates: list[tuple[float, dict[str, str]]] = []
    for row in rows:
        number = _simple_numeric(row.get(field))
        if number is not None:
            candidates.append((number, row))

    if not candidates:
        return {
            "target": target,
            "count": 0,
            "distinct_count": 0,
            "lower": None,
            "upper": None,
            "exact": False,
            "exact_match": None,
        }

    candidates.sort(key=lambda item: item[0])
    tolerance = max(1e-9, abs(target) * 1e-9)

    lower_items = [item for item in candidates if item[0] < target - tolerance]
    upper_items = [item for item in candidates if item[0] > target + tolerance]
    exact_items = [item for item in candidates if math.isclose(item[0], target, rel_tol=1e-9, abs_tol=tolerance)]

    lo = max(lower_items, key=lambda item: item[0]) if lower_items else None
    hi = min(upper_items, key=lambda item: item[0]) if upper_items else None
    exact_item = exact_items[0] if exact_items else None

    def pack(item):
        if not item:
            return None
        number, row = item
        return {
            "value": number,
            "distance": abs(number - target),
            "row": row,
        }

    distinct_values = {number for number, _ in candidates}
    return {
        "target": target,
        "count": len(candidates),
        "distinct_count": len(distinct_values),
        "lower": pack(lo),
        "upper": pack(hi),
        "exact": bool(exact_item),
        "exact_match": pack(exact_item),
    }


def register_industria_routes(bp) -> None:
    @bp.route("/industria", methods=["GET"], endpoint="industria_index")
    def industria_index():
        counts = {slug: len(_load(section)) for slug, section in SECTIONS.items()}
        return render_template(
            "industria/index.html",
            sections=SECTIONS,
            counts=counts,
            industria_page=True,
            meta={},
        )

    @bp.route("/industria/<slug>", methods=["GET"], endpoint="industria_section")
    def industria_section(slug: str):
        section = SECTIONS.get(slug)
        if not section:
            return redirect(url_for("industria.industria_index"))

        all_rows = _load(section)
        rows = list(all_rows)
        q = request.args.get("q", "").strip()

        proximity_counts = _proximity_counts(all_rows) if slug == "cs" else {}
        available_numeric_fields = (
            {key: label for key, label in NUMERIC_CS_FIELDS.items() if proximity_counts.get(key, 0) > 0}
            if slug == "cs"
            else {}
        )

        proximity_field = ""
        proximity_value = ""
        proximity_active = False
        proximity_exact = False
        proximity_exact_ids: set[str] = set()

        if slug == "cs":
            proximity_field = request.args.get("prox_criterio", "").strip()
            proximity_value = request.args.get("prox_valor", "").strip()
            if proximity_field and proximity_value:
                if proximity_field in available_numeric_fields:
                    try:
                        target = float(proximity_value.replace(",", "."))
                        if not math.isfinite(target):
                            raise ValueError
                        result = _proximity(all_rows, proximity_field, target)
                        keep_values: list[float] = []
                        if result.get("lower"):
                            keep_values.append(float(result["lower"]["value"]))
                        if result.get("exact_match"):
                            keep_values.append(float(result["exact_match"]["value"]))
                        if result.get("upper"):
                            keep_values.append(float(result["upper"]["value"]))

                        tolerance = max(1e-9, abs(target) * 1e-9)
                        filtered_rows: list[dict[str, str]] = []
                        for row in all_rows:
                            number = _simple_numeric(row.get(proximity_field))
                            if number is None:
                                continue
                            if any(math.isclose(number, value, rel_tol=1e-9, abs_tol=tolerance) for value in keep_values):
                                filtered_rows.append(row)
                            if math.isclose(number, target, rel_tol=1e-9, abs_tol=tolerance):
                                proximity_exact_ids.add(_text(row.get("id")))

                        rows = sorted(
                            filtered_rows,
                            key=lambda row: (_simple_numeric(row.get(proximity_field)) is None, _simple_numeric(row.get(proximity_field)) or 0.0),
                        )
                        proximity_active = True
                        proximity_exact = bool(proximity_exact_ids)
                    except ValueError:
                        flash("Ingresa un valor numerico valido para la busqueda de proximidad.", "error")
                else:
                    flash("El criterio seleccionado no tiene valores numericos disponibles.", "error")

        if q:
            nq = _norm(q)
            rows = [r for r in rows if nq in _norm(" ".join(_text(r.get(f.key)) for f in section.fields))]

        # Código de equipos contiene miles de registros. Se pagina para mantener
        # la tabla ágil sin cambiar el comportamiento de los demás catálogos.
        total_rows = len(rows)
        pagination_enabled = slug == "codigo-equipos"
        page_size = 150 if pagination_enabled else max(total_rows, 1)
        try:
            page = max(1, int(request.args.get("page", "1")))
        except ValueError:
            page = 1
        total_pages = max(1, math.ceil(total_rows / page_size)) if pagination_enabled else 1
        if page > total_pages:
            page = total_pages
        if pagination_enabled:
            start = (page - 1) * page_size
            rows = rows[start:start + page_size]

        template = "industria/cs.html" if slug == "cs" else "industria/section.html"
        return render_template(
            template,
            section=section,
            rows=rows,
            q=q,
            numeric_fields=available_numeric_fields,
            proximity_counts=proximity_counts,
            proximity_field=proximity_field,
            proximity_value=proximity_value,
            proximity_active=proximity_active,
            proximity_exact=proximity_exact,
            proximity_exact_ids=proximity_exact_ids,
            total_rows=total_rows,
            pagination_enabled=pagination_enabled,
            page=page,
            page_size=page_size,
            total_pages=total_pages,
            industria_page=True,
            meta={},
        )

    @bp.route("/industria/<slug>/guardar", methods=["POST"], endpoint="industria_save")
    def industria_save(slug: str):
        section = SECTIONS.get(slug)
        if not section:
            return redirect(url_for("industria.industria_index"))
        new_row = {f.key: _text(request.form.get(f.key)) for f in section.fields}
        if not any(new_row.values()):
            flash("Completa al menos un campo antes de guardar.", "error")
            return redirect(url_for("industria.industria_section", slug=slug))
        new_row["id"] = uuid.uuid4().hex
        with _LOCK:
            rows = _load(section)
            rows.append(new_row)
            _atomic_save(section, rows)
        flash("Registro agregado correctamente.", "success")
        return redirect(url_for("industria.industria_section", slug=slug))

    @bp.route("/industria/<slug>/<row_id>/editar", methods=["POST"], endpoint="industria_edit")
    def industria_edit(slug: str, row_id: str):
        section = SECTIONS.get(slug)
        if not section:
            return redirect(url_for("industria.industria_index"))
        with _LOCK:
            rows = _load(section)
            target = next((r for r in rows if r.get("id") == row_id), None)
            if target is None:
                flash("No se encontró el registro.", "error")
            else:
                for field in section.fields:
                    target[field.key] = _text(request.form.get(field.key))
                _atomic_save(section, rows)
                flash("Registro actualizado.", "success")
        return redirect(url_for("industria.industria_section", slug=slug))

    @bp.route("/industria/<slug>/<row_id>/eliminar", methods=["POST"], endpoint="industria_delete")
    def industria_delete(slug: str, row_id: str):
        section = SECTIONS.get(slug)
        if not section:
            return redirect(url_for("industria.industria_index"))
        with _LOCK:
            rows = _load(section)
            new_rows = [r for r in rows if r.get("id") != row_id]
            if len(rows) == len(new_rows):
                flash("No se encontró el registro.", "error")
            else:
                _atomic_save(section, new_rows)
                flash("Registro eliminado.", "success")
        return redirect(url_for("industria.industria_section", slug=slug))

    @bp.route("/industria/<slug>/subida-masiva", methods=["POST"], endpoint="industria_bulk")
    def industria_bulk(slug: str):
        section = SECTIONS.get(slug)
        if not section:
            return redirect(url_for("industria.industria_index"))
        archivo = request.files.get("archivo")
        mode = request.form.get("mode", "replace")
        if not archivo or not archivo.filename:
            flash("Selecciona un archivo Excel.", "error")
            return redirect(url_for("industria.industria_section", slug=slug))
        try:
            incoming = _parse_source(section, archivo)
            if not incoming:
                raise ValueError(f"La hoja {section.sheet} no contiene registros para importar.")
            with _LOCK:
                if mode == "merge":
                    rows, added, updated = _merge(section, _load(section), incoming)
                    _atomic_save(section, rows)
                    flash(f"Carga completada: {added} registros agregados y {updated} actualizados.", "success")
                else:
                    _atomic_save(section, incoming)
                    flash(f"Catálogo reemplazado correctamente: {len(incoming)} registros cargados.", "success")
        except Exception as exc:
            flash(f"No fue posible importar {section.title}: {exc}", "error")
        return redirect(url_for("industria.industria_section", slug=slug))

    @bp.route("/industria/cs/proximidad", methods=["GET"], endpoint="industria_cs_proximity")
    def industria_cs_proximity():
        field = request.args.get("criterio", "")
        raw_value = request.args.get("valor", "")
        if field not in NUMERIC_CS_FIELDS:
            return jsonify({"ok": False, "error": "Criterio numérico no válido."}), 400
        try:
            target = float(raw_value.replace(",", "."))
            if not math.isfinite(target):
                raise ValueError
        except (ValueError, AttributeError):
            return jsonify({"ok": False, "error": "Ingresa un valor numérico válido."}), 400
        result = _proximity(_load(SECTIONS["cs"]), field, target)
        result.update({
            "ok": True,
            "criterion": field,
            "criterion_label": NUMERIC_CS_FIELDS[field],
            "unit": CS_FIELD_UNITS.get(field, ""),
        })
        return jsonify(result)


register_industria_routes(bp)
