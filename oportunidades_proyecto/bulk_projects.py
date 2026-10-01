"""Carga estricta de proyectos con la estructura pública SISTEMAS A–AB."""
from __future__ import annotations

from datetime import datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
import json
from pathlib import Path
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import db
from .import_excel import _cell_text, _header_key, _parse_project_matrix_workbook


def fields():
    path = Path(__file__).resolve().parent.parent / 'documentacion/catalogos_sistemas_20260930.json'
    return json.loads(path.read_text(encoding='utf-8'))['fields']


class BulkValidationError(ValueError):
    def __init__(self, errors):
        self.errors = errors
        super().__init__(f'El Excel contiene {len(errors)} errores. Corríjalos y vuelva a analizarlo.')


def split_project(value):
    match = re.fullmatch(r'(\d+)\s*([A-Za-z]{1,2}\d*)?', _cell_text(value))
    if not match:
        return '', ''
    return match[1], 'Principal' + (' ' + match[2].upper() if match[2] else '')


def existing_projects(conn=None):
    if conn is None:
        with db.connection() as opened:
            return existing_projects(opened)
    rows = conn.execute('SELECT proyecto FROM opportunities WHERE eliminado_en IS NULL').fetchall()
    result = {}
    for row in rows:
        key = db.normalize_search(row['proyecto'])
        result[key] = result.get(key, 0) + 1
    return result


def validate_targets(payload, conn=None):
    """Se ejecuta nuevamente al confirmar; la base puede haber cambiado."""
    if conn is None:
        with db.connection() as opened:
            return validate_targets(payload, opened)
    mode = payload.get('import_mode')
    errors = []
    if mode not in {'add', 'update'} or payload.get('validation_version') != 1:
        raise ValueError('Analice nuevamente el Excel desde Agregar / actualizar desde Excel.')
    current = existing_projects(conn)
    for item in payload.get('opportunities') or []:
        project = item['data']['proyecto']
        count = current.get(db.normalize_search(project), 0)
        reason = ''
        if count > 1:
            reason = 'Hay varios proyectos con este código en la base; debe resolver esa duplicidad.'
        elif mode == 'add' and count:
            reason = 'El proyecto ya existe. Utilice Actualizar proyectos para modificarlo.'
        elif mode == 'update' and not count:
            reason = 'El proyecto no existe. Utilice Agregar proyectos para crearlo.'
        if reason:
            for row in item['source']['rows']:
                errors.append(dict(row=row, column='A', field='Proyecto', value=project, reason=reason))
    for item in payload.get('opportunities') or []:
        rows = conn.execute("SELECT s.nombre FROM subsystems s JOIN opportunities o ON o.id=s.opportunity_id WHERE o.proyecto=? AND o.eliminado_en IS NULL AND s.eliminado_en IS NULL", (item['data']['proyecto'],)).fetchall()
        suffixes = [db.subsystem_suffix(row['nombre']) for row in rows]
        for sub in item['subsystems']:
            if suffixes.count(db.subsystem_suffix(sub['data']['nombre'])) > 1:
                errors.append(dict(row=sub['source_row'], column='A', field='Proyecto', value=item['data']['proyecto'], reason='Hay varios subsistemas con este sufijo en la base. Resuelva la duplicidad antes de actualizar.'))
    return errors


def preview_changes(payload, comparison):
    changes = []
    if payload.get('import_mode') != 'update':
        return changes
    ids = {row['proyecto']: row['existing_id'] for row in comparison['rows']}
    labels = {field['key']: field['label'] for field in fields()}
    labels['fecha_inicio'] = 'Inicio'

    def compare(project, subsystem, before, after, keys):
        for key in keys:
            incoming = str(after.get(key) or '').strip()
            current = str(before[key] or '').strip()
            if incoming and incoming != current:
                changes.append(dict(project=project, subsystem=subsystem,
                                    field=labels.get(key, key), before=current, after=incoming))

    with db.connection() as conn:
        for item in payload['opportunities']:
            project = item['data']['proyecto']
            oid = ids.get(project)
            if oid is None:
                continue
            current = conn.execute('SELECT * FROM opportunities WHERE id=?', (oid,)).fetchone()
            compare(project, '', current, item['data'], ('cliente', 'industria', 'fecha_inicio'))
            existing = {db.subsystem_suffix(row['nombre']): row for row in conn.execute(
                'SELECT * FROM subsystems WHERE opportunity_id=? AND eliminado_en IS NULL', (oid,))}
            for sub in item['subsystems']:
                suffix = db.subsystem_suffix(sub['data']['nombre'])
                current_sub = existing.get(suffix)
                if current_sub is None:
                    continue
                compare(project, suffix or 'Principal', current_sub, sub['data'],
                        (key for key in labels if key in db.SUBSYSTEM_FIELDS and key != 'nombre'))
                for table, staged_key, field in [('entry_points', 'entry_points', 'punto_ingreso'),
                                                  ('exit_points', 'exit_points', 'punto_destino')]:
                    staged = sub.get(staged_key) or []
                    if not staged:
                        continue
                    row = conn.execute(f'SELECT tipo FROM {table} WHERE subsystem_id=? ORDER BY id LIMIT 1', (current_sub['id'],)).fetchone()
                    compare(project, suffix or 'Principal', {field: row['tipo'] if row else ''},
                            {field: staged[0]['tipo']}, (field,))
    return changes


def parse_projects(stream, filename, mode, catalogs):
    schema = fields()
    errors = []

    def error(row, col, value, reason):
        errors.append(dict(row=row, column=get_column_letter(col),
                           field=schema[col-1]['label'] if col <= 28 else 'Columna adicional',
                           value=_cell_text(value), reason=reason))

    if mode not in {'add', 'update'}:
        raise ValueError('Seleccione Agregar proyectos o Actualizar proyectos.')
    stream.seek(0)
    try:
        workbook = load_workbook(stream, read_only=True, data_only=False)
    except Exception as exc:
        raise ValueError('No fue posible abrir el archivo .xlsx.') from exc
    normalized = Workbook()
    output = normalized.active
    output.title = 'SISTEMAS'
    try:
        if 'SISTEMAS' not in workbook.sheetnames:
            raise BulkValidationError([dict(row=1, column='A', field='Hoja', value='', reason='Falta la hoja SISTEMAS.')])
        sheet = workbook['SISTEMAS']
        # Admite la plantilla descargada y el histórico con descriptores en fila 1.
        header_row = max((1, 2), key=lambda row: sum(
            _header_key(sheet.cell(row, col).value) == _header_key(field['label'])
            for col, field in enumerate(schema, 1)
        ))
        for col, field in enumerate(schema, 1):
            value = sheet.cell(header_row, col).value
            if _header_key(value) != _header_key(field['label']):
                error(header_row, col, value, f"Se esperaba '{field['label']}' en esta posición. No cambie los nombres ni el orden A–AB.")
        for col in range(29, sheet.max_column + 1):
            if sheet.cell(header_row, col).value is not None:
                error(header_row, col, sheet.cell(header_row, col).value, 'La plantilla debe tener exactamente 28 columnas, A–AB.')
        if errors:
            raise BulkValidationError(errors)
        output.append([field['label'] for field in schema])
        seen = {}
        general = {}
        row_count = 0
        for row_number, cells in enumerate(sheet.iter_rows(min_row=header_row + 1), header_row + 1):
            if not any(cell.value is not None and _cell_text(cell.value) for cell in cells):
                output.append([])  # conserva los números de fila del archivo
                continue
            row_count += 1
            values = [_cell_text(cell.value) for cell in cells[:28]]
            values += [''] * (28-len(values))
            project, subsystem = split_project(values[0])
            if not project:
                error(row_number, 1, values[0], 'Falta el código o su formato es inválido. Ejemplos: 3876, 3876A, 3876B, 3876A1.')
            elif (project, subsystem) in seen:
                error(row_number, 1, values[0], f'Proyecto y subsistema repetidos; aparecen también en la fila {seen[(project, subsystem)]}.')
            else:
                seen[(project, subsystem)] = row_number
            if mode == 'add' and not values[1]:
                error(row_number, 2, '', 'El cliente es obligatorio al agregar un proyecto.')
            for col, field in enumerate(schema, 1):
                value = values[col-1]
                cell = cells[col-1] if col <= len(cells) else None
                if cell is not None and cell.data_type in {'f', 'e'}:
                    error(row_number, col, value, 'Use un valor directo; no se admiten fórmulas ni errores de Excel.')
                    continue
                if not value:
                    continue
                if field['kind'] == 'number':
                    try:
                        number = Decimal(value.replace(',', '.'))
                        if not number.is_finite() or number < 0:
                            raise InvalidOperation
                        if field['key'] in {'curvas_90', 'curvas_unidad_soplado'} and number != number.to_integral_value():
                            raise InvalidOperation
                        values[col-1] = format(number, 'f')
                    except InvalidOperation:
                        reason = 'Debe ser un entero mayor o igual a cero.' if field['key'] in {'curvas_90', 'curvas_unidad_soplado'} else 'Debe ser un número mayor o igual a cero, sin unidades; use coma o punto decimal.'
                        error(row_number, col, value, reason)
                elif field['kind'] == 'date':
                    try:
                        parsed = None
                        for fmt in ('%Y-%m-%d', '%d/%m/%Y'):
                            try:
                                parsed = datetime.strptime(value, fmt).date()
                                break
                            except ValueError:
                                pass
                        if parsed is None:
                            raise ValueError
                        values[col-1] = parsed.isoformat()
                    except ValueError:
                        error(row_number, col, value, 'Fecha inválida. Use una fecha de Excel, DD/MM/AAAA o AAAA-MM-DD.')
                if field['key'] in catalogs:
                    options = catalogs[field['key']]
                    if field['kind'] == 'number':
                        try:
                            valid = any(Decimal(str(v).replace(',', '.')) == Decimal(values[col-1]) for v in options)
                        except InvalidOperation:
                            valid = False
                    else:
                        matches = [v for v in options if db.normalize_search(v) == db.normalize_search(value)]
                        valid = bool(matches)
                        if valid:
                            values[col-1] = matches[0]
                    if not valid:
                        location = {'clientes': 'Proyectos > Clientes', 'materiales': 'Datos > Materiales',
                                    'puntos-origen': 'Datos > Puntos de origen y destino',
                                    'cb-matriz': 'Datos > Criterios base', 'cs': 'Datos > Criterios de salida',
                                    'sistemas': 'Datos > Otros datos de SISTEMAS'}[field['section']]
                        error(row_number, col, value, f'El valor no está registrado en {location}. Agréguelo allí o use un valor del catálogo.')
            # Cliente e industria son datos del proyecto y deben coincidir entre A/B.
            for index in (1, 12):
                value = values[index]
                if not project or not value:
                    continue
                key = (project, index)
                previous = general.get(key)
                if previous and db.normalize_search(value) != db.normalize_search(previous[0]):
                    error(row_number, index+1, value, f'No coincide con el valor del mismo proyecto en la fila {previous[1]}.')
                else:
                    general[key] = (value, row_number)
            for col, cell in enumerate(cells[28:], 29):
                if _cell_text(cell.value):
                    error(row_number, col, cell.value, 'Hay datos fuera de las columnas A–AB.')
            output.append(values)
        if not row_count:
            error(header_row+1, 1, '', 'No hay filas de proyectos para cargar.')
        if errors:
            raise BulkValidationError(errors)
        payload = _parse_project_matrix_workbook(normalized, filename)
        # La hoja normalizada comienza en fila 1; reconstruye las filas reales.
        offset = header_row-1
        for item in payload['opportunities']:
            item['source']['rows'] = [row+offset for row in item['source']['rows']]
            for sub in item['subsystems']:
                sub['source_row'] += offset
        payload.update(import_mode=mode, validation_version=1)
        errors = validate_targets(payload)
        if errors:
            raise BulkValidationError(errors)
        return payload
    finally:
        workbook.close()
        normalized.close()


def template_bytes():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'SISTEMAS'
    schema = fields()
    sheet.append([field['label'] for field in schema])
    sheet.freeze_panes = 'C2'
    sheet.auto_filter.ref = 'A1:AB1001'
    sheet.row_dimensions[1].height = 72
    for col, field in enumerate(schema, 1):
        cell = sheet.cell(1, col)
        cell.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='143550')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        sheet.column_dimensions[get_column_letter(col)].width = 24 if col != 2 else 40
    instructions = workbook.create_sheet('INSTRUCCIONES')
    for line in [
        'Carga masiva de proyectos OGA',
        'Complete SISTEMAS desde la fila 2. Mantenga las 28 columnas A–AB y sus encabezados.',
        'Una fila por proyecto y subsistema. Ejemplos: 3876A, 3876B, 3876A1.',
        'Agregar: proyecto y cliente obligatorios. El código no debe existir en la base.',
        'Actualizar: el proyecto debe existir. Las celdas vacías conservan el dato actual.',
        'Cliente e industria deben coincidir entre los subsistemas del mismo proyecto.',
        'Los campos de catálogo deben existir en Datos. Los números pueden ser nuevos.',
        'Inicio: fecha de Excel, DD/MM/AAAA o AAAA-MM-DD. Números sin unidades y mayores o iguales a cero.',
        'ATEX admite SI, NO y clasificaciones históricas registradas en Datos.',
        'Revise la previsualización. Si hay errores, ninguna fila se guarda.',
        'Esta carga procesa SISTEMAS. No importa equipos ni otras hojas.',
    ]:
        instructions.append([line])
        instructions.cell(instructions.max_row, 1).alignment = Alignment(wrap_text=True, vertical='center')
        instructions.row_dimensions[instructions.max_row].height = 36
    instructions.column_dimensions['A'].width = 110
    instructions['A1'].font = Font(bold=True, size=16, color='143550')
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    stream.seek(0)
    return stream
