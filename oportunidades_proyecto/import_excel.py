from __future__ import annotations

from collections import OrderedDict
from datetime import date, datetime, timezone
from pathlib import Path
import json
import re
import secrets
import unicodedata
from typing import Any

from openpyxl import load_workbook

from .db import DATA_DIR, normalize_search

STAGE_DIR = DATA_DIR / "importaciones_temporales"
MAX_STAGE_AGE_HOURS = 24

REQUIRED_SHEETS = (
    "SISTEMAS",
    "PUNTOS DE ENTRADA",
    "PUNTOS DE SALIDA",
    "Equipos",
)

SYSTEM_FIELD_HEADERS = {
    "radicado": ("Número de Radicado", "Numero de Radicado"),
    "proyecto": ("Proyecto",),
    "cliente": ("Cliente",),
    "planta": ("Planta",),
    "contacto": ("Contacto",),
    "inicio": ("Inicio",),
    "estado_origen": ("Estado",),
    "tipo_oportunidad": ("Tipo de Oportunidad",),
    "descripcion": ("Descripción", "Descripcion"),
    "nombre_subsistema": ("Nombre del Proceso",),
    "descripcion_proceso": ("Descripción del Proceso", "Descripcion del Proceso"),
    "voltaje_potencia": ("Voltaje de potencia", "Voltaje de potencia (V)"),
    "material_contacto": ("Material de contacto con el producto", "Material en contacto con el producto"),
    "material_estructural": ("Material estructuras", "Material estructural"),
    "material_transportado": ("Material a Transportar", "Material transportado"),
    "flujo_kg_h": ("Flujo en KG/Hora", "Flujo (kg/h)"),
    "distancia_horizontal_m": ("Distancia Horizontal de Transporte (m)", "Distancia horizontal (m)"),
    "distancia_vertical_m": ("Distancia Vertical de Transporte (m)", "Distancia vertical (m)"),
    "curvas_90": ("Cantidad de curvas de tubería x 90 grados", "Curvas de tubería x 90°"),
    "distancia_unidad_soplado_m": ("Distancia Unidad de Soplado (m)", "Distancia unidad de soplado/vacío (m)"),
    "curvas_unidad_soplado": ("Curvas Tubería Unidad de Soplado / Vacío (Ud)", "Curvas unidad de soplado/vacío"),
    "preferencia_tipologia": ("Preferencia de tipología", "Preferencia de tipologia"),
    "preferencia_acoples": ("Preferencia de acoples",),
    "tipo_flujo": ("Tipo de flujo",),
    "pesaje_oga": ("Pesaje en proceso OGA", "Pesaje OGA"),
    "atex": ("ATEX", "Clasificación ATEX", "Clasificacion ATEX"),
    "nec": ("NEC", "Clasificación NEC", "Clasificacion NEC"),
    "ubicacion": ("Exterior", "Ubicación (interior/exterior)", "Ubicacion (interior/exterior)"),
    "aire_comprimido": ("Disponibilidad de aire comprimido",),
    "tipo_transporte": ("Tipo de Transporte",),
    "diametro_tuberia": ("Diámetro Tubería Transporte", "Diametro Tuberia Transporte"),
    "tipo_acople": ("Tipo de Acople",),
    "potencia_hp": ("Potencia (hp)",),
    "caudal_cfm": ("Caudal (CFM)",),
    "diferencial_presion_psi": ("Diferencial de Presión (PSI)", "Diferencial de Presion (PSI)"),
    "tipo_bomba": ("Tipo de Bomba",),
    "area_filtracion_m2": ("Área de Filtración (m²)", "Area de Filtracion (m2)"),
    "micraje_filtracion": ("Micraje de área de filtración", "Micraje de area de filtracion"),
    "consumo_aire_cfm": ("Consumo de Aire (cfm)",),
    "presion_alimentacion_psi": ("Presión de Alimentación (PSI)", "Presion de Alimentacion (PSI)"),
    "es_multiequipos": ("Es multiequipos",),
}

SUBSYSTEM_FIELDS = (
    "nombre", "nombre_proceso", "descripcion_proceso", "voltaje_potencia",
    "material_contacto", "material_estructural", "material_transportado", "flujo_kg_h",
    "distancia_horizontal_m", "distancia_vertical_m", "curvas_90",
    "distancia_unidad_soplado_m", "curvas_unidad_soplado", "preferencia_tipologia",
    "preferencia_acoples", "tipo_flujo", "pesaje_oga", "atex", "nec", "ubicacion",
    "aire_comprimido", "tipo_transporte", "diametro_tuberia", "tipo_acople",
    "potencia_hp", "caudal_cfm", "diferencial_presion_psi", "tipo_bomba",
    "area_filtracion_m2", "micraje_filtracion", "consumo_aire_cfm",
    "presion_alimentacion_psi", "es_multiequipos", "observaciones",
)


def _header_key(value: Any) -> str:
    text = str(value or "").strip().casefold()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.replace("²", "2")
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return " ".join(text.split())


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "SI" if value else "NO"
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return format(value, ".15g")
    return str(value).strip()


def _raw_relation_text(value: Any) -> str:
    """Texto de relación que conserva espacios finales del Excel.

    Algunos radicados históricos reutilizan el nombre visual "Principal" y se
    distinguen en el archivo fuente por un espacio final. Para enlazar entradas,
    salidas y equipos primero conservamos ese valor exacto; solo si no existe una
    coincidencia exacta usamos una coincidencia normalizada y no ambigua.
    """
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).replace("\xa0", " ").casefold()


def _identifier(value: Any) -> str:
    return _cell_text(value)


def _find_sheet(workbook, expected_name: str):
    wanted = _header_key(expected_name)
    for sheet in workbook.worksheets:
        if _header_key(sheet.title) == wanted:
            return sheet
    raise ValueError(f"No se encontró la hoja obligatoria '{expected_name}'.")


def _find_header_row(sheet, required_headers: tuple[str, ...], max_scan: int = 8) -> tuple[int, dict[str, int]]:
    required = {_header_key(item) for item in required_headers}
    for row_number in range(1, min(sheet.max_row, max_scan) + 1):
        current: dict[str, int] = {}
        for column, cell in enumerate(sheet[row_number], start=1):
            key = _header_key(cell.value)
            if key and key not in current:
                current[key] = column
        if required.issubset(set(current)):
            return row_number, current
    raise ValueError(f"No se reconocieron los encabezados de la hoja '{sheet.title}'.")


def _column(headers: dict[str, int], aliases: tuple[str, ...], *, required: bool = False) -> int | None:
    for alias in aliases:
        index = headers.get(_header_key(alias))
        if index is not None:
            return index
    if required:
        raise ValueError(f"Falta la columna obligatoria '{aliases[0]}'.")
    return None


def _read_cell(row: tuple[Any, ...], column: int | None) -> Any:
    if column is None or column < 1 or column > len(row):
        return None
    return row[column - 1]


def _map_status(raw: Any) -> tuple[str, bool]:
    original = _cell_text(raw)
    value = normalize_search(original)
    if not value:
        return "Nueva", False
    if "cerrada" in value and ("no satisfactoria" in value or "insatisfactoria" in value or "perdida" in value):
        return "Perdida", True
    if "cerrada" in value and "satisfactoria" in value:
        return "Ganada", True
    if "ganad" in value or "adjudic" in value:
        return "Ganada", True
    if "perdid" in value or "rechaz" in value:
        return "Perdida", True
    if "oferta" in value and ("envi" in value or "present" in value):
        return "Oferta enviada", True
    if "cotiz" in value:
        return "Cotizando", True
    if "anal" in value or "estudio" in value:
        return "En análisis", True
    if "abiert" in value or "nueva" in value:
        return "Nueva", True
    return "Nueva", False


def _first_nonempty(items: list[str]) -> str:
    for item in items:
        if str(item or "").strip():
            return str(item).strip()
    return ""



PROJECT_MATRIX_HEADERS = {
    "proyecto": ("Proyecto",),
    "cliente": ("Cliente",),
    "inicio": ("Inicio",),
    "material_transportado": ("Material a Transportar",),
    "punto_ingreso": ("PUNTO DE INGRESO", "Punto de ingreso"),
    "punto_destino": ("PUNTO DESTINO", "Punto destino"),
    "flujo_kg_h": ("Flujo en KG/Hora",),
    "distancia_horizontal_m": ("Distancia Horizontal de Transporte (m)",),
    "distancia_vertical_m": ("Distancia Vertical de Transporte (m)",),
    "curvas_90": ("Cantidad de curvas de tubería x 90 grados",),
    "distancia_unidad_soplado_m": ("Distancia Unidad de Soplado (m)",),
    "curvas_unidad_soplado": ("Curvas Tubería Unidad de Soplado / Vacío (Ud)",),
    "industria": ("TIPO DE INDUSTRIA", "Tipo de industria"),
    "tipo_flujo": ("Tipo de flujo",),
    "atex": ("ATEX",),
    "material_contacto": ("Material de contacto con el producto",),
    "voltaje_potencia": ("Voltaje de potencia", "Voltaje de potencia (V)"),
    "tipo_transporte": ("Tipo de Transporte",),
    "diametro_tuberia": ("Diámetro Tubería Transporte", "Diametro Tuberia Transporte"),
    "tipo_acople": ("Tipo de Acople",),
    "potencia_hp": ("Potencia (hp)",),
    "caudal_cfm": ("Caudal (CFM)",),
    "diferencial_presion_psi": ("Diferencial de Presión (PSI)", "Diferencial de Presion (PSI)"),
    "tipo_bomba": ("Tipo de Bomba",),
    "area_filtracion_m2": ("Área de Filtración (m²)", "Area de Filtracion (m2)"),
    "micraje_filtracion": ("Micraje de área de filtración", "Micraje de area de filtracion"),
    "consumo_aire_cfm": ("Consumo de Aire (cfm)",),
    "presion_alimentacion_psi": ("Presión de Alimentación (PSI)", "Presion de Alimentacion (PSI)"),
}


def _split_project_subsystem(value: Any) -> tuple[str, str]:
    """Convierte `3107 A` en proyecto `3107` y subsistema `Principal A`."""
    raw = _cell_text(value)
    match = re.match(r"^\s*([0-9]+)\s+(.+?)\s*$", raw)
    if not match:
        return raw, "Principal"
    suffix = match.group(2).strip()
    return match.group(1), f"Principal {suffix}" if suffix else "Principal"


def _parse_project_matrix_workbook(workbook, source_filename: str) -> dict[str, Any]:
    """Lee la plantilla SISTEMAS A:AB usada actualmente por OGA.

    La hoja SISTEMAS contiene una fila por subsistema. Por ejemplo `3107 A` y
    `3107 B` pertenecen a la misma oportunidad 3107 y se convierten en
    `Principal A` y `Principal B`. Las columnas A:AB se conservan completas,
    incluidos los valores vacíos.
    """
    systems = _find_sheet(workbook, "SISTEMAS")
    header_row, headers = _find_header_row(systems, ("Proyecto", "Cliente", "Inicio", "Material a Transportar"))
    cols = {
        field: _column(headers, aliases, required=field in {"proyecto", "cliente", "inicio"})
        for field, aliases in PROJECT_MATRIX_HEADERS.items()
    }

    opportunities: "OrderedDict[str, dict[str, Any]]" = OrderedDict()
    subsystem_targets: dict[tuple[str, str], dict[str, Any]] = {}
    warnings: list[str] = []
    system_rows = entry_count = exit_count = 0

    subsystem_mapping = (
        "material_transportado", "flujo_kg_h", "distancia_horizontal_m",
        "distancia_vertical_m", "curvas_90", "distancia_unidad_soplado_m",
        "curvas_unidad_soplado", "tipo_flujo", "atex", "material_contacto",
        "voltaje_potencia", "tipo_transporte", "diametro_tuberia", "tipo_acople",
        "potencia_hp", "caudal_cfm", "diferencial_presion_psi", "tipo_bomba",
        "area_filtracion_m2", "micraje_filtracion", "consumo_aire_cfm",
        "presion_alimentacion_psi",
    )

    for row_number, row in enumerate(systems.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        raw_project = _read_cell(row, cols["proyecto"])
        project, subsystem_name = _split_project_subsystem(raw_project)
        if not project:
            continue
        system_rows += 1
        project_key = normalize_search(project)
        client = _cell_text(_read_cell(row, cols["cliente"]))
        start = _cell_text(_read_cell(row, cols["inicio"]))
        industry = _cell_text(_read_cell(row, cols["industria"]))

        opportunity = opportunities.get(project_key)
        if opportunity is None:
            opportunity = {
                "data": {
                    # radicado/nombre/estado se mantienen solo por compatibilidad
                    # interna con el esquema histórico; no son campos del formulario.
                    "radicado": project,
                    "proyecto": project,
                    "nombre": project,
                    "cliente": client,
                    "descripcion": "",
                    "planta": "",
                    "ciudad": "",
                    "pais": "",
                    "industria": industry,
                    "tipo_oportunidad": "",
                    "estado": "Nueva",
                    "fecha_inicio": start,
                    "valor_estimado": "",
                    "observaciones": "",
                },
                "subsystems": [],
                "source": {"rows": [row_number]},
            }
            opportunities[project_key] = opportunity
        else:
            opportunity["source"].setdefault("rows", []).append(row_number)
            if not opportunity["data"].get("cliente") and client:
                opportunity["data"]["cliente"] = client
            elif client and normalize_search(client) != normalize_search(opportunity["data"].get("cliente")):
                warnings.append(f"SISTEMAS fila {row_number}: el proyecto {project} tiene un cliente diferente; se conservó el primero.")
            # La fecha de proyecto representa el inicio más temprano de sus subsistemas.
            existing_start = str(opportunity["data"].get("fecha_inicio") or "")
            if start and (not existing_start or start < existing_start):
                opportunity["data"]["fecha_inicio"] = start
            if not opportunity["data"].get("industria") and industry:
                opportunity["data"]["industria"] = industry

        sub_data = {field: "" for field in SUBSYSTEM_FIELDS}
        sub_data["nombre"] = subsystem_name
        for field in subsystem_mapping:
            sub_data[field] = _cell_text(_read_cell(row, cols[field]))

        entry_value = _cell_text(_read_cell(row, cols["punto_ingreso"]))
        exit_value = _cell_text(_read_cell(row, cols["punto_destino"]))
        entry_points = []
        exit_points = []
        if entry_value:
            entry_points.append({"tipo": entry_value, "cantidad": "1", "restriccion_altura": ""})
            entry_count += 1
        if exit_value:
            exit_points.append({"tipo": exit_value, "cantidad": "1", "restriccion_altura": ""})
            exit_count += 1

        subsystem = {
            "data": sub_data,
            "source_raw_name": subsystem_name.casefold(),
            "source_row": row_number,
            "entry_points": entry_points,
            "exit_points": exit_points,
            "equipment": [],
            "synthetic": False,
        }
        opportunity["subsystems"].append(subsystem)
        subsystem_targets[(project_key, normalize_search(subsystem_name))] = subsystem

    equipment_rows = linked_equipment = 0
    try:
        equipment = _find_sheet(workbook, "EQUIPOS")
        eq_header_row, eq_headers = _find_header_row(
            equipment,
            ("Proyecto", "Nombre del subsistema", "Tipo de Equipo", "Referencia", "Cantidad"),
        )
        e_project = _column(eq_headers, ("Proyecto",), required=True)
        e_sub = _column(eq_headers, ("Nombre del subsistema",), required=True)
        e_type = _column(eq_headers, ("Tipo de Equipo",), required=True)
        e_ref = _column(eq_headers, ("Referencia",), required=True)
        e_qty = _column(eq_headers, ("Cantidad",), required=True)
        equipment_records: list[tuple[int, str, str, str, str, str]] = []
        groups_by_project: dict[str, list[str]] = {}
        for row_number, row in enumerate(equipment.iter_rows(min_row=eq_header_row + 1, values_only=True), start=eq_header_row + 1):
            project = _identifier(_read_cell(row, e_project))
            sub_name = _cell_text(_read_cell(row, e_sub))
            eq_type = _cell_text(_read_cell(row, e_type))
            reference = _cell_text(_read_cell(row, e_ref))
            quantity = _cell_text(_read_cell(row, e_qty))
            if not project and not sub_name and not eq_type and not reference:
                continue
            equipment_rows += 1
            project_key = normalize_search(project)
            group_key = normalize_search(sub_name)
            equipment_records.append((row_number, project_key, group_key, sub_name, eq_type, reference, quantity))
            groups = groups_by_project.setdefault(project_key, [])
            if group_key and group_key not in groups:
                groups.append(group_key)

        # En el histórico, los nombres de EQUIPOS no siempre coinciden literalmente
        # con SISTEMAS (p. ej. A1/A2 frente a A/B). Para la gran mayoría de los
        # proyectos existe la misma cantidad de grupos en ambas hojas; en ese caso
        # se relacionan por el orden en que aparecen, sin perder los matches exactos.
        equipment_aliases: dict[tuple[str, str], dict[str, Any]] = {}
        for project_key, opportunity in opportunities.items():
            targets = list(opportunity.get("subsystems") or [])
            groups = list(groups_by_project.get(project_key) or [])
            if not targets or not groups:
                continue
            used_target_ids: set[int] = set()
            for group_key in groups:
                exact = subsystem_targets.get((project_key, group_key))
                if exact is not None:
                    equipment_aliases[(project_key, group_key)] = exact
                    used_target_ids.add(id(exact))
            remaining_targets = [target for target in targets if id(target) not in used_target_ids]
            remaining_groups = [group for group in groups if (project_key, group) not in equipment_aliases]
            if len(targets) == 1:
                for group in remaining_groups:
                    equipment_aliases[(project_key, group)] = targets[0]
            elif len(remaining_groups) >= len(remaining_targets):
                for group, target in zip(remaining_groups, remaining_targets):
                    equipment_aliases[(project_key, group)] = target

            # Si EQUIPOS contiene un grupo adicional que no existe en SISTEMAS,
            # se conserva creando un grupo técnico con A-AB vacío. Esto evita
            # perder equipos válidos y hace visibles los campos faltantes como “—”.
            still_unmapped = [group for group in groups if (project_key, group) not in equipment_aliases]
            if still_unmapped:
                display_by_key = {
                    record[2]: record[3]
                    for record in equipment_records
                    if record[1] == project_key
                }
                for group in still_unmapped:
                    synthetic_data = {field: "" for field in SUBSYSTEM_FIELDS}
                    synthetic_data["nombre"] = display_by_key.get(group) or "Principal"
                    synthetic = {
                        "data": synthetic_data,
                        "source_raw_name": group,
                        "source_row": None,
                        "entry_points": [],
                        "exit_points": [],
                        "equipment": [],
                        "synthetic": True,
                    }
                    opportunity["subsystems"].append(synthetic)
                    equipment_aliases[(project_key, group)] = synthetic
                    subsystem_targets[(project_key, group)] = synthetic
                    warnings.append(
                        f"EQUIPOS: se creó el grupo '{synthetic_data['nombre']}' del proyecto {project_key} porque contiene equipos pero no tiene una fila A-AB en SISTEMAS."
                    )

        for row_number, project_key, group_key, sub_name, eq_type, reference, quantity in equipment_records:
            target = equipment_aliases.get((project_key, group_key)) or subsystem_targets.get((project_key, group_key))
            if not target:
                if len(warnings) < 100:
                    warnings.append(f"EQUIPOS fila {row_number}: no se pudo asociar proyecto {project_key} / subsistema '{sub_name}'.")
                continue
            target["equipment"].append({"tipo_equipo": eq_type, "referencia": reference, "cantidad": quantity})
            linked_equipment += 1
    except ValueError:
        # EQUIPOS es complementaria en esta versión de la plantilla. Los 28
        # campos A:AB de SISTEMAS siguen siendo importables si la hoja no existe.
        pass

    opportunity_list = list(opportunities.values())
    summary = {
        "oportunidades": len(opportunity_list),
        "radicados_unicos": len(opportunity_list),
        "subsistemas": sum(len(item["subsystems"]) for item in opportunity_list),
        "entradas": entry_count,
        "salidas": exit_count,
        "equipos": linked_equipment,
        "filas_sistemas": system_rows,
        "filas_entradas": entry_count,
        "filas_salidas": exit_count,
        "filas_equipos": equipment_rows,
        "no_asociados": max(0, equipment_rows - linked_equipment),
    }
    return {
        "version": 2,
        "source_format": "SISTEMAS_A_AB",
        "source_filename": str(source_filename or "SISTEMAS.xlsx"),
        "parsed_at": datetime.now(timezone.utc).replace(microsecond=0).isoformat(),
        "summary": summary,
        "warning_counts": {"equipo_sin_subsistema": max(0, equipment_rows - linked_equipment)},
        "warnings": warnings,
        "opportunities": opportunity_list,
    }


def parse_workbook(file_obj, source_filename: str) -> dict[str, Any]:
    """Detecta automáticamente la plantilla actual A:AB o la histórica."""
    try:
        if hasattr(file_obj, "seek"):
            file_obj.seek(0)
        workbook = load_workbook(file_obj, data_only=True, read_only=True)
    except Exception as exc:
        raise ValueError(f"No fue posible abrir el Excel: {exc}") from exc

    try:
        systems = _find_sheet(workbook, "SISTEMAS")
        try:
            _, headers = _find_header_row(systems, ("Proyecto", "Cliente", "Inicio", "Material a Transportar"))
            is_current = _header_key("Número de Radicado") not in headers and _header_key("PUNTO DE INGRESO") in headers
        except ValueError:
            is_current = False
        if is_current:
            return _parse_project_matrix_workbook(workbook, source_filename)
    finally:
        workbook.close()

    if hasattr(file_obj, "seek"):
        file_obj.seek(0)
    return _parse_legacy_workbook(file_obj, source_filename)


def _parse_legacy_workbook(file_obj, source_filename: str) -> dict[str, Any]:
    try:
        if hasattr(file_obj, "seek"):
            file_obj.seek(0)
        workbook = load_workbook(file_obj, data_only=True, read_only=True)
    except Exception as exc:
        raise ValueError(f"No fue posible abrir el Excel: {exc}") from exc

    warnings: list[str] = []
    warning_counts: dict[str, int] = {
        "estado_no_reconocido": 0,
        "entrada_sin_subsistema": 0,
        "salida_sin_subsistema": 0,
        "equipo_sin_subsistema": 0,
        "asociacion_corregida": 0,
        "subsistema_sintetico": 0,
        "datos_generales_inconsistentes": 0,
    }

    systems = _find_sheet(workbook, "SISTEMAS")
    entries = _find_sheet(workbook, "PUNTOS DE ENTRADA")
    exits = _find_sheet(workbook, "PUNTOS DE SALIDA")
    equipment = _find_sheet(workbook, "Equipos")

    header_row, system_headers = _find_header_row(systems, ("Número de Radicado", "Proyecto", "Cliente", "Nombre del Proceso"))
    sys_cols: dict[str, int | None] = {}
    for field, aliases in SYSTEM_FIELD_HEADERS.items():
        sys_cols[field] = _column(system_headers, aliases, required=field in {"radicado", "proyecto", "cliente", "nombre_subsistema"})

    opportunities: "OrderedDict[tuple[str, str], dict[str, Any]]" = OrderedDict()
    exact_targets: dict[tuple[str, str], list[dict[str, Any]]] = {}
    normalized_targets: dict[tuple[str, str], list[dict[str, Any]]] = {}
    targets_by_radicado: dict[str, list[dict[str, Any]]] = {}
    system_rows = 0

    for row_number, row in enumerate(systems.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        radicado = _identifier(_read_cell(row, sys_cols["radicado"]))
        proyecto = _identifier(_read_cell(row, sys_cols["proyecto"]))
        if not radicado and not proyecto:
            continue
        if not radicado or not proyecto:
            warnings.append(f"SISTEMAS fila {row_number}: se omitió una fila sin radicado o proyecto.")
            continue
        system_rows += 1
        key = (normalize_search(radicado), normalize_search(proyecto))

        cliente = _cell_text(_read_cell(row, sys_cols["cliente"]))
        planta = _cell_text(_read_cell(row, sys_cols["planta"]))
        contacto = _cell_text(_read_cell(row, sys_cols["contacto"]))
        inicio = _cell_text(_read_cell(row, sys_cols["inicio"]))
        estado_original = _cell_text(_read_cell(row, sys_cols["estado_origen"]))
        estado, status_known = _map_status(estado_original)
        if estado_original and not status_known:
            warning_counts["estado_no_reconocido"] += 1
            if warning_counts["estado_no_reconocido"] <= 20:
                warnings.append(f"SISTEMAS fila {row_number}: estado '{estado_original}' no reconocido; se importará como Nueva.")
        tipo_oportunidad = _cell_text(_read_cell(row, sys_cols["tipo_oportunidad"]))
        descripcion = _cell_text(_read_cell(row, sys_cols["descripcion"]))
        raw_subsystem_value = _read_cell(row, sys_cols["nombre_subsistema"])
        subsystem_name = _cell_text(raw_subsystem_value) or "Principal"

        if key not in opportunities:
            obs_parts: list[str] = []
            if contacto:
                obs_parts.append(f"Contacto histórico: {contacto}.")
            opportunities[key] = {
                "data": {
                    "radicado": radicado,
                    "proyecto": proyecto,
                    "nombre": descripcion or f"Proyecto {proyecto} - {cliente}",
                    "cliente": cliente,
                    "descripcion": descripcion,
                    "planta": planta,
                    "ciudad": "",
                    "pais": "",
                    "industria": "",
                    "tipo_oportunidad": tipo_oportunidad,
                    "estado": estado,
                    "fecha_inicio": inicio,
                    "valor_estimado": "",
                    "observaciones": " ".join(obs_parts),
                },
                "source": {
                    "estado_original": estado_original,
                    "contactos": [contacto] if contacto else [],
                },
                "subsystems": [],
            }
        else:
            opportunity = opportunities[key]
            base = opportunity["data"]
            comparisons = {
                "cliente": cliente,
                "planta": planta,
                "tipo_oportunidad": tipo_oportunidad,
                "fecha_inicio": inicio,
                "estado": estado,
            }
            differences = [name for name, value in comparisons.items() if value and base.get(name) and normalize_search(value) != normalize_search(base.get(name))]
            if differences:
                warning_counts["datos_generales_inconsistentes"] += 1
                if warning_counts["datos_generales_inconsistentes"] <= 20:
                    warnings.append(
                        f"Radicado {radicado} / proyecto {proyecto}: hay diferencias entre subsistemas en {', '.join(differences)}; se conservará el primer valor."
                    )
            if contacto and contacto not in opportunity["source"]["contactos"]:
                opportunity["source"]["contactos"].append(contacto)

        sub_data: dict[str, str] = {field: "" for field in SUBSYSTEM_FIELDS}
        sub_data["nombre"] = subsystem_name
        sub_data["nombre_proceso"] = descripcion
        sub_data["descripcion_proceso"] = _cell_text(_read_cell(row, sys_cols["descripcion_proceso"]))
        for field in SUBSYSTEM_FIELDS:
            if field in {"nombre", "nombre_proceso", "descripcion_proceso", "observaciones"}:
                continue
            if field in sys_cols:
                sub_data[field] = _cell_text(_read_cell(row, sys_cols[field]))
        subsystem = {
            "data": sub_data,
            "source_raw_name": _raw_relation_text(raw_subsystem_value),
            "source_row": row_number,
            "entry_points": [],
            "exit_points": [],
            "equipment": [],
        }
        opportunities[key]["subsystems"].append(subsystem)

        rel_radicado = normalize_search(radicado)
        exact_key = (rel_radicado, subsystem["source_raw_name"])
        exact_targets.setdefault(exact_key, []).append(subsystem)
        norm_key = (rel_radicado, normalize_search(subsystem_name))
        normalized_targets.setdefault(norm_key, []).append(subsystem)
        targets_by_radicado.setdefault(rel_radicado, []).append(subsystem)

    if not opportunities:
        raise ValueError("La hoja SISTEMAS no contiene oportunidades válidas para importar.")

    opportunity_keys_by_radicado: dict[str, list[tuple[str, str]]] = {}
    for opportunity_key in opportunities:
        opportunity_keys_by_radicado.setdefault(opportunity_key[0], []).append(opportunity_key)
    fallback_notes: set[tuple[str, str, str]] = set()

    def resolve_subsystem(radicado_value: Any, subsystem_value: Any, source_name: str, row_number: int) -> dict[str, Any] | None:
        rad = normalize_search(_identifier(radicado_value))
        raw_name = _raw_relation_text(subsystem_value)
        display_name = _cell_text(subsystem_value) or "Sin asignar"
        exact = exact_targets.get((rad, raw_name), [])
        if len(exact) == 1:
            return exact[0]
        normalized_name = normalize_search(display_name)
        normal = normalized_targets.get((rad, normalized_name), [])
        if len(normal) == 1:
            return normal[0]

        # Si el radicado tiene un solo subsistema, la discrepancia de nombre se
        # interpreta como un error de digitación del histórico y se conserva el
        # dato asociándolo al único destino posible.
        rad_targets = targets_by_radicado.get(rad, [])
        if len(rad_targets) == 1:
            note_key = (rad, normalized_name, "sole")
            if note_key not in fallback_notes:
                fallback_notes.add(note_key)
                warning_counts["asociacion_corregida"] += 1
                warnings.append(
                    f"{source_name} fila {row_number}: '{display_name}' no coincide con el subsistema de radicado {_cell_text(radicado_value)}; se asoció automáticamente al único subsistema disponible '{rad_targets[0]['data']['nombre']}'."
                )
            return rad_targets[0]

        # Si el radicado pertenece a una única OD pero hay varios subsistemas y
        # ninguno coincide, no descartamos la información: se crea un subsistema
        # técnico claramente marcado como proveniente de filas huérfanas.
        opportunity_keys = opportunity_keys_by_radicado.get(rad, [])
        if len(opportunity_keys) == 1:
            opportunity = opportunities[opportunity_keys[0]]
            synthetic_data: dict[str, str] = {field: "" for field in SUBSYSTEM_FIELDS}
            synthetic_data["nombre"] = display_name
            synthetic_data["observaciones"] = (
                "Subsistema creado automáticamente durante el cargue histórico porque "
                "el nombre no coincidía con los subsistemas de la hoja SISTEMAS."
            )
            synthetic = {
                "data": synthetic_data,
                "source_raw_name": raw_name,
                "source_row": None,
                "entry_points": [],
                "exit_points": [],
                "equipment": [],
                "synthetic": True,
            }
            opportunity["subsystems"].append(synthetic)
            exact_targets.setdefault((rad, raw_name), []).append(synthetic)
            normalized_targets.setdefault((rad, normalized_name), []).append(synthetic)
            targets_by_radicado.setdefault(rad, []).append(synthetic)
            warning_counts["subsistema_sintetico"] += 1
            warnings.append(
                f"{source_name} fila {row_number}: se creó el subsistema '{display_name}' para radicado {_cell_text(radicado_value)} con el fin de conservar registros que no tenían correspondencia en SISTEMAS."
            )
            return synthetic
        return None

    def parse_points(sheet, target_key: str, warning_key: str) -> tuple[int, int]:
        hdr_row, headers = _find_header_row(sheet, ("Número de radicado", "Nombre del subsistema", "Tipo", "Cantidad"))
        c_rad = _column(headers, ("Número de radicado", "Numero de radicado"), required=True)
        c_sub = _column(headers, ("Nombre del subsistema",), required=True)
        c_type = _column(headers, ("Tipo",), required=True)
        c_qty = _column(headers, ("Cantidad",), required=True)
        c_height = _column(headers, ("Restriccion de altura", "Restricción de altura"), required=False)
        total = linked = 0
        for row_number, row in enumerate(sheet.iter_rows(min_row=hdr_row + 1, values_only=True), start=hdr_row + 1):
            rad = _read_cell(row, c_rad)
            sub = _read_cell(row, c_sub)
            kind = _cell_text(_read_cell(row, c_type))
            if not _cell_text(rad) and not _cell_text(sub) and not kind:
                continue
            total += 1
            target = resolve_subsystem(rad, sub, sheet.title, row_number)
            if not target:
                warning_counts[warning_key] += 1
                if warning_counts[warning_key] <= 25:
                    warnings.append(f"{sheet.title} fila {row_number}: no se pudo asociar radicado {_cell_text(rad)} / subsistema '{_cell_text(sub)}'.")
                continue
            target[target_key].append({
                "tipo": kind,
                "cantidad": _cell_text(_read_cell(row, c_qty)),
                "restriccion_altura": _cell_text(_read_cell(row, c_height)),
            })
            linked += 1
        return total, linked

    entry_rows, linked_entries = parse_points(entries, "entry_points", "entrada_sin_subsistema")
    exit_rows, linked_exits = parse_points(exits, "exit_points", "salida_sin_subsistema")

    eq_header_row, eq_headers = _find_header_row(equipment, ("Número de radicado", "Nombre del subsistema", "Tipo de Equipo", "Referencia", "Cantidad"))
    e_rad = _column(eq_headers, ("Número de radicado", "Numero de radicado"), required=True)
    e_sub = _column(eq_headers, ("Nombre del subsistema",), required=True)
    e_type = _column(eq_headers, ("Tipo de Equipo",), required=True)
    e_ref = _column(eq_headers, ("Referencia",), required=True)
    e_qty = _column(eq_headers, ("Cantidad",), required=True)
    equipment_rows = linked_equipment = 0
    for row_number, row in enumerate(equipment.iter_rows(min_row=eq_header_row + 1, values_only=True), start=eq_header_row + 1):
        rad = _read_cell(row, e_rad)
        sub = _read_cell(row, e_sub)
        eq_type = _cell_text(_read_cell(row, e_type))
        reference = _cell_text(_read_cell(row, e_ref))
        if not _cell_text(rad) and not _cell_text(sub) and not eq_type and not reference:
            continue
        equipment_rows += 1
        target = resolve_subsystem(rad, sub, equipment.title, row_number)
        if not target:
            warning_counts["equipo_sin_subsistema"] += 1
            if warning_counts["equipo_sin_subsistema"] <= 25:
                warnings.append(f"Equipos fila {row_number}: no se pudo asociar radicado {_cell_text(rad)} / subsistema '{_cell_text(sub)}'.")
            continue
        target["equipment"].append({
            "tipo_equipo": eq_type,
            "referencia": reference,
            "cantidad": _cell_text(_read_cell(row, e_qty)),
        })
        linked_equipment += 1

    opportunity_list = list(opportunities.values())
    # Al finalizar consolidamos los contactos históricos sin crear campos nuevos en la Fase 2.
    for item in opportunity_list:
        contacts = [value for value in item["source"].get("contactos", []) if value]
        if len(contacts) > 1:
            item["data"]["observaciones"] += " Contactos históricos: " + ", ".join(contacts) + "."

    unique_radicados = len({normalize_search(item["data"]["radicado"]) for item in opportunity_list})
    subsystem_count = sum(len(item["subsystems"]) for item in opportunity_list)
    summary = {
        "oportunidades": len(opportunity_list),
        "radicados_unicos": unique_radicados,
        "subsistemas": subsystem_count,
        "entradas": linked_entries,
        "salidas": linked_exits,
        "equipos": linked_equipment,
        "filas_sistemas": system_rows,
        "filas_entradas": entry_rows,
        "filas_salidas": exit_rows,
        "filas_equipos": equipment_rows,
        "no_asociados": (entry_rows - linked_entries) + (exit_rows - linked_exits) + (equipment_rows - linked_equipment),
    }
    return {
        "version": 1,
        "source_filename": str(source_filename or "plantilla.xlsx"),
        "parsed_at": datetime.now(timezone.utc).replace(microsecond=0).isoformat(),
        "summary": summary,
        "warning_counts": warning_counts,
        "warnings": warnings,
        "opportunities": opportunity_list,
    }


def _cleanup_old_stages() -> None:
    STAGE_DIR.mkdir(parents=True, exist_ok=True)
    now = datetime.now(timezone.utc).timestamp()
    max_age = MAX_STAGE_AGE_HOURS * 3600
    for path in STAGE_DIR.glob("*.json"):
        try:
            if now - path.stat().st_mtime > max_age:
                path.unlink(missing_ok=True)
        except OSError:
            pass


def save_stage(payload: dict[str, Any], owner_user_id: Any) -> str:
    _cleanup_old_stages()
    STAGE_DIR.mkdir(parents=True, exist_ok=True)
    token = secrets.token_urlsafe(24)
    wrapper = {
        "owner_user_id": str(owner_user_id or ""),
        "saved_at": datetime.now(timezone.utc).replace(microsecond=0).isoformat(),
        "payload": payload,
    }
    target = STAGE_DIR / f"{token}.json"
    target.write_text(json.dumps(wrapper, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    return token


def load_stage(token: str, owner_user_id: Any) -> dict[str, Any]:
    _cleanup_old_stages()
    if not re.fullmatch(r"[A-Za-z0-9_-]{20,80}", str(token or "")):
        raise ValueError("Importación temporal no válida.")
    target = STAGE_DIR / f"{token}.json"
    if not target.is_file():
        raise ValueError("La previsualización expiró o ya fue procesada.")
    try:
        wrapper = json.loads(target.read_text(encoding="utf-8"))
    except Exception as exc:
        raise ValueError("No fue posible leer la previsualización temporal.") from exc
    if str(wrapper.get("owner_user_id") or "") != str(owner_user_id or ""):
        raise PermissionError("Esta importación pertenece a otra sesión de usuario.")
    return dict(wrapper.get("payload") or {})


def delete_stage(token: str, owner_user_id: Any | None = None) -> None:
    if not re.fullmatch(r"[A-Za-z0-9_-]{20,80}", str(token or "")):
        return
    target = STAGE_DIR / f"{token}.json"
    if not target.is_file():
        return
    if owner_user_id is not None:
        try:
            wrapper = json.loads(target.read_text(encoding="utf-8"))
            if str(wrapper.get("owner_user_id") or "") != str(owner_user_id or ""):
                return
        except Exception:
            return
    target.unlink(missing_ok=True)
