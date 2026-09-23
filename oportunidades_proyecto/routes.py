from __future__ import annotations

from functools import wraps
from pathlib import Path
import json
import mimetypes
import re
import shutil
import unicodedata
import uuid

from flask import Blueprint, abort, flash, g, jsonify, redirect, render_template, request, send_file, url_for
from werkzeug.utils import secure_filename
from openpyxl import load_workbook

from auth_store import STATUS_ACTIVE, audit_event, list_users
from .db import (
    ATTACHMENT_IMAGE, ATTACHMENT_OFFER, DATA_DIR, FILES_DIR, STATUSES, SUBSYSTEM_FIELDS, VISIBLE_TECH_FIELDS,
    add_equipment, add_point, bulk_import_payload, create_attachment, create_opportunity, create_subsystem,
    delete_child, ensure_database, get_attachment, get_opportunity, get_opportunity_detail,
    get_subsystem, preview_bulk_import, search_opportunities, soft_delete_attachment, soft_delete_opportunity,
    soft_delete_subsystem, stats, update_opportunity, update_opportunity_with_visible_technical,
    update_subsystem, update_visible_subsystem,
)
from .import_excel import delete_stage, load_stage, parse_workbook, save_stage
from .search import parse_query, query_help_fields

bp = Blueprint("oportunidades", __name__, url_prefix="/oportunidades-proyecto")

MAX_PDF_BYTES = 50 * 1024 * 1024
MAX_IMAGE_BYTES = 20 * 1024 * 1024
MAX_IMPORT_BYTES = 25 * 1024 * 1024
ALLOWED_IMAGES = {".jpg", ".jpeg", ".png", ".webp"}
COUNTRIES_FILE = Path(__file__).resolve().parent.parent / "data" / "paises_y_10_ciudades_importantes.json"
INDUSTRIA_DATA_DIR = Path(__file__).resolve().parent.parent / "data" / "INDUSTRIA"
_TECH_OPTIONS_CACHE: dict[str, tuple[int, int, list[str]]] = {}

CRITERIA_LABELS = (
    ("nombre_proceso", "Nombre del proceso"),
    ("descripcion_proceso", "Descripción del proceso"),
    ("material_transportado", "Material transportado"),
    ("flujo_kg_h", "Flujo (kg/h)"),
    ("distancia_horizontal_m", "Distancia horizontal (m)"),
    ("distancia_vertical_m", "Distancia vertical (m)"),
    ("curvas_90", "Curvas de tubería x 90°"),
    ("voltaje_potencia", "Voltaje de potencia (V)"),
    ("material_contacto", "Material en contacto con producto"),
    ("material_estructural", "Material estructural"),
    ("tipo_flujo", "Tipo de flujo"),
    ("pesaje_oga", "Pesaje OGA"),
    ("atex", "Clasificación ATEX"),
    ("nec", "Clasificación NEC"),
    ("ubicacion", "Ubicación"),
    ("aire_comprimido", "Disponibilidad aire comprimido"),
    ("distancia_unidad_soplado_m", "Distancia unidad de soplado/vacío (m)"),
    ("curvas_unidad_soplado", "Curvas unidad de soplado/vacío"),
    ("preferencia_tipologia", "Preferencia de tipología"),
    ("preferencia_acoples", "Preferencia de acoples"),
    ("tipo_transporte", "Tipo de transporte"),
    ("diametro_tuberia", "Diámetro tubería transporte"),
    ("tipo_acople", "Tipo de acople"),
    ("potencia_hp", "Potencia (hp)"),
    ("caudal_cfm", "Caudal (CFM)"),
    ("diferencial_presion_psi", "Diferencial de presión (PSI)"),
    ("tipo_bomba", "Tipo de bomba"),
    ("area_filtracion_m2", "Área de filtración (m²)"),
    ("micraje_filtracion", "Micraje de filtración"),
    ("consumo_aire_cfm", "Consumo de aire (CFM)"),
    ("presion_alimentacion_psi", "Presión de alimentación (PSI)"),
    ("es_multiequipos", "Es multiequipos"),
    ("observaciones", "Observaciones técnicas"),
)


def _admin_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        user = getattr(g, "current_user", None)
        if not user or not user.get("is_admin"):
            abort(403)
        return view(*args, **kwargs)
    return wrapped


def _user() -> dict:
    return dict(getattr(g, "current_user", None) or {})


def _client_ip() -> str:
    return request.remote_addr or "DESCONOCIDA"


def _country_options() -> list[str]:
    """Carga los países del JSON maestro. Ciudad permanece como texto libre."""
    try:
        payload = json.loads(COUNTRIES_FILE.read_text(encoding="utf-8-sig"))
    except (OSError, json.JSONDecodeError):
        return []
    countries = [
        str(name).strip()
        for name, cities in payload.items()
        if not str(name).startswith("_") and str(name).strip() and isinstance(cities, list)
    ]
    return sorted(dict.fromkeys(countries), key=lambda value: value.casefold())


def _norm_header(value: object) -> str:
    text_value = "" if value is None else str(value).strip()
    text_value = unicodedata.normalize("NFKD", text_value).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^A-Z0-9]+", " ", text_value.upper()).strip()


def _catalog_values(filename: str, aliases: tuple[str, ...], *, numeric: bool = False) -> list[str]:
    """Lee una columna de los Excel maestros de Industria sin duplicar sus datos en OD.

    El cache depende de fecha/tamaño del archivo: cuando Industria guarda o importa el
    Excel, la siguiente carga de Oportunidades toma automáticamente el catálogo nuevo.
    """
    path = INDUSTRIA_DATA_DIR / filename
    try:
        stat = path.stat()
    except OSError:
        return []

    cache_key = f"{filename}|{'|'.join(aliases)}|{int(numeric)}"
    cached = _TECH_OPTIONS_CACHE.get(cache_key)
    if cached and cached[0] == stat.st_mtime_ns and cached[1] == stat.st_size:
        return list(cached[2])

    alias_norms = {_norm_header(alias) for alias in aliases}
    values: list[str] = []
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return []
    try:
        ws = wb.active
        header_row = None
        value_col = None
        # Tanto los Excel históricos como los migrados de Industria tienen el
        # encabezado dentro de las primeras filas.
        for row_num, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 12), values_only=True), start=1):
            for idx, cell in enumerate(row):
                if _norm_header(cell) in alias_norms:
                    header_row = row_num
                    value_col = idx
                    break
            if value_col is not None:
                break
        if value_col is None:
            return []

        seen: set[str] = set()
        for row in ws.iter_rows(min_row=(header_row or 1) + 1, values_only=True):
            if value_col >= len(row):
                continue
            raw = row[value_col]
            if raw is None:
                continue
            if isinstance(raw, float) and raw.is_integer():
                value = str(int(raw))
            else:
                value = str(raw).strip()
            if not value:
                continue
            identity = value.casefold()
            if identity in seen:
                continue
            seen.add(identity)
            values.append(value)
    finally:
        wb.close()

    if numeric:
        def numeric_key(value: str):
            try:
                return (0, float(value.replace(",", ".")), value.casefold())
            except ValueError:
                return (1, float("inf"), value.casefold())
        values.sort(key=numeric_key)
    else:
        values.sort(key=str.casefold)

    _TECH_OPTIONS_CACHE[cache_key] = (stat.st_mtime_ns, stat.st_size, list(values))
    return values


def _technical_catalog_options() -> dict[str, list[str]]:
    """Catálogos de Industria usados como selección en datos técnicos de OD.

    Solo se convierten en selector los campos que realmente tienen un catálogo
    de valores. Los criterios numéricos de CS son reglas/valores de cálculo y
    permanecen como entrada libre en Oportunidades.
    """
    return {
        "voltaje_potencia": _catalog_values(
            "CB_MATRIZ.xlsx",
            ("Voltaje De Potencia", "Voltaje de potencia (V)", "Voltaje de potencia"),
            numeric=True,
        ),
        "material_transportado": _catalog_values(
            "TIPOS_DE_MATERIALES.xlsx",
            ("Nombre común",),
        ),
    }


def _active_users() -> list[dict]:
    return [user for user in list_users() if user.get("status") == STATUS_ACTIVE and user.get("is_active")]


def _selected_users(ids: list[str]) -> list[dict]:
    selected = {str(item) for item in ids if item}
    return [user for user in _active_users() if str(user.get("id")) in selected]


def _opportunity_form() -> dict[str, str]:
    return {
        "radicado": (request.form.get("radicado") or "").strip(),
        "proyecto": (request.form.get("proyecto") or "").strip(),
        "nombre": (request.form.get("nombre") or "").strip(),
        "cliente": (request.form.get("cliente") or "").strip(),
        "descripcion": (request.form.get("descripcion") or "").strip(),
        "planta": (request.form.get("planta") or "").strip(),
        "ciudad": (request.form.get("ciudad") or "").strip(),
        "pais": (request.form.get("pais") or "").strip(),
        "industria": (request.form.get("industria") or "").strip(),
        "tipo_oportunidad": (request.form.get("tipo_oportunidad") or "").strip(),
        "estado": (request.form.get("estado") or "Nueva").strip(),
        "fecha_inicio": (request.form.get("fecha_inicio") or "").strip(),
        "valor_estimado": (request.form.get("valor_estimado") or "").strip(),
        "observaciones": (request.form.get("observaciones") or "").strip(),
    }


def _subsystem_form(prefix: str = "") -> dict[str, str]:
    values: dict[str, str] = {}
    for field in SUBSYSTEM_FIELDS:
        values[field] = (request.form.get(f"{prefix}{field}") or "").strip()
    values["nombre"] = values["nombre"] or "Principal"
    return values


def _visible_technical_form(prefix: str = "") -> dict[str, str]:
    return {
        field: (request.form.get(f"{prefix}{field}") or "").strip()
        for field in VISIBLE_TECH_FIELDS
    }


def _visible_technical_updates_from_edit() -> dict[int, dict[str, str]]:
    updates: dict[int, dict[str, str]] = {}
    for raw_id in request.form.getlist("technical_subsystem_id"):
        try:
            subsystem_id = int(raw_id)
        except (TypeError, ValueError):
            continue
        updates[subsystem_id] = _visible_technical_form(f"tech_{subsystem_id}_")
    return updates


def _validate_opportunity(data: dict[str, str]) -> None:
    if not data["radicado"]:
        raise ValueError("Ingrese el número de radicado.")
    if not data["proyecto"]:
        raise ValueError("Ingrese el número o código de proyecto.")
    if not data["cliente"]:
        raise ValueError("Ingrese el cliente.")
    if not data["nombre"]:
        raise ValueError("Ingrese el nombre de la oportunidad.")
    if data["estado"] not in STATUSES:
        raise ValueError("Estado de oportunidad no válido.")
    limits = {"radicado": 60, "proyecto": 60, "nombre": 180, "cliente": 180, "descripcion": 1200,
              "planta": 120, "ciudad": 120, "pais": 120, "industria": 140, "tipo_oportunidad": 100,
              "valor_estimado": 80, "observaciones": 2000}
    for key, limit in limits.items():
        if len(data.get(key, "")) > limit:
            raise ValueError(f"El campo {key.replace('_', ' ')} supera el tamaño permitido.")


def _decorate_detail(detail: dict | None) -> dict | None:
    if not detail:
        return None
    for subsystem in detail.get("subsystems", []):
        subsystem["criteria"] = [
            {"key": key, "label": label, "value": subsystem.get(key)}
            for key, label in CRITERIA_LABELS if str(subsystem.get(key) or "").strip()
        ]
    return detail


def _safe_attachment_path(relative_path: str) -> Path:
    candidate = (DATA_DIR / str(relative_path or "")).resolve()
    root = DATA_DIR.resolve()
    try:
        candidate.relative_to(root)
    except ValueError:
        abort(404)
    return candidate


def _validate_file(file_storage, kind: str) -> tuple[str, str, bytes]:
    if not file_storage or not file_storage.filename:
        raise ValueError("Seleccione un archivo.")
    original = secure_filename(Path(file_storage.filename).name)
    if not original:
        raise ValueError("El nombre del archivo no es válido.")
    ext = Path(original).suffix.lower()
    stream = file_storage.stream
    position = stream.tell()
    header = stream.read(32)
    stream.seek(position)
    if kind == ATTACHMENT_OFFER:
        if ext != ".pdf" or not header.startswith(b"%PDF-"):
            raise ValueError("Las ofertas deben adjuntarse en formato PDF válido.")
        return original, "application/pdf", header
    if ext not in ALLOWED_IMAGES:
        raise ValueError("Las imágenes deben ser JPG, PNG o WEBP.")
    valid = (
        header.startswith(b"\xff\xd8\xff") or
        header.startswith(b"\x89PNG\r\n\x1a\n") or
        (header.startswith(b"RIFF") and b"WEBP" in header[:16])
    )
    if not valid:
        raise ValueError("El archivo no parece ser una imagen válida.")
    return original, mimetypes.guess_type(original)[0] or "application/octet-stream", header


def _save_attachment(file_storage, kind: str, opportunity_id: int, subsystem_id: int, title: str = "") -> int:
    original, mime, _ = _validate_file(file_storage, kind)
    ext = Path(original).suffix.lower()
    folder_name = "ofertas" if kind == ATTACHMENT_OFFER else "imagenes"
    folder = FILES_DIR / str(opportunity_id) / str(subsystem_id) / folder_name
    folder.mkdir(parents=True, exist_ok=True)
    target = folder / f"{uuid.uuid4().hex}{ext}"
    file_storage.save(target)
    size = target.stat().st_size
    max_size = MAX_PDF_BYTES if kind == ATTACHMENT_OFFER else MAX_IMAGE_BYTES
    if size <= 0 or size > max_size:
        target.unlink(missing_ok=True)
        label = "PDF" if kind == ATTACHMENT_OFFER else "imagen"
        raise ValueError(f"El {label} está vacío o supera el límite permitido.")
    relative = str(target.relative_to(DATA_DIR)).replace("\\", "/")
    try:
        return create_attachment(
            subsystem_id,
            tipo=kind,
            titulo=(title or Path(original).stem).strip()[:180],
            original_name=original,
            relative_path=relative,
            mime_type=mime,
            size=size,
            user=_user(),
        )
    except Exception:
        target.unlink(missing_ok=True)
        raise


@bp.get("/")
def index():
    ensure_database()
    query = (request.args.get("q") or "").strip()
    rows = search_opportunities(parse_query(query), limit=500)
    selected_id = request.args.get("od", type=int)
    detail = _decorate_detail(get_opportunity_detail(selected_id)) if selected_id else None
    return render_template(
        "oportunidades/index.html",
        opportunities=rows,
        selected=detail,
        query=query,
        od_stats=stats(),
        statuses=STATUSES,
        countries=_country_options(),
        technical_options=_technical_catalog_options(),
        active_users=_active_users(),
        search_fields=list(query_help_fields()),
        opportunities_page=True,
    )



@bp.get("/importar", endpoint="import_page")
@_admin_required
def import_page():
    ensure_database()
    return render_template(
        "oportunidades/import.html",
        opportunities_page=True,
        import_payload=None,
        comparison=None,
        import_token=None,
    )


@bp.post("/importar/analizar")
@_admin_required
def import_analyze():
    upload = request.files.get("archivo_excel")
    if not upload or not upload.filename:
        flash("Seleccione el archivo Excel de cargue masivo.", "error")
        return redirect(url_for("oportunidades.import_page"))
    filename = secure_filename(Path(upload.filename).name)
    if Path(filename).suffix.lower() != ".xlsx":
        flash("La subida masiva acepta únicamente archivos .xlsx.", "error")
        return redirect(url_for("oportunidades.import_page"))
    try:
        stream = upload.stream
        position = stream.tell()
        stream.seek(0, 2)
        size = stream.tell()
        stream.seek(position)
        if size <= 0 or size > MAX_IMPORT_BYTES:
            raise ValueError("El Excel está vacío o supera el límite de 25 MB.")
        payload = parse_workbook(stream, filename)
        token = save_stage(payload, _user().get("id"))
        return redirect(url_for("oportunidades.import_preview", token=token))
    except Exception as exc:
        audit_event("ANALIZAR CARGUE MASIVO", "OPORTUNIDADES DE PROYECTO", str(exc), "ERROR", _user(), _client_ip())
        g.audit_logged = True
        flash(str(exc), "error")
        return redirect(url_for("oportunidades.import_page"))


@bp.get("/importar/<token>")
@_admin_required
def import_preview(token: str):
    try:
        payload = load_stage(token, _user().get("id"))
    except PermissionError:
        abort(403)
    except Exception as exc:
        flash(str(exc), "error")
        return redirect(url_for("oportunidades.import_page"))
    comparison = preview_bulk_import(payload)
    return render_template(
        "oportunidades/import.html",
        opportunities_page=True,
        import_payload=payload,
        comparison=comparison,
        import_token=token,
    )


@bp.post("/importar/<token>/confirmar")
@_admin_required
def import_confirm(token: str):
    try:
        payload = load_stage(token, _user().get("id"))
        result = bulk_import_payload(payload, _user())
        source = str(payload.get("source_filename") or "Excel")
        detail = (
            f"{source} · {result['oportunidades_creadas']} OD nuevas · "
            f"{result['oportunidades_combinadas']} existentes combinadas · "
            f"{result['subsistemas_creados']} subsistemas nuevos · "
            f"{result['equipos_agregados']} equipos agregados"
        )
        audit_event("CARGUE MASIVO", "OPORTUNIDADES DE PROYECTO", detail, "OK", _user(), _client_ip())
        g.audit_logged = True
        delete_stage(token, _user().get("id"))
        flash(
            "Carga masiva completada: "
            f"{result['oportunidades_creadas']} oportunidades nuevas, "
            f"{result['oportunidades_combinadas']} existentes revisadas, "
            f"{result['subsistemas_creados']} subsistemas, "
            f"{result['entradas_agregadas']} entradas, "
            f"{result['salidas_agregadas']} salidas y "
            f"{result['equipos_agregados']} equipos agregados.",
            "ok",
        )
        return redirect(url_for("oportunidades.index"))
    except PermissionError:
        abort(403)
    except Exception as exc:
        audit_event("CARGUE MASIVO", "OPORTUNIDADES DE PROYECTO", str(exc), "ERROR", _user(), _client_ip())
        g.audit_logged = True
        flash(f"No se realizó la importación: {exc}", "error")
        return redirect(url_for("oportunidades.import_preview", token=token))


@bp.post("/importar/<token>/cancelar")
@_admin_required
def import_cancel(token: str):
    delete_stage(token, _user().get("id"))
    flash("Previsualización descartada. No se modificó la base de datos.", "ok")
    return redirect(url_for("oportunidades.import_page"))


@bp.get("/api/buscar")
def api_search():
    query = (request.args.get("q") or "").strip()
    rows = search_opportunities(parse_query(query), limit=500)
    return jsonify({"ok": True, "q": query, "rows": rows, "total": len(rows)})


@bp.get("/api/<int:opportunity_id>/detalle")
def api_detail_html(opportunity_id: int):
    detail = _decorate_detail(get_opportunity_detail(opportunity_id))
    if not detail:
        return jsonify({"ok": False, "error": "Oportunidad no encontrada."}), 404
    html = render_template("oportunidades/_detail.html", opportunity=detail)
    return jsonify({"ok": True, "id": opportunity_id, "html": html})


@bp.get("/api/<int:opportunity_id>/datos")
@_admin_required
def api_detail_data(opportunity_id: int):
    detail = get_opportunity_detail(opportunity_id)
    if not detail:
        return jsonify({"ok": False, "error": "Oportunidad no encontrada."}), 404
    return jsonify({"ok": True, "opportunity": detail})


@bp.post("/crear")
@_admin_required
def create():
    data = _opportunity_form()
    # Radicado y Estado dejaron de ser campos de usuario. Se conservan
    # internamente para compatibilidad con la BD, importaciones y relaciones.
    data["radicado"] = f"AUTO-{uuid.uuid4().hex[:12].upper()}"
    data["estado"] = STATUSES[0]
    try:
        _validate_opportunity(data)
        responsibles = _selected_users(request.form.getlist("responsables"))
        offer = request.files.get("oferta_inicial")
        image = request.files.get("imagen_inicial")
        if offer and offer.filename:
            _validate_file(offer, ATTACHMENT_OFFER)
        if image and image.filename:
            _validate_file(image, ATTACHMENT_IMAGE)
        opportunity_id = create_opportunity(data, responsibles, _user())
        subsystem_data = _subsystem_form("sub_")
        subsystem_id = create_subsystem(opportunity_id, subsystem_data, _user())
        warnings: list[str] = []
        if offer and offer.filename:
            try:
                _save_attachment(offer, ATTACHMENT_OFFER, opportunity_id, subsystem_id, request.form.get("oferta_titulo") or "Oferta inicial")
                audit_event("AGREGAR OFERTA", "OPORTUNIDADES DE PROYECTO", f"OD {opportunity_id} · {Path(offer.filename).name}", "OK", _user(), _client_ip())
            except Exception as exc:
                warnings.append(f"Oferta: {exc}")
        if image and image.filename:
            try:
                _save_attachment(image, ATTACHMENT_IMAGE, opportunity_id, subsystem_id, request.form.get("imagen_titulo") or "Imagen inicial")
                audit_event("AGREGAR IMAGEN", "OPORTUNIDADES DE PROYECTO", f"OD {opportunity_id} · {Path(image.filename).name}", "OK", _user(), _client_ip())
            except Exception as exc:
                warnings.append(f"Imagen: {exc}")
        audit_event("CREAR OPORTUNIDAD", "OPORTUNIDADES DE PROYECTO", f"Proyecto {data['proyecto']}", "OK", _user(), _client_ip())
        g.audit_logged = True
        if warnings:
            flash("La OD se creó, pero hubo archivos que no pudieron guardarse: " + " | ".join(warnings), "error")
        else:
            flash("Oportunidad creada correctamente.", "ok")
        return redirect(url_for("oportunidades.index", od=opportunity_id))
    except Exception as exc:
        audit_event("CREAR OPORTUNIDAD", "OPORTUNIDADES DE PROYECTO", str(exc), "ERROR", _user(), _client_ip())
        g.audit_logged = True
        flash(str(exc), "error")
        return redirect(url_for("oportunidades.index"))


@bp.post("/<int:opportunity_id>/editar")
@_admin_required
def edit(opportunity_id: int):
    existing = get_opportunity(opportunity_id)
    if not existing:
        abort(404)
    data = _opportunity_form()
    # Mantener los valores internos históricos aunque ya no sean editables/visibles.
    data["radicado"] = str(existing.get("radicado") or f"AUTO-{uuid.uuid4().hex[:12].upper()}")
    data["estado"] = str(existing.get("estado") or STATUSES[0])
    try:
        _validate_opportunity(data)
        update_opportunity_with_visible_technical(
            opportunity_id,
            data,
            _selected_users(request.form.getlist("responsables")),
            _visible_technical_updates_from_edit(),
            _user(),
        )
        audit_event("EDITAR OPORTUNIDAD", "OPORTUNIDADES DE PROYECTO", f"OD {opportunity_id} · Proyecto {data['proyecto']}", "OK", _user(), _client_ip())
        g.audit_logged = True
        flash("Oportunidad y datos técnicos actualizados.", "ok")
    except Exception as exc:
        flash(str(exc), "error")
    return redirect(url_for("oportunidades.index", od=opportunity_id))


@bp.post("/<int:opportunity_id>/eliminar")
@_admin_required
def delete(opportunity_id: int):
    opportunity = get_opportunity(opportunity_id)
    if not opportunity:
        abort(404)
    if soft_delete_opportunity(opportunity_id, _user()):
        audit_event("ELIMINAR OPORTUNIDAD", "OPORTUNIDADES DE PROYECTO", f"Proyecto {opportunity['proyecto']}", "OK", _user(), _client_ip())
        g.audit_logged = True
        flash("Oportunidad eliminada.", "ok")
    return redirect(url_for("oportunidades.index"))


@bp.post("/<int:opportunity_id>/subsistemas/crear")
@_admin_required
def create_subsystem_route(opportunity_id: int):
    if not get_opportunity(opportunity_id):
        abort(404)
    try:
        subsystem_id = create_subsystem(opportunity_id, _subsystem_form(), _user())
        audit_event("CREAR SUBSISTEMA", "OPORTUNIDADES DE PROYECTO", f"OD {opportunity_id} · Subsistema {subsystem_id}", "OK", _user(), _client_ip())
        g.audit_logged = True
        flash("Subsistema agregado.", "ok")
    except Exception as exc:
        flash(str(exc), "error")
    return redirect(url_for("oportunidades.index", od=opportunity_id))


@bp.post("/subsistemas/<int:subsystem_id>/editar")
@_admin_required
def edit_subsystem_route(subsystem_id: int):
    subsystem = get_subsystem(subsystem_id)
    if not subsystem:
        abort(404)
    try:
        opportunity_id = update_visible_subsystem(subsystem_id, _visible_technical_form())
        audit_event("EDITAR SUBSISTEMA", "OPORTUNIDADES DE PROYECTO", f"OD {opportunity_id} · {subsystem.get('nombre')}", "OK", _user(), _client_ip())
        g.audit_logged = True
        flash("Subsistema actualizado.", "ok")
        return redirect(url_for("oportunidades.index", od=opportunity_id))
    except Exception as exc:
        flash(str(exc), "error")
        return redirect(url_for("oportunidades.index", od=subsystem["opportunity_id"]))


@bp.post("/subsistemas/<int:subsystem_id>/eliminar")
@_admin_required
def delete_subsystem_route(subsystem_id: int):
    subsystem = get_subsystem(subsystem_id)
    if not subsystem:
        abort(404)
    opportunity_id = soft_delete_subsystem(subsystem_id)
    if opportunity_id:
        audit_event("ELIMINAR SUBSISTEMA", "OPORTUNIDADES DE PROYECTO", f"OD {opportunity_id} · {subsystem.get('nombre')}", "OK", _user(), _client_ip())
        g.audit_logged = True
        flash("Subsistema eliminado.", "ok")
    return redirect(url_for("oportunidades.index", od=opportunity_id or subsystem["opportunity_id"]))


@bp.post("/subsistemas/<int:subsystem_id>/ofertas/agregar")
@_admin_required
def add_offer(subsystem_id: int):
    subsystem = get_subsystem(subsystem_id)
    if not subsystem:
        abort(404)
    try:
        _save_attachment(request.files.get("archivo"), ATTACHMENT_OFFER, int(subsystem["opportunity_id"]), subsystem_id, request.form.get("titulo") or "")
        audit_event("AGREGAR OFERTA", "OPORTUNIDADES DE PROYECTO", f"OD {subsystem['opportunity_id']} · {subsystem['nombre']}", "OK", _user(), _client_ip())
        g.audit_logged = True
        flash("Oferta agregada.", "ok")
    except Exception as exc:
        flash(str(exc), "error")
    return redirect(url_for("oportunidades.index", od=subsystem["opportunity_id"]))


@bp.post("/subsistemas/<int:subsystem_id>/imagenes/agregar")
@_admin_required
def add_image(subsystem_id: int):
    subsystem = get_subsystem(subsystem_id)
    if not subsystem:
        abort(404)
    try:
        _save_attachment(request.files.get("archivo"), ATTACHMENT_IMAGE, int(subsystem["opportunity_id"]), subsystem_id, request.form.get("titulo") or "")
        audit_event("AGREGAR IMAGEN", "OPORTUNIDADES DE PROYECTO", f"OD {subsystem['opportunity_id']} · {subsystem['nombre']}", "OK", _user(), _client_ip())
        g.audit_logged = True
        flash("Imagen agregada.", "ok")
    except Exception as exc:
        flash(str(exc), "error")
    return redirect(url_for("oportunidades.index", od=subsystem["opportunity_id"]))


@bp.get("/archivos/<int:attachment_id>")
def attachment_view(attachment_id: int):
    item = get_attachment(attachment_id)
    if not item:
        abort(404)
    path = _safe_attachment_path(item["ruta_relativa"])
    if not path.exists():
        abort(404)
    return send_file(path, mimetype=item.get("mime_type") or None, conditional=True)


@bp.get("/archivos/<int:attachment_id>/descargar")
def attachment_download(attachment_id: int):
    item = get_attachment(attachment_id)
    if not item:
        abort(404)
    path = _safe_attachment_path(item["ruta_relativa"])
    if not path.exists():
        abort(404)
    return send_file(path, as_attachment=True, download_name=item.get("nombre_original") or path.name, conditional=True)


@bp.post("/archivos/<int:attachment_id>/eliminar")
@_admin_required
def attachment_delete(attachment_id: int):
    item = get_attachment(attachment_id)
    if not item:
        abort(404)
    data = soft_delete_attachment(attachment_id, _user())
    if data:
        path = _safe_attachment_path(data["ruta_relativa"])
        path.unlink(missing_ok=True)
        action = "ELIMINAR OFERTA" if data["tipo"] == ATTACHMENT_OFFER else "ELIMINAR IMAGEN"
        audit_event(action, "OPORTUNIDADES DE PROYECTO", f"OD {data['opportunity_id']} · {data['nombre_original']}", "OK", _user(), _client_ip())
        g.audit_logged = True
        flash("Archivo eliminado.", "ok")
    return redirect(url_for("oportunidades.index", od=item["opportunity_id"]))


@bp.post("/subsistemas/<int:subsystem_id>/puntos/<kind>/agregar")
@_admin_required
def add_point_route(subsystem_id: int, kind: str):
    subsystem = get_subsystem(subsystem_id)
    if not subsystem:
        abort(404)
    try:
        add_point(subsystem_id, kind, request.form.get("tipo") or "", request.form.get("cantidad") or "", request.form.get("restriccion_altura") or "")
        action = "AGREGAR PUNTO DE ENTRADA" if kind == "entrada" else "AGREGAR PUNTO DE SALIDA"
        audit_event(action, "OPORTUNIDADES DE PROYECTO", f"OD {subsystem['opportunity_id']} · {subsystem['nombre']}", "OK", _user(), _client_ip())
        g.audit_logged = True
        flash("Punto agregado.", "ok")
    except Exception as exc:
        flash(str(exc), "error")
    return redirect(url_for("oportunidades.index", od=subsystem["opportunity_id"]))


@bp.post("/subsistemas/<int:subsystem_id>/equipos/agregar")
@_admin_required
def add_equipment_route(subsystem_id: int):
    subsystem = get_subsystem(subsystem_id)
    if not subsystem:
        abort(404)
    try:
        add_equipment(subsystem_id, request.form.get("tipo_equipo") or "", request.form.get("referencia") or "", request.form.get("cantidad") or "")
        audit_event("AGREGAR EQUIPO", "OPORTUNIDADES DE PROYECTO", f"OD {subsystem['opportunity_id']} · {subsystem['nombre']}", "OK", _user(), _client_ip())
        g.audit_logged = True
        flash("Equipo agregado.", "ok")
    except Exception as exc:
        flash(str(exc), "error")
    return redirect(url_for("oportunidades.index", od=subsystem["opportunity_id"]))


@bp.post("/detalle/<table>/<int:row_id>/eliminar")
@_admin_required
def delete_child_route(table: str, row_id: int):
    try:
        opportunity_id = delete_child(table, row_id)
    except ValueError:
        abort(404)
    if opportunity_id:
        audit_event("ELIMINAR DETALLE", "OPORTUNIDADES DE PROYECTO", f"{table} #{row_id} · OD {opportunity_id}", "OK", _user(), _client_ip())
        g.audit_logged = True
        flash("Registro eliminado.", "ok")
        return redirect(url_for("oportunidades.index", od=opportunity_id))
    abort(404)
