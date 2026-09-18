"""Revisor de Planos — integrado como Blueprint dentro de la app GENERADOR RQ.

Todas las rutas quedan bajo el prefijo /planos (por ejemplo /planos/upload,
/planos/review, etc). El HTML vive en templates/planos.html y usa el mismo
sidebar y estilos (static/css/style.css) que el resto de la aplicacion.
"""

import os
import re
import json
import uuid
from datetime import datetime
from collections import defaultdict

from flask import Blueprint, render_template, request, jsonify, send_file
from werkzeug.utils import secure_filename

try:
    import fitz  # PyMuPDF
except Exception as e:
    fitz = None
    FITZ_ERROR = str(e)
else:
    FITZ_ERROR = None

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from storage import get_user_data_dir

planos_bp = Blueprint("planos", __name__, url_prefix="/planos")

PLANOS_DIR = os.path.dirname(os.path.abspath(__file__))


def current_planos_paths():
    root = get_user_data_dir() / "PLANOS"
    uploads = root / "CARGAS"
    reports = root / "REPORTES"
    uploads.mkdir(parents=True, exist_ok=True)
    reports.mkdir(parents=True, exist_ok=True)
    return {
        "root": str(root),
        "uploads": str(uploads),
        "reports": str(reports),
        "master": str(reports / "lista_planos_revisados.xlsx"),
        "state": str(root / "estado.json"),
    }

CODE_RE = re.compile(r"\b([0-9]{4}[A-Z0-9]{3,}(?:EN|LM|PF|PZ)?[A-Z0-9]{0,})\b")
PART_LINE_RE = re.compile(r"^\s*(\d+)\s+([A-Z0-9]{6,})\s+(.+?)\s+(\d+(?:[\.,]\d+)?)\s*$")
DIM_RE = re.compile(r"(\d+(?:[\.,]\d+)?)\s*mm\s*[xX]\s*(\d+(?:[\.,]\d+)?)\s*mm", re.I)
ESP_RE = re.compile(r"ESP\s*=\s*(CAL\s*\d+|\d+/\d+|\d+(?:[\.,]\d+)?\s*mm)", re.I)
LEN_RE = re.compile(r"(?:L\s*[=-]\s*|L=)(\d+(?:[\.,]\d+)?)\s*mm", re.I)
DIA_RE = re.compile(r"[Øø]\s*([0-9/\-\.]+)(?:\"|''|\s|$)")

STANDARD_WORDS = [
    "TUERCA", "TOR-", "TORNILLO", "ARAND", "WASSA", "ANILLO", "SEEGER",
    "RESORTE", "ACT LIN", "ACTUADOR", "ROTULA", "FERULA", "CODO", "EMP", "EMPAQUE",
    "ABRAZADERA", "RODAMIENTO", "CHUMACERA", "MOTOR", "VALVULA", "SENSOR"
]


def clean_number(value):
    if value is None:
        return None
    try:
        return float(str(value).replace(",", "."))
    except Exception:
        return None


def natural_key(value):
    text = str(value or "").upper()
    return [int(part) if part.isdigit() else part for part in re.split(r"(\d+)", text)]


def classify(code, desc):
    text = f"{code} {desc}".upper()
    # Regla solicitada: estandar = codigo compuesto SOLO por numeros.
    if code.isdigit():
        return "ESTANDAR"
    if "EN" in code:
        return "EN"
    if "LM" in code or text.startswith("LM ") or "LAMINA" in text:
        return "LM"
    if "PF" in code:
        return "PF"
    if "PZ" in code:
        return "PZ"
    return "FRAME"


def requires_drawing(tipo):
    return tipo in ["EN", "LM", "PF", "PZ", "FRAME"]


def extract_pdf_text(pdf_path):
    if fitz is None:
        raise RuntimeError("Falta PyMuPDF. Instale con: python -m pip install PyMuPDF")
    doc = fitz.open(pdf_path)
    pages = []
    for i, page in enumerate(doc, start=1):
        text = page.get_text("text") or ""
        pages.append({"page": i, "text": text})
    doc.close()
    return pages


def detect_sheet_codes(pages):
    """Detecta codigos que tienen plano en cada hoja.

    Puede haber una hoja con varias piezas detalladas. Se aceptan codigos que esten
    cerca de textos CANT/ESCALA/PESO; se excluyen codigos ubicados dentro de PARTS LIST.
    """
    sheet_map = {}
    for p in pages:
        page_no = p["page"]
        lines = [x.strip() for x in p["text"].splitlines() if x.strip()]

        # Rango de la PARTS LIST para evitar confundir referencias BOM con planos propios.
        parts_start = None
        for idx, line in enumerate(lines):
            if "PARTS LIST" in line.upper():
                parts_start = idx
                break

        def in_parts_list(idx):
            if parts_start is None:
                return False
            # En estos PDFs la lista termina cuando empiezan marcos/titulo/propiedad.
            tail_stop_words = ("ESTE PLANO", "MODELADO POR", "APROBADO", "HSEQ", "TODAS LAS DIMENSIONES")
            if idx < parts_start:
                return False
            for j in range(parts_start, idx + 1):
                if lines[j].upper().startswith(tail_stop_words):
                    return False
            return True

        for idx, line in enumerate(lines):
            if in_parts_list(idx):
                continue
            codes = CODE_RE.findall(line)
            if not codes:
                continue
            window = " ".join(lines[max(0, idx - 3): min(len(lines), idx + 6)]).upper()
            looks_like_title = ("CANT" in window and ("ESCALA" in window or "PESO" in window)) or idx < 6
            if not looks_like_title:
                continue
            for code in codes:
                if code.isdigit():
                    continue
                if code.upper() in ["CODIGO", "ITEM"]:
                    continue
                if any(tag in code for tag in ["EN", "LM", "PF", "PZ", "BOQ", "BQU"]):
                    sheet_map.setdefault(code, page_no)
    return sheet_map


def parse_parts(pages):
    """Lee todas las PARTS LIST de todas las hojas.

    PyMuPDF extrae las tablas de estos planos por columnas verticales, por ejemplo:
    CANT / DESCRIPCION / CODIGO / ITEM y luego: cantidad, descripcion, codigo, item.
    Por eso no basta leer una fila completa; se reconstruyen grupos de 4 lineas.
    """
    items = []
    seen = set()

    def is_int_line(x):
        return bool(re.fullmatch(r"\d+", x.strip()))

    def looks_like_code(x):
        x = x.strip()
        return bool(re.fullmatch(r"[A-Z0-9]{6,}", x))

    for p in pages:
        page_no = p["page"]
        lines = [x.strip() for x in p["text"].splitlines() if x.strip()]
        if not any("PARTS LIST" in x.upper() for x in lines):
            continue

        # Tomar solo lo que viene despues de PARTS LIST
        try:
            start_idx = next(i for i, x in enumerate(lines) if "PARTS LIST" in x.upper()) + 1
        except StopIteration:
            continue
        tail = lines[start_idx:]

        # Quitar encabezados y ruido simple
        skip_words = {"CANT", "DESCRIPCION", "DESCRIPCIÓN", "CODIGO", "CÓDIGO", "ITEM", "ITEM "}
        data = [x for x in tail if x.upper().strip() not in skip_words]

        i = 0
        while i <= len(data) - 4:
            # formato vertical: cantidad, descripcion, codigo, item
            if is_int_line(data[i]) and looks_like_code(data[i + 2]) and is_int_line(data[i + 3]):
                cant_raw = data[i]
                desc = data[i + 1]
                code = data[i + 2]
                item = data[i + 3]
                i += 4
            else:
                # respaldo para textos que si salgan como fila horizontal
                m = PART_LINE_RE.match(data[i])
                if m:
                    item, code, desc, cant_raw = m.groups()
                    i += 1
                else:
                    i += 1
                    continue

            if len(desc) < 3 or code.upper() in ["CODIGO", "ITEM"]:
                continue

            cant = clean_number(cant_raw) or 1
            key = (page_no, item, code, desc, cant)
            if key in seen:
                continue
            seen.add(key)

            tipo = classify(code, desc)
            dim = DIM_RE.search(desc)
            largo = ancho = area_unit = area_total = None
            if tipo == "LM" and dim:
                largo = clean_number(dim.group(1))
                ancho = clean_number(dim.group(2))
                if largo and ancho:
                    area_unit = largo * ancho / 1000000
                    area_total = area_unit * cant

            esp_match = ESP_RE.search(desc)
            esp = esp_match.group(1).replace('"', '').strip() if esp_match else ""
            if "CAL" in desc.upper() and not esp:
                cal = re.search(r"CAL\s*\d+", desc.upper())
                esp = cal.group(0) if cal else ""

            len_match = LEN_RE.search(desc)
            longitud = clean_number(len_match.group(1)) if len_match else None
            dia_match = DIA_RE.search(desc)
            diametro = dia_match.group(1) if dia_match else ""

            items.append({
                "hoja_bom": page_no,
                "item": item,
                "codigo": code,
                "descripcion": desc,
                "cantidad": cant,
                "tipo": tipo,
                "requiere_plano": requires_drawing(tipo),
                "tiene_plano": False,
                "hoja_plano": "",
                "observacion": "",
                "espesor": esp,
                "largo_mm": largo,
                "ancho_mm": ancho,
                "area_m2_unit": area_unit,
                "area_m2_total": area_total,
                "diametro": diametro,
                "longitud_mm": longitud,
            })
    return items


def deduplicate_items(items):
    grouped = {}
    for it in items:
        key = (it["codigo"], it["descripcion"], it["tipo"])
        if key not in grouped:
            grouped[key] = it.copy()
        else:
            grouped[key]["cantidad"] += it["cantidad"]
            if grouped[key].get("area_m2_unit") is not None:
                grouped[key]["area_m2_total"] = grouped[key]["area_m2_unit"] * grouped[key]["cantidad"]
            grouped[key]["hoja_bom"] = f"{grouped[key]['hoja_bom']}, {it['hoja_bom']}"
    return list(grouped.values())


def review_pdf(pdf_path):
    pages = extract_pdf_text(pdf_path)
    sheet_map = detect_sheet_codes(pages)
    items = deduplicate_items(parse_parts(pages))
    for it in items:
        hoja = sheet_map.get(it["codigo"])
        it["tiene_plano"] = bool(hoja)
        it["hoja_plano"] = hoja or ""
        if it["requiere_plano"] and not hoja:
            it["observacion"] = "FALTA PLANO"
        elif it["requiere_plano"] and hoja:
            it["observacion"] = "OK - plano hoja %s" % hoja
        else:
            it["observacion"] = "Estandar, no requiere plano"
    lamina_group = defaultdict(lambda: {"espesor": "", "area_m2_total": 0, "cantidad_piezas": 0})
    for it in items:
        if it["tipo"] == "LM":
            esp = it.get("espesor") or "SIN ESPESOR"
            lamina_group[esp]["espesor"] = esp
            lamina_group[esp]["area_m2_total"] += it.get("area_m2_total") or 0
            lamina_group[esp]["cantidad_piezas"] += it.get("cantidad") or 0
    lamina_resumen = sorted(lamina_group.values(), key=lambda x: natural_key(x.get("espesor", "")))
    items = sorted(items, key=lambda x: (x.get("tipo", ""), natural_key(x.get("codigo", "")), natural_key(x.get("descripcion", ""))))
    return {"items": items, "sheet_map": sheet_map, "lamina_resumen": lamina_resumen, "pages": len(pages)}


def style_ws(ws):
    header_fill = PatternFill("solid", fgColor="1266C3")
    green_fill = PatternFill("solid", fgColor="1FB64D")
    white_font = Font(color="FFFFFF", bold=True)
    bold = Font(bold=True)
    thin = Side(style="thin", color="DDDDDD")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows():
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = white_font
        cell.alignment = Alignment(horizontal="center")
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        max_len = 10
        for cell in ws[letter]:
            val = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, min(len(val) + 2, 45))
        ws.column_dimensions[letter].width = max_len
    ws.freeze_panes = "A2"


def export_excel(result, pdf_name):
    paths = current_planos_paths()
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe = re.sub(r"[^A-Za-z0-9_\-]+", "_", os.path.splitext(pdf_name)[0])[:60]
    path = os.path.join(paths["reports"], f"revision_{safe}_{ts}.xlsx")
    wb = Workbook()

    # Hoja principal solicitada: lista completa para compras/materiales.
    # Incluye ESTANDAR + LAMINAS + PERFILES PF + PIEZAS PZ.
    ws = wb.active
    ws.title = "Lista materiales"
    ws.append(["Grupo", "Codigo", "Descripcion", "Cantidad", "Espesor", "Diametro", "Longitud mm", "Largo mm", "Ancho mm", "Area total m2", "Tiene plano", "Hoja plano", "Observacion"])
    materiales = [it for it in result["items"] if it.get("tipo") in ("ESTANDAR", "LM", "PF", "PZ", "FRAME")]
    orden_grupo = {"LM": 1, "PF": 2, "PZ": 3, "FRAME": 4, "ESTANDAR": 5}
    materiales.sort(key=lambda x: (orden_grupo.get(x.get("tipo", ""), 99), natural_key(x.get("codigo", "")), natural_key(x.get("descripcion", ""))))
    for it in materiales:
        grupo = {"LM": "LAMINA", "PF": "PERFIL", "PZ": "PIEZA", "FRAME": "PIEZA FRAME", "ESTANDAR": "ESTANDAR"}.get(it.get("tipo"), it.get("tipo"))
        ws.append([
            grupo,
            it.get("codigo"),
            it.get("descripcion"),
            it.get("cantidad"),
            it.get("espesor"),
            it.get("diametro"),
            it.get("longitud_mm"),
            it.get("largo_mm"),
            it.get("ancho_mm"),
            it.get("area_m2_total"),
            "SI" if it.get("tiene_plano") else "NO",
            it.get("hoja_plano"),
            it.get("observacion"),
        ])
    style_ws(ws)

    ws = wb.create_sheet("Control de planos")
    headers = ["Codigo", "Descripcion", "Cantidad", "Tipo", "Requiere plano", "Tiene plano", "Hoja plano", "Hoja BOM", "Observacion", "Espesor", "Largo mm", "Ancho mm", "Area unit m2", "Area total m2", "Diametro", "Longitud mm"]
    ws.append(headers)
    for it in result["items"]:
        ws.append([
            it["codigo"], it["descripcion"], it["cantidad"], it["tipo"], "SI" if it["requiere_plano"] else "NO",
            "SI" if it["tiene_plano"] else "NO", it["hoja_plano"], it["hoja_bom"], it["observacion"],
            it.get("espesor"), it.get("largo_mm"), it.get("ancho_mm"), it.get("area_m2_unit"), it.get("area_m2_total"),
            it.get("diametro"), it.get("longitud_mm")
        ])
    style_ws(ws)
    ws2 = wb.create_sheet("Laminas agrupadas")
    ws2.append(["Espesor", "Cantidad piezas", "Area total m2"])
    for lg in result["lamina_resumen"]:
        ws2.append([lg["espesor"], lg["cantidad_piezas"], lg["area_m2_total"]])
    style_ws(ws2)
    ws3 = wb.create_sheet("Planos detectados")
    ws3.append(["Codigo", "Hoja"])
    for code, page in sorted(result["sheet_map"].items(), key=lambda x: (x[1], x[0])):
        ws3.append([code, page])
    style_ws(ws3)
    wb.save(path)
    update_master(pdf_name, path, result)
    return path


def update_master(pdf_name, report_path, result):
    master_xlsx = current_planos_paths()["master"]
    if os.path.exists(master_xlsx):
        wb = load_workbook(master_xlsx)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Planos revisados"
        ws.append(["Fecha", "Plano PDF", "Archivo Excel", "Total items", "Faltan planos", "Area laminas m2"])
    missing = sum(1 for it in result["items"] if it["requiere_plano"] and not it["tiene_plano"])
    area = sum((lg["area_m2_total"] or 0) for lg in result["lamina_resumen"])
    ws.append([datetime.now().strftime("%Y-%m-%d %H:%M:%S"), pdf_name, os.path.basename(report_path), len(result["items"]), missing, area])
    style_ws(ws)
    wb.save(master_xlsx)


def save_state(data):
    state_path = current_planos_paths()["state"]
    temporary = state_path + ".tmp"
    with open(temporary, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    os.replace(temporary, state_path)


def load_state():
    state_path = current_planos_paths()["state"]
    if not os.path.exists(state_path):
        return {}
    try:
        with open(state_path, "r", encoding="utf-8") as f:
            return json.load(f)
    except (OSError, json.JSONDecodeError):
        return {}


# =========================
# REVISION DISTRIBUCION DE LAMINA
# =========================
def norm_code(value):
    return str(value or "").strip().upper().replace(" ", "")


def norm_esp(value):
    text = str(value or "").strip().upper().replace('"', '').replace("''", "")
    text = re.sub(r"\s+", " ", text)
    return text


def read_material_excel(path):
    """Lee un Excel con columnas codigo, cantidad y espesor.
    Acepta encabezados como CODIGO/CÓDIGO, CANT/CANTIDAD y ESP/ESPESOR.
    """
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    header_row = None
    colmap = {}
    aliases = {
        "codigo": ["CODIGO", "CÓDIGO", "CODE", "REFERENCIA", "REF"],
        "cantidad": ["CANT", "CANTIDAD", "QTY", "QUANTITY"],
        "espesor": ["ESPESOR", "ESP", "CALIBRE", "THICKNESS"]
    }
    for r in range(1, min(ws.max_row, 12) + 1):
        vals = [str(ws.cell(r, c).value or "").strip().upper() for c in range(1, min(ws.max_column, 15) + 1)]
        found = {}
        for c, val in enumerate(vals, start=1):
            val_clean = val.replace(".", "").replace(":", "")
            for key, names in aliases.items():
                if val_clean in names or any(n in val_clean for n in names):
                    found[key] = c
        if "codigo" in found and ("cantidad" in found or "espesor" in found):
            header_row = r
            colmap = found
            break
    if header_row is None:
        header_row = 1
        colmap = {"codigo": 1, "cantidad": 2, "espesor": 3}

    data = {}
    for r in range(header_row + 1, ws.max_row + 1):
        code = norm_code(ws.cell(r, colmap.get("codigo", 1)).value)
        if not code or code in ["NONE", "CODIGO", "CÓDIGO"]:
            continue
        qty = clean_number(ws.cell(r, colmap.get("cantidad", 2)).value) or 0
        esp = norm_esp(ws.cell(r, colmap.get("espesor", 3)).value)
        if code not in data:
            data[code] = {"codigo": code, "cantidad": 0, "espesor": esp}
        data[code]["cantidad"] += qty
        if esp and not data[code].get("espesor"):
            data[code]["espesor"] = esp
        elif esp and data[code].get("espesor") and esp != data[code].get("espesor"):
            data[code]["espesor"] += " / " + esp
    return data


def detect_lamina_pdf_codes(pdf_path):
    pages = extract_pdf_text(pdf_path)
    page_map = {}
    all_text = "\n".join(p["text"] for p in pages)
    # Codigos con LM en el PDF; incluye planos de corte o referencias de lamina.
    candidates = set(re.findall(r"\b([A-Z0-9]{4,}LM[A-Z0-9]{0,})\b", all_text.upper()))
    for p in pages:
        txt = p["text"].upper()
        for code in candidates:
            if code in txt and code not in page_map:
                page_map[code] = p["page"]
    return page_map, len(pages)


def review_distribution(pdf_path, lm_path, dxf_path):
    lm = read_material_excel(lm_path) if lm_path and os.path.exists(lm_path) else {}
    dxf = read_material_excel(dxf_path) if dxf_path and os.path.exists(dxf_path) else {}
    pdf_map, pages_count = detect_lamina_pdf_codes(pdf_path) if pdf_path and os.path.exists(pdf_path) else ({}, 0)

    codes = set(lm) | set(dxf) | set(pdf_map)
    rows = []
    for code in sorted(codes, key=natural_key):
        a = lm.get(code, {})
        b = dxf.get(code, {})
        q_lm = a.get("cantidad", "")
        q_dxf = b.get("cantidad", "")
        e_lm = a.get("espesor", "")
        e_dxf = b.get("espesor", "")
        in_pdf = code in pdf_map
        diffs = []
        if code not in lm:
            diffs.append("No esta en lista LM")
        if code not in dxf:
            diffs.append("No esta en lista DXF")
        if not in_pdf:
            diffs.append("No se encontro plano en PDF")
        if code in lm and code in dxf:
            if (q_lm or 0) != (q_dxf or 0):
                diffs.append("Diferencia cantidad")
            if norm_esp(e_lm) != norm_esp(e_dxf):
                diffs.append("Diferencia espesor")
        rows.append({
            "codigo": code,
            "tiene_plano": in_pdf,
            "hoja": pdf_map.get(code, ""),
            "cant_lm": q_lm,
            "cant_dxf": q_dxf,
            "espesor_lm": e_lm,
            "espesor_dxf": e_dxf,
            "diferencia": "; ".join(diffs),
            "ok": not diffs,
        })
    return {"rows": rows, "pages": pages_count, "lm_count": len(lm), "dxf_count": len(dxf)}


def save_upload_file(file_obj, prefix, allowed_ext):
    if not file_obj or not file_obj.filename:
        raise ValueError("No se recibio archivo")
    ext = os.path.splitext(file_obj.filename)[1].lower()
    if ext not in allowed_ext:
        raise ValueError("Extension no permitida")
    filename = secure_filename(file_obj.filename)
    path = os.path.join(current_planos_paths()["uploads"], f"{prefix}_{uuid.uuid4().hex}_{filename}")
    file_obj.save(path)
    return path, filename


# =========================
# RUTAS (todas bajo /planos)
# =========================

@planos_bp.route("/")
def planos_home():
    return render_template("planos.html")


@planos_bp.route("/upload", methods=["POST"])
def upload():
    if "pdf" not in request.files:
        return jsonify({"ok": False, "error": "No se recibio PDF"}), 400
    f = request.files["pdf"]
    if not f.filename.lower().endswith(".pdf"):
        return jsonify({"ok": False, "error": "Debe cargar un archivo PDF"}), 400
    filename = secure_filename(f.filename)
    unique = f"{uuid.uuid4().hex}_{filename}"
    path = os.path.join(current_planos_paths()["uploads"], unique)
    f.save(path)
    save_state({"pdf_path": path, "pdf_name": filename})
    return jsonify({"ok": True, "pdf_name": filename})


@planos_bp.route("/review", methods=["POST"])
def review():
    state = load_state()
    pdf_path = state.get("pdf_path")
    pdf_name = state.get("pdf_name")
    if not pdf_path or not os.path.exists(pdf_path):
        return jsonify({"ok": False, "error": "Primero cargue un plano PDF"}), 400
    try:
        result = review_pdf(pdf_path)
        excel_path = export_excel(result, pdf_name)
        result["excel_file"] = os.path.basename(excel_path)
        result["pdf_name"] = pdf_name
        state.update({"last_result": result, "last_excel": excel_path})
        save_state(state)
        return jsonify({"ok": True, "result": result})
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 500


@planos_bp.route("/clear", methods=["POST"])
def clear():
    save_state({})
    return jsonify({"ok": True})


@planos_bp.route("/download/current")
def download_current():
    state = load_state()
    path = state.get("last_excel")
    if not path or not os.path.exists(path):
        return jsonify({"ok": False, "error": "No hay reporte actual"}), 404
    return send_file(path, as_attachment=True)


@planos_bp.route("/download/master")
def download_master():
    master_xlsx = current_planos_paths()["master"]
    if not os.path.exists(master_xlsx):
        wb = Workbook()
        ws = wb.active
        ws.title = "Planos revisados"
        ws.append(["Fecha", "Plano PDF", "Archivo Excel", "Total items", "Faltan planos", "Area laminas m2"])
        style_ws(ws)
        wb.save(master_xlsx)
    return send_file(master_xlsx, as_attachment=True)


@planos_bp.route("/history")
def history():
    paths = current_planos_paths()
    review_dir = paths["reports"]
    master_xlsx = paths["master"]
    files = []
    for name in os.listdir(review_dir):
        if name.lower().endswith(".xlsx") and name != os.path.basename(master_xlsx):
            path = os.path.join(review_dir, name)
            files.append({"name": name, "size": os.path.getsize(path), "mtime": datetime.fromtimestamp(os.path.getmtime(path)).strftime("%Y-%m-%d %H:%M"), "mtime_raw": os.path.getmtime(path)})
    files.sort(key=lambda x: x["mtime_raw"], reverse=True)
    for f in files:
        f.pop("mtime_raw", None)
    return jsonify({"ok": True, "files": files})


@planos_bp.route("/download/report/<name>")
def download_report(name):
    name = secure_filename(name)
    path = os.path.join(current_planos_paths()["reports"], name)
    if not os.path.exists(path):
        return jsonify({"ok": False, "error": "No existe el archivo"}), 404
    return send_file(path, as_attachment=True)


@planos_bp.route("/delete/report/<name>", methods=["POST"])
def delete_report(name):
    name = secure_filename(name)
    path = os.path.join(current_planos_paths()["reports"], name)
    if not os.path.exists(path):
        return jsonify({"ok": False, "error": "No existe el archivo"}), 404
    os.remove(path)
    # Si era el reporte actual, limpiar esa referencia.
    state = load_state()
    if state.get("last_excel") == path:
        state.pop("last_excel", None)
        state.pop("last_result", None)
        save_state(state)
    return jsonify({"ok": True})


@planos_bp.route("/dist/upload_pdf", methods=["POST"])
def dist_upload_pdf():
    try:
        path, name = save_upload_file(request.files.get("pdf"), "dist_pdf", [".pdf"])
        state = load_state()
        dist = state.get("distribution", {})
        dist.update({"pdf_path": path, "pdf_name": name})
        state["distribution"] = dist
        save_state(state)
        return jsonify({"ok": True, "name": name})
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 400


@planos_bp.route("/dist/upload_lm", methods=["POST"])
def dist_upload_lm():
    try:
        path, name = save_upload_file(request.files.get("excel"), "dist_lm", [".xlsx", ".xlsm"])
        state = load_state()
        dist = state.get("distribution", {})
        dist.update({"lm_path": path, "lm_name": name})
        state["distribution"] = dist
        save_state(state)
        preview = list(read_material_excel(path).values())
        return jsonify({"ok": True, "name": name, "rows": preview})
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 400


@planos_bp.route("/dist/upload_dxf", methods=["POST"])
def dist_upload_dxf():
    try:
        path, name = save_upload_file(request.files.get("excel"), "dist_dxf", [".xlsx", ".xlsm"])
        state = load_state()
        dist = state.get("distribution", {})
        dist.update({"dxf_path": path, "dxf_name": name})
        state["distribution"] = dist
        save_state(state)
        preview = list(read_material_excel(path).values())
        return jsonify({"ok": True, "name": name, "rows": preview})
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 400


@planos_bp.route("/dist/review", methods=["POST"])
def dist_review():
    state = load_state()
    dist = state.get("distribution", {})
    pdf_path = dist.get("pdf_path")
    lm_path = dist.get("lm_path")
    dxf_path = dist.get("dxf_path")
    if not pdf_path or not os.path.exists(pdf_path):
        return jsonify({"ok": False, "error": "Cargue primero el PDF de distribucion de lamina"}), 400
    if not lm_path or not os.path.exists(lm_path):
        return jsonify({"ok": False, "error": "Cargue primero la lista Excel LM"}), 400
    if not dxf_path or not os.path.exists(dxf_path):
        return jsonify({"ok": False, "error": "Cargue primero la lista Excel DXF"}), 400
    try:
        result = review_distribution(pdf_path, lm_path, dxf_path)
        dist["last_result"] = result
        state["distribution"] = dist
        save_state(state)
        return jsonify({"ok": True, "result": result})
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 500


@planos_bp.route("/dist/clear", methods=["POST"])
def dist_clear():
    state = load_state()
    state.pop("distribution", None)
    save_state(state)
    return jsonify({"ok": True})
