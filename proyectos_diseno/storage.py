from __future__ import annotations

import os
import shutil
import tempfile
import threading
from copy import deepcopy
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


DATASETS: Dict[str, Dict[str, Any]] = {
    "designers": {
        "filename": "Disenadores.xlsx",
        "sheet": "DISENADORES",
        "id_field": "id",
        "columns": [
            "id", "nombre", "color", "costo_mensual_empresa", "hora_entrada",
            "hora_salida", "almuerzo_inicio", "almuerzo_fin", "activo",
        ],
    },
    "pmps": {
        "filename": "PMP.xlsx",
        "sheet": "PMP",
        "id_field": "id",
        "columns": ["id", "nombre", "activo"],
    },
    "activities": {
        "filename": "Actividades.xlsx",
        "sheet": "ACTIVIDADES",
        "id_field": "id",
        "columns": ["id", "nombre", "etapa", "porcentaje", "tipos_proyecto", "activa"],
    },
    "equipment": {
        "filename": "Equipos.xlsx",
        "sheet": "EQUIPOS",
        "id_field": "id",
        "columns": [
            "id", "codigo", "nombre", "dias_estandar", "dias_medio", "dias_no_estandar", "dias_complejo", "activo",
        ],
    },
    "project_sizes": {
        "filename": "Tamanos_Proyectos.xlsx",
        "sheet": "TAMANOS_PROYECTOS",
        "id_field": "id",
        "columns": ["id", "nombre", "minimo", "maximo", "activo"],
    },
    "project_revisions": {
        "filename": "Revisiones_Proyectos.xlsx",
        "sheet": "REVISIONES_PROYECTOS",
        "id_field": "id",
        "columns": ["id", "project_id", "numero_revision", "fecha", "creado_en", "actualizado_en"],
    },
    "holidays": {
        "filename": "Festivos.xlsx",
        "sheet": "FESTIVOS",
        "id_field": "id",
        "columns": ["id", "fecha", "nombre", "origen", "activo"],
    },
    "projects": {
        "filename": "Proyectos.xlsx",
        "sheet": "PROYECTOS",
        "id_field": "id",
        "columns": [
            "id", "numero", "cliente", "tipo_proyecto", "designer_id", "pmp_id", "pmp_nombre", "bodega",
            "tamano_id", "tamano_nombre", "fecha_contra_actual", "imagen_path",
            "fecha_recepcion_alcance", "fecha_inicio",
            "etapa1_inicio", "etapa1_fin", "fecha_aprobacion",
            "etapa2_inicio", "etapa2_fin", "etapa3_inicio", "etapa3_fin",
            "dias_referencia", "dias_aprobacion", "retraso_etapa1", "retraso_etapa2",
            "retraso_etapa3", "notas_historico", "estatus_historico", "estado", "fecha_finalizacion",
            "creado_en", "actualizado_en",
        ],
    },
    "stage_periods": {
        "filename": "Periodos_Etapas.xlsx",
        "sheet": "PERIODOS_ETAPAS",
        "id_field": "id",
        "columns": ["id", "project_id", "etapa", "inicio", "fin", "orden"],
    },
    "upcoming_projects": {
        "filename": "Proyectos_Por_Empezar.xlsx",
        "sheet": "PROYECTOS_POR_EMPEZAR",
        "id_field": "id",
        "columns": [
            "id", "numero", "cliente", "designer_id", "fecha_inicio",
            "dias_referencia", "creado_en", "actualizado_en",
        ],
    },
    "additionals": {
        "filename": "Adicionales.xlsx",
        "sheet": "ADICIONALES",
        "id_field": "id",
        "columns": [
            "id", "titulo", "proyecto", "solicita", "descripcion", "fecha_inicio",
            "fecha_finalizacion", "dias_duracion", "creado_en", "actualizado_en",
        ],
    },
    "project_equipment": {
        "filename": "Equipos_Proyectos.xlsx",
        "sheet": "EQUIPOS_PROYECTOS",
        "id_field": "id",
        "columns": [
            "id", "project_id", "equipment_id", "codigo", "nombre", "cantidad",
            "tipo", "dias_aplicados",
        ],
    },
    "project_progress": {
        "filename": "Avance_Proyectos.xlsx",
        "sheet": "AVANCE_PROYECTOS",
        "id_field": "id",
        "columns": [
            "id", "project_id", "activity_id", "actividad", "etapa", "porcentaje",
            "cumplida", "fecha_cumplimiento",
        ],
    },
}

BOOL_FIELDS = {"activo", "activa", "cumplida"}
NUMERIC_FIELDS = {
    "costo_mensual_empresa", "porcentaje", "dias_estandar", "dias_medio", "dias_no_estandar", "dias_complejo",
    "minimo", "maximo",
    "dias_referencia", "dias_aprobacion", "retraso_etapa1", "retraso_etapa2", "retraso_etapa3",
    "cantidad", "dias_aplicados", "dias_duracion", "etapa", "orden",
}
DATE_FIELDS = {
    "fecha", "fecha_contra_actual", "fecha_recepcion_alcance", "fecha_inicio", "etapa1_inicio", "etapa1_fin", "fecha_aprobacion",
    "etapa2_inicio", "etapa2_fin", "etapa3_inicio", "etapa3_fin", "fecha_finalizacion",
    "fecha_cumplimiento", "creado_en", "actualizado_en", "inicio", "fin",
}
PROJECT_TYPES = {"T1", "T2", "T3", "T4"}


class StorageError(RuntimeError):
    pass


class ExcelStorage:
    """Almacenamiento sencillo: los archivos Excel son la única base de datos."""

    def __init__(self, base_dir: Path, backup_dir: Path):
        self.base_dir = Path(base_dir)
        self.backup_dir = Path(backup_dir)
        self._lock = threading.RLock()
        self._cache: Dict[str, List[Dict[str, Any]]] = {}
        self.base_dir.mkdir(parents=True, exist_ok=True)
        self.backup_dir.mkdir(parents=True, exist_ok=True)
        self.ensure_all_workbooks()
        self.migrate_workbook_schemas()
        self.reload_from_excel(validate=False)
        self.migrate_rev11_defaults()
        self.migrate_rev17_defaults()
        self.migrate_rev19_defaults()
        self.migrate_rev20_defaults()
        self.reload_from_excel(validate=True)

    def path_for(self, dataset: str) -> Path:
        if dataset not in DATASETS:
            raise StorageError(f"Conjunto desconocido: {dataset}")
        return self.base_dir / DATASETS[dataset]["filename"]

    @staticmethod
    def _normalize_value(field: str, value: Any) -> Any:
        if value is None:
            return ""
        if isinstance(value, datetime):
            if field in DATE_FIELDS:
                return value.strftime("%Y-%m-%d")
            return value.isoformat(timespec="seconds")
        if isinstance(value, date):
            return value.isoformat()
        if field in BOOL_FIELDS:
            if isinstance(value, str):
                return value.strip().lower() in {"1", "true", "si", "sí", "x", "yes"}
            return bool(value)
        if field in NUMERIC_FIELDS:
            if value == "":
                return 0
            try:
                numeric = float(value)
                return int(numeric) if numeric.is_integer() else numeric
            except (TypeError, ValueError):
                return value
        return str(value).strip() if not isinstance(value, (dict, list)) else value

    def ensure_all_workbooks(self) -> None:
        with self._lock:
            for dataset in DATASETS:
                if not self.path_for(dataset).exists():
                    self._write_rows(dataset, [])

    def migrate_workbook_schemas(self) -> None:
        """Agrega columnas nuevas a los Excel existentes sin perder registros."""
        with self._lock:
            for dataset, cfg in DATASETS.items():
                path = self.path_for(dataset)
                try:
                    wb = load_workbook(path, data_only=False, read_only=False)
                except Exception as exc:
                    raise StorageError(f"No se pudo revisar {path.name}: {exc}") from exc
                if cfg["sheet"] not in wb.sheetnames:
                    wb.close()
                    raise StorageError(f"{path.name} no contiene la hoja {cfg['sheet']}")
                ws = wb[cfg["sheet"]]
                headers = [str(c.value or "").strip() for c in ws[1]]
                missing = [column for column in cfg["columns"] if column not in headers]
                if not missing:
                    wb.close()
                    continue
                existing_rows = []
                header_index = {name: index for index, name in enumerate(headers)}
                for values in ws.iter_rows(min_row=2, values_only=True):
                    if not any(value not in (None, "") for value in values):
                        continue
                    row = {}
                    for column in cfg["columns"]:
                        index = header_index.get(column)
                        value = values[index] if index is not None and index < len(values) else ""
                        row[column] = self._normalize_value(column, value)
                    existing_rows.append(row)
                wb.close()
                self._write_rows(dataset, existing_rows)

    def migrate_rev11_defaults(self) -> None:
        """Conserva los datos anteriores y crea la estructura nueva de Rev11."""
        activities = self.read("activities")
        activities_changed = False
        for row in activities:
            if not str(row.get("tipos_proyecto", "")).strip():
                row["tipos_proyecto"] = "T1,T2,T3,T4"
                activities_changed = True
        if activities_changed:
            self.write("activities", activities)

        projects = self.read("projects")
        projects_changed = False
        for row in projects:
            if str(row.get("tipo_proyecto", "")).strip().upper() not in PROJECT_TYPES:
                row["tipo_proyecto"] = "T1"
                projects_changed = True
        if projects_changed:
            self.write("projects", projects)

        periods = self.read("stage_periods")
        period_project_ids = {str(r.get("project_id", "")) for r in periods}
        additions: List[Dict[str, Any]] = []
        import uuid
        for project in projects:
            project_id = str(project.get("id", ""))
            if not project_id or project_id in period_project_ids:
                continue
            for stage in (1, 2, 3):
                start = str(project.get(f"etapa{stage}_inicio", "")).strip()
                end = str(project.get(f"etapa{stage}_fin", "")).strip()
                if start and end:
                    additions.append({
                        "id": str(uuid.uuid4()), "project_id": project_id, "etapa": stage,
                        "inicio": start[:10], "fin": end[:10], "orden": 1,
                    })
        if additions:
            periods.extend(additions)
            self.write("stage_periods", periods)

    def migrate_rev17_defaults(self) -> None:
        """Inicializa el ESTATUS de históricos anteriores sin perder datos."""
        projects = self.read("projects")
        changed = False
        for row in projects:
            if str(row.get("estado", "")).strip().lower() == "finalizado" and not str(row.get("estatus_historico", "")).strip():
                row["estatus_historico"] = "PRODUCCION"
                changed = True
        if changed:
            self.write("projects", projects)


    def migrate_rev19_defaults(self) -> None:
        """Agrega valores iniciales de v18.1 sin modificar información existente."""
        equipment = self.read("equipment")
        changed = False
        for row in equipment:
            if row.get("dias_medio", "") in ("", None):
                row["dias_medio"] = 0
                changed = True
            if row.get("dias_complejo", "") in ("", None):
                row["dias_complejo"] = 0
                changed = True
        if changed:
            self.write("equipment", equipment)

        sizes = self.read("project_sizes")
        if not sizes:
            import uuid
            defaults = [
                ("BAJO", 1, 3),
                ("MEDIO BAJO", 4, 10),
                ("MEDIO", 11, 15),
                ("MEDIO ALTO", 16, 20),
                ("ALTO", 21, 100),
            ]
            self.write("project_sizes", [
                {"id": str(uuid.uuid4()), "nombre": name, "minimo": low, "maximo": high, "activo": True}
                for name, low, high in defaults
            ])

    def migrate_rev20_defaults(self) -> None:
        """Convierte la aprobación legada de v18.1 en un lapso explícito de etapa 0."""
        periods=self.read("stage_periods"); projects=self.read("projects"); import uuid
        changed=False
        from datetime import date as _date, timedelta as _timedelta
        for project in projects:
            pid=str(project.get("id", "")); pperiods=[r for r in periods if str(r.get("project_id", ""))==pid]
            if not pid or any(int(float(r.get("etapa") or 0))==0 for r in pperiods): continue
            approval=str(project.get("fecha_aprobacion", "")).strip()[:10]
            s1=sorted([r for r in pperiods if int(float(r.get("etapa") or 0))==1],key=lambda r:str(r.get("fin", "")))
            if not approval or not s1: continue
            try:
                start=_date.fromisoformat(str(s1[-1].get("fin", ""))[:10])+_timedelta(days=1); end=_date.fromisoformat(approval)
            except Exception: continue
            if end < start: continue
            periods.append({"id":str(uuid.uuid4()),"project_id":pid,"etapa":0,"inicio":start.isoformat(),"fin":end.isoformat(),"orden":1}); changed=True
        if changed: self.write("stage_periods", periods)

    def _format_workbook(self, ws, columns: List[str], row_count: int) -> None:
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(1, row_count + 1)}"
        header_fill = "0B5DAA"
        thin_gray = Side(style="thin", color="D9E1F2")
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor=header_fill)
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(bottom=thin_gray)
        ws.row_dimensions[1].height = 25
        widths = {
            "id": 38, "nombre": 32, "actividad": 36, "cliente": 30, "pmp_nombre": 30, "bodega": 16,
            "tamano_id": 38, "tamano_nombre": 20, "fecha_contra_actual": 18, "imagen_path": 48,
            "numero_revision": 18, "minimo": 12, "maximo": 12,
            "titulo": 30, "proyecto": 20, "solicita": 24, "descripcion": 48, "estatus_historico": 22,
            "numero": 16, "codigo": 18, "color": 13, "etapa": 15,
            "porcentaje": 13, "tipos_proyecto": 24, "tipo_proyecto": 20,
            "fecha": 14, "fecha_recepcion_alcance": 22, "fecha_inicio": 14,
            "etapa1_inicio": 14, "etapa1_fin": 14, "fecha_aprobacion": 16,
            "dias_aprobacion": 17, "retraso_etapa1": 17, "retraso_etapa2": 17,
            "retraso_etapa3": 17, "notas_historico": 50, "dias_duracion": 16, "pmp_id": 38,
            "etapa2_inicio": 14, "etapa2_fin": 14, "etapa3_inicio": 14,
            "etapa3_fin": 14, "fecha_finalizacion": 17, "creado_en": 20,
            "actualizado_en": 20, "hora_entrada": 14, "hora_salida": 14,
            "almuerzo_inicio": 16, "almuerzo_fin": 16,
            "costo_mensual_empresa": 23, "dias_estandar": 16, "dias_medio": 15,
            "dias_no_estandar": 19, "dias_complejo": 16, "dias_referencia": 18,
            "project_id": 38, "designer_id": 38, "equipment_id": 38,
            "activity_id": 38, "fecha_cumplimiento": 20,
            "inicio": 14, "fin": 14, "orden": 10,
        }
        for idx, col in enumerate(columns, 1):
            ws.column_dimensions[get_column_letter(idx)].width = widths.get(col, 18)
        for row in ws.iter_rows(min_row=2, max_row=max(2, row_count + 1)):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        if row_count > 0:
            ref = f"A1:{get_column_letter(len(columns))}{row_count + 1}"
            table_name = "TB_" + "".join(ch for ch in ws.title[:20] if ch.isalnum() or ch == "_")
            table = Table(displayName=table_name, ref=ref)
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
                showRowStripes=True, showColumnStripes=False,
            )
            ws.add_table(table)

    def _write_rows(self, dataset: str, rows: Iterable[Dict[str, Any]]) -> None:
        cfg = DATASETS[dataset]
        columns = cfg["columns"]
        normalized_rows = [deepcopy(r) for r in rows]
        path = self.path_for(dataset)
        wb = Workbook()
        ws = wb.active
        ws.title = cfg["sheet"]
        ws.append(columns)
        for raw in normalized_rows:
            ws.append([self._normalize_value(col, raw.get(col, "")) for col in columns])
        self._format_workbook(ws, columns, len(normalized_rows))
        with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx", dir=path.parent) as tmp:
            temp_path = Path(tmp.name)
        try:
            wb.save(temp_path)
            os.replace(temp_path, path)
        finally:
            temp_path.unlink(missing_ok=True)

    def _read_disk(self, dataset: str) -> List[Dict[str, Any]]:
        cfg = DATASETS[dataset]
        path = self.path_for(dataset)
        try:
            wb = load_workbook(path, data_only=True, read_only=True)
        except Exception as exc:
            raise StorageError(f"No se pudo abrir {path.name}: {exc}") from exc
        if cfg["sheet"] not in wb.sheetnames:
            wb.close()
            raise StorageError(f"{path.name} no contiene la hoja {cfg['sheet']}")
        ws = wb[cfg["sheet"]]
        headers = [str(c.value or "").strip() for c in ws[1]]
        missing = [c for c in cfg["columns"] if c not in headers]
        if missing:
            wb.close()
            raise StorageError(f"{path.name}: faltan columnas {', '.join(missing)}")
        index = {name: headers.index(name) for name in cfg["columns"]}
        rows: List[Dict[str, Any]] = []
        for values in ws.iter_rows(min_row=2, values_only=True):
            if not any(v not in (None, "") for v in values):
                continue
            row: Dict[str, Any] = {}
            for field in cfg["columns"]:
                pos = index[field]
                value = values[pos] if pos < len(values) else ""
                row[field] = self._normalize_value(field, value)
            rows.append(row)
        wb.close()
        return rows

    def read(self, dataset: str, from_disk: bool = False) -> List[Dict[str, Any]]:
        with self._lock:
            if from_disk:
                return self._read_disk(dataset)
            if dataset not in self._cache:
                self._cache[dataset] = self._read_disk(dataset)
            return deepcopy(self._cache[dataset])

    def reload_from_excel(self, validate: bool = True) -> Dict[str, int]:
        with self._lock:
            disk_data = {dataset: self._read_disk(dataset) for dataset in DATASETS}
            if validate:
                self._validate_data(disk_data)
            self._cache = deepcopy(disk_data)
            return {dataset: len(rows) for dataset, rows in disk_data.items()}

    def write(self, dataset: str, rows: Iterable[Dict[str, Any]]) -> None:
        with self._lock:
            normalized = [deepcopy(r) for r in rows]
            self._write_rows(dataset, normalized)
            self._cache[dataset] = normalized

    def get(self, dataset: str, row_id: str) -> Dict[str, Any] | None:
        id_field = DATASETS[dataset]["id_field"]
        return next((r for r in self.read(dataset) if str(r.get(id_field)) == str(row_id)), None)

    def upsert(self, dataset: str, row: Dict[str, Any]) -> Dict[str, Any]:
        id_field = DATASETS[dataset]["id_field"]
        row_id = str(row.get(id_field, "")).strip()
        if not row_id:
            raise StorageError(f"Falta {id_field} para guardar en {dataset}")
        with self._lock:
            rows = self.read(dataset)
            for idx, current in enumerate(rows):
                if str(current.get(id_field)) == row_id:
                    merged = dict(current)
                    merged.update(row)
                    rows[idx] = merged
                    self.write(dataset, rows)
                    return merged
            rows.append(row)
            self.write(dataset, rows)
            return row

    def delete(self, dataset: str, row_id: str) -> bool:
        id_field = DATASETS[dataset]["id_field"]
        with self._lock:
            rows = self.read(dataset)
            new_rows = [r for r in rows if str(r.get(id_field)) != str(row_id)]
            if len(new_rows) == len(rows):
                return False
            self.write(dataset, new_rows)
            return True

    def replace_where(self, dataset: str, predicate, replacement: Iterable[Dict[str, Any]]) -> None:
        with self._lock:
            rows = [r for r in self.read(dataset) if not predicate(r)]
            rows.extend(list(replacement))
            self.write(dataset, rows)

    @staticmethod
    def _activity_types(value: Any) -> set[str]:
        if isinstance(value, list):
            parts = value
        else:
            parts = str(value or "").replace(";", ",").split(",")
        return {str(item).strip().upper() for item in parts if str(item).strip().upper() in PROJECT_TYPES}

    def _validate_data(self, all_data: Dict[str, List[Dict[str, Any]]]) -> Dict[str, int]:
        counts: Dict[str, int] = {}
        seen_project_numbers: set[str] = set()
        seen_upcoming_numbers: set[str] = set()
        project_ids = {str(r.get("id", "")) for r in all_data.get("projects", [])}
        for dataset, cfg in DATASETS.items():
            rows = all_data[dataset]
            counts[dataset] = len(rows)
            id_field = cfg["id_field"]
            ids: set[str] = set()
            for n, row in enumerate(rows, 2):
                rid = str(row.get(id_field, "")).strip()
                if not rid:
                    raise StorageError(f"{cfg['filename']} fila {n}: falta {id_field}")
                if rid in ids:
                    raise StorageError(f"{cfg['filename']}: ID duplicado {rid}")
                ids.add(rid)
                if dataset == "projects":
                    number = str(row.get("numero", "")).strip().upper()
                    if not number:
                        raise StorageError(f"Proyectos.xlsx fila {n}: falta numero")
                    if number in seen_project_numbers:
                        raise StorageError(f"Proyectos.xlsx: proyecto duplicado {number}")
                    seen_project_numbers.add(number)
                    if str(row.get("tipo_proyecto", "")).strip().upper() not in PROJECT_TYPES:
                        raise StorageError(f"Proyectos.xlsx fila {n}: tipo_proyecto debe ser T1, T2, T3 o T4")
                if dataset == "upcoming_projects":
                    number = str(row.get("numero", "")).strip().upper()
                    if not number:
                        raise StorageError(f"Proyectos_Por_Empezar.xlsx fila {n}: falta numero")
                    if number in seen_upcoming_numbers:
                        raise StorageError(f"Proyectos_Por_Empezar.xlsx: proyecto duplicado {number}")
                    seen_upcoming_numbers.add(number)
                if dataset == "activities":
                    stage = int(float(row.get("etapa") or 0))
                    if stage not in {1, 2, 3}:
                        raise StorageError(f"Actividades.xlsx fila {n}: etapa debe ser 1, 2 o 3")
                    pct = float(row.get("porcentaje") or 0)
                    if pct <= 0 or pct > 100:
                        raise StorageError(f"Actividades.xlsx fila {n}: porcentaje inválido")
                    if not self._activity_types(row.get("tipos_proyecto")):
                        raise StorageError(f"Actividades.xlsx fila {n}: debe asignar al menos un tipo de proyecto")
                if dataset == "stage_periods":
                    if str(row.get("project_id", "")) not in project_ids:
                        raise StorageError(f"Periodos_Etapas.xlsx fila {n}: project_id no existe")
                    stage = int(float(row.get("etapa") or 0))
                    if stage not in {0, 1, 2, 3}:
                        raise StorageError(f"Periodos_Etapas.xlsx fila {n}: etapa debe ser 0 (Aprobación), 1, 2 o 3")
                    try:
                        start = date.fromisoformat(str(row.get("inicio", ""))[:10])
                        end = date.fromisoformat(str(row.get("fin", ""))[:10])
                    except ValueError as exc:
                        raise StorageError(f"Periodos_Etapas.xlsx fila {n}: fecha inválida") from exc
                    if end < start:
                        raise StorageError(f"Periodos_Etapas.xlsx fila {n}: fin anterior al inicio")
                if dataset == "additionals":
                    if not str(row.get("titulo", "")).strip() or not str(row.get("proyecto", "")).strip():
                        raise StorageError(f"Adicionales.xlsx fila {n}: título y proyecto son obligatorios")
                    try:
                        start = date.fromisoformat(str(row.get("fecha_inicio", ""))[:10])
                        end = date.fromisoformat(str(row.get("fecha_finalizacion", ""))[:10])
                    except ValueError as exc:
                        raise StorageError(f"Adicionales.xlsx fila {n}: fecha inválida") from exc
                    if end < start:
                        raise StorageError(f"Adicionales.xlsx fila {n}: finalización anterior al inicio")
                if dataset == "holidays" and row.get("fecha"):
                    try:
                        date.fromisoformat(str(row["fecha"])[:10])
                    except ValueError as exc:
                        raise StorageError(f"Festivos.xlsx fila {n}: fecha inválida") from exc

        return counts

    def validate_all(self, from_disk: bool = False) -> Dict[str, int]:
        data = {dataset: self.read(dataset, from_disk=from_disk) for dataset in DATASETS}
        return self._validate_data(data)

    def create_snapshot(self) -> Path:
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        snapshot = self.backup_dir / f"copia_excel_{stamp}"
        snapshot.mkdir(parents=True, exist_ok=True)
        for dataset in DATASETS:
            shutil.copy2(self.path_for(dataset), snapshot / self.path_for(dataset).name)
        return snapshot
