from __future__ import annotations
from datetime import timedelta
from functools import wraps
from io import StringIO
import csv
import json
import uuid
from flask import Blueprint, abort, flash, g, redirect, render_template, request, Response, url_for
from auth_store import ROLE_ADMIN, ROLE_CHIEF, ROLE_COLLABORATOR, create_user, get_user_by_id, list_users, audit_event, reset_password, set_user_status, unlock_user
from permisos import MODULES
from . import db

bp=Blueprint('reportes',__name__,url_prefix='/reportes')

@bp.before_request
def chief_reporting_permissions():
    user=getattr(g,'current_user',None) or {}
    if not (user.get('is_chief') or user.get('role')==ROLE_CHIEF): return
    if request.endpoint=='reportes.index': return redirect(url_for('reportes.dashboard'))
    if request.endpoint in {'reportes.submit','reportes.plan_personal','reportes.cancel_personal',
                            'reportes.complete_assigned_task','reportes.my_excel'}:
        abort(403,description='El Jefe de Diseño administra y revisa las tareas del equipo; no registra tareas personales.')


def manager_required(view):
    @wraps(view)
    def wrapped(*args,**kwargs):
        if not (g.current_user.get('is_admin') or g.current_user.get('is_chief')): abort(403)
        return view(*args,**kwargs)
    return wrapped

def _audit(action,detail):
    audit_event(action,'REPORTES',detail,'OK',g.current_user,request.remote_addr or '')
    g.audit_logged=True

def _users():
    actor=g.current_user; rows=[]
    for user in list_users():
        if user['role']!=ROLE_COLLABORATOR: continue
        profile=db.get_profile(user['id'])
        if db.can_manage(actor,profile): rows.append({**user,'profile':profile})
    return rows

def _configured_users(): return [u for u in _users() if u['profile']]

def _range():
    try:
        start=db.parse_day(request.args.get('desde') or db.today().isoformat())
        end=db.parse_day(request.args.get('hasta') or start.isoformat())
        if end<start or (end-start).days>30 or end>db.today(): raise ValueError('Seleccione hasta 31 días, sin fechas futuras.')
        return start,end
    except ValueError as e: abort(400,description=str(e))

def _render(template,**context):
    return render_template(template,reportes_page=True,admin_page=True,modules=MODULES,
      today=db.today().isoformat(),time_label=db.time_label,duration_label=db.duration_label,**context)

@bp.get('/')
def index():
    day=db.today()
    if request.args.get('fecha'):
        try:
            day=db.parse_day(request.args['fecha'])
            if day>db.today()+timedelta(days=31): raise ValueError('Puede programar hasta 31 días hacia adelante.')
        except ValueError as e: abort(400,description=str(e))
    user=g.current_user
    db.sync_project_activities()
    return _render('reportes/mis_reportes.html',profile=db.get_profile(user['id']),
      summary=db.day_summary(user['id'],day),selected_day=day.isoformat(),
      reports=_my_report_rows(user['id'],day),assignments=_my_assignment_options(user['id'],day),
      pending_tasks=db.personal_tasks(user['id'],day),analysis=_my_analysis(user['id']) if request.args.get('analizar') else None)

@bp.post('/enviar')
def submit():
    try:
        db.sync_project_activities()
        report_id=db.submit_report(request.form,g.current_user)
        _audit('ENVIAR REPORTE',report_id)
        flash('Tarea completada y enviada al Jefe de Diseño.','ok')
    except ValueError as e: flash(str(e),'error')
    return redirect(url_for('reportes.index',fecha=request.form.get('day') or db.today().isoformat()))

@bp.get('/equipo')
@manager_required
def team():
    from proyectos_diseno.routes import storage
    users=_users(); selected=None
    user_id=request.args.get('usuario','')
    if user_id:
        selected=next((u for u in users if u['id']==user_id),None)
        if not selected: abort(403)
    board=storage.read('designers')
    selected_designer=request.args.get('disenador','')
    board_designer=next((d for d in board if d['id']==selected_designer),None)
    if selected_designer:
        if not board_designer: abort(404)
        linked=next((u for u in users if u['profile'] and u['profile']['designer_id']==selected_designer),None)
        occupied=next((p for p in db.profiles() if p['designer_id']==selected_designer),None)
        if occupied and not db.can_manage(g.current_user,occupied): abort(403)
        if linked: selected=linked
    elif selected and selected['profile']:
        board_designer=next((d for d in board if d['id']==selected['profile']['designer_id']),None)
    unlinked=[d for d in board if not any(p['designer_id']==d['id'] for p in db.profiles())]
    return _render('reportes/equipo.html',users=users,selected=selected,board_designer=board_designer,unlinked=unlinked,
      designers=storage.read('designers'),chiefs=[u for u in list_users() if u['role']==ROLE_CHIEF and u['is_active']])

@bp.post('/equipo/guardar')
@manager_required
def save_team():
    from proyectos_diseno.routes import storage
    actor=g.current_user; user_id=request.form.get('user_id','')
    chosen_role=request.form.get('role') or ROLE_COLLABORATOR
    if chosen_role not in {ROLE_COLLABORATOR,ROLE_CHIEF,ROLE_ADMIN}: abort(400)
    if chosen_role!=ROLE_COLLABORATOR and not actor.get('is_admin'): abort(403)
    raw={'start':request.form.get('start'),'end':request.form.get('end'),'breaks':request.form.get('breaks',''),
         'minimum_reports':0,'weekdays':request.form.getlist('weekdays'),
         'modules':request.form.getlist('modules')}
    try:
        config=db.validate_config(raw)
        db.finite_number(request.form.get('monthly_cost') or 0,'Costo mensual',0,1000000000)
        manager_id=actor['id'] if actor.get('is_chief') else request.form.get('manager_id','')
        manager=get_user_by_id(manager_id)
        if not manager or not manager['is_active'] or manager['role'] not in {ROLE_CHIEF,ROLE_ADMIN}:
            raise ValueError('Seleccione un Jefe de Diseño activo (o el Desarrollador responsable).')
        old=db.get_profile(user_id) if user_id else None
        if not db.can_manage(actor,old): abort(403)
        if old and old['version']!=int(request.form.get('version','0')): raise ValueError('El perfil cambió. Recargue antes de guardar.')
        designer_id=request.form.get('designer_id','')
        linked=storage.get('designers',designer_id) if designer_id else None
        if designer_id and not linked and not (old and designer_id==old['designer_id']): raise ValueError('Seleccione un diseñador existente válido.')
        if not designer_id:
            existing_user=get_user_by_id(user_id) if user_id else None
            name=(existing_user or {}).get('name') or request.form.get('name','')
            matches=[d for d in storage.read('designers') if str(d.get('nombre','')).strip().casefold()==str(name).strip().casefold()]
            if len(matches)>1: raise ValueError('Seleccione el diseñador existente: hay varias coincidencias de nombre.')
            if len(matches)==1: designer_id=matches[0]['id']; linked=matches[0]
        if designer_id and any(p['designer_id']==designer_id and p['user_id']!=user_id for p in db.profiles()):
            raise ValueError('Ese diseñador del tablero ya tiene otra cuenta vinculada.')
        if old and old['designer_id'] and designer_id!=old['designer_id']:
            raise ValueError('Conserve el vínculo actual del perfil con el diseñador del tablero.')
        if not user_id and chosen_role!=ROLE_COLLABORATOR and not linked:
            raise ValueError('Seleccione un diseñador existente para crear su cuenta con otro rol.')
        # Comprueba primero que el catálogo del tablero es accesible.
        storage.read('designers')
        if user_id:
            target=get_user_by_id(user_id)
            if not target or target['role']!=ROLE_COLLABORATOR: raise ValueError('Solo se configuran perfiles de Diseñador.')
        else:
            target=create_user(request.form.get('username'),request.form.get('name'),request.form.get('password'),chosen_role,actor['username'])
            user_id=target['id']
        if not designer_id:
            matches=[d for d in storage.read('designers') if str(d.get('nombre','')).strip().casefold()==target['name'].strip().casefold()]
            if len(matches)==1:
                designer_id=matches[0]['id']; linked=matches[0]
            elif len(matches)>1: raise ValueError('Hay diseñadores existentes con el mismo nombre. Seleccione el registro que corresponde.')
            else: designer_id=str(uuid.uuid4())
        if any(p['designer_id']==designer_id and p['user_id']!=user_id for p in db.profiles()):
            raise ValueError('El diseñador ya está vinculado a otra cuenta; no se creará un duplicado.')
        profile=db.save_profile(user_id,manager_id,designer_id,config,actor,request.form.get('version'))
        designer=dict(linked or {'id':designer_id,'nombre':target['name'],'color':'#1268c9','costo_mensual_empresa':0,'activo':True})
        first_break=config['breaks'][0] if config['breaks'] else ['','']
        designer.update(nombre=target['name'],activo=target['is_active'],color=request.form.get('color') or designer.get('color','#1268c9'),
                        costo_mensual_empresa=db.finite_number(request.form.get('monthly_cost') or designer.get('costo_mensual_empresa',0),'Costo mensual',0,1000000000))
        designer.update(hora_entrada=config['start'],hora_salida=config['end'],almuerzo_inicio=first_break[0],almuerzo_fin=first_break[1],
                        descansos=json.dumps(config['breaks']),dias_laborales=json.dumps(config['weekdays']))
        try:
            storage.upsert('designers',designer)
        except OSError:
            flash('El perfil fue guardado, pero no se pudo sincronizar el tablero. Cierre los Excel abiertos y vuelva a guardar este perfil.','error')
            return redirect(url_for('admin_users') if target['role']!=ROLE_COLLABORATOR else url_for('reportes.team',usuario=user_id))
        _audit('CONFIGURAR DISEÑADOR',f"{target['username']} · horario, reportes y módulos")
        db.sync_project_activities()
        flash('Cuenta configurada y vinculada al diseñador existente. Sus permisos se aplican según su rol.','ok')
        if target['role']!=ROLE_COLLABORATOR:
            return redirect(url_for('admin_users'))
    except PermissionError: abort(403)
    except ValueError as e: flash(str(e),'error')
    except OSError:
        flash('No se pudo sincronizar el diseñador con el tablero. Cierre los Excel abiertos y vuelva a guardar el perfil. La cuenta y el perfil se conservan.','error')
    return redirect(url_for('reportes.team',usuario=user_id) if user_id else url_for('reportes.team'))

@bp.post('/equipo/<user_id>/cuenta')
@manager_required
def account(user_id):
    target=get_user_by_id(user_id); profile=db.get_profile(user_id)
    if not target or target['role']!=ROLE_COLLABORATOR or not db.can_manage(g.current_user,profile): abort(403)
    action=request.form.get('action')
    try:
        if action=='password': reset_password(user_id,request.form.get('password'))
        elif action=='unlock': unlock_user(user_id)
        elif action=='status': set_user_status(user_id,request.form.get('status'))
        else: raise ValueError('Acción inválida.')
        _audit('ADMINISTRAR CUENTA DISEÑADOR',f'{target["username"]} · {action}')
        flash('Cuenta actualizada.','ok')
    except ValueError as e: flash(str(e),'error')
    return redirect(url_for('reportes.team',usuario=user_id))

@bp.get('/asignaciones')
def tasks():
    db.sync_project_activities()
    manager=g.current_user.get('is_admin') or g.current_user.get('is_chief')
    users=_configured_users() if manager else [{'id':g.current_user['id'],'name':g.current_user['name']}]
    items=[task for task in db.assignments([u['id'] for u in users]) if not task['source_id']]
    names={u['id']:u['name'] for u in users}
    selected=None
    if request.args.get('editar'):
        selected=next((t for t in items if t['id']==request.args['editar']),None)
        if not selected or not manager: abort(403)
        if selected.get('source_id'): return redirect(url_for('proyectos_diseno.index',view='projects',project_id=selected['project_id']))
    return _render('reportes/asignaciones.html',users=users,assignments=items,names=names,selected=selected,manager=manager)

@bp.post('/asignaciones/guardar')
@manager_required
def save_task():
    try:
        user=get_user_by_id(request.form.get('user_id'))
        if not user or not user['is_active'] or user['role']!=ROLE_COLLABORATOR: raise ValueError('Seleccione un diseñador activo.')
        task_id=db.save_assignment(request.form,g.current_user,request.form.get('id') or None)
        _audit('GUARDAR ASIGNACIÓN',task_id); flash('Asignación guardada.','ok')
    except PermissionError: abort(403)
    except ValueError as e: flash(str(e),'error')
    return redirect(url_for('reportes.tasks'))

@bp.get('/tablero')
@manager_required
def dashboard():
    db.sync_project_activities()
    start,end=_range(); users=_configured_users(); filter_user=request.args.get('usuario','')
    if filter_user:
        if not any(u['id']==filter_user for u in users): abort(403)
        users=[u for u in users if u['id']==filter_user]
    ids=[u['id'] for u in users]; names={u['id']:u['name'] for u in users}
    days=[start+timedelta(days=i) for i in range((end-start).days+1)]
    summaries=[]
    for user in users:
        rows=[db.day_summary(user['id'],d) for d in days]
        summaries.append({'user':user,'days':rows,'expected':sum(r['expected'] for r in rows),'actual':sum(r['actual'] for r in rows),
          'pending_days':sum(r['status']=='PENDIENTE' for r in rows),'missing_reports':sum(r['missing_reports'] for r in rows if r['status']=='PENDIENTE')})
    tasks=[t for t in db.assignments(ids,end) if not t['source_id'] and db.parse_day(t['start_date'])<=end]
    for task in tasks:
        if task['status']=='CANCELADA': state='CANCELADA'
        elif task['progress']>=100: state='COMPLETADA'
        else:
            due=db.parse_day(task['due_date']);profile=db.get_profile(task['user_id'])
            closed=due<end or (due==end and (end<db.today() or db.now().hour*60+db.now().minute>=db.time_minute(profile['end'])))
            state='ATRASADA' if closed else 'PENDIENTE'
        task['task_state']=state
    personal=[]
    for user in users:
        for day in days:
            for task in db.personal_tasks(user['id'],day):
                task['task_state']='POR CORREGIR' if task.get('report_status')=='DEVUELTO' else task['status']
                task['net_minutes']=db.net_minutes(task['start_minute'],task['end_minute'],db.day_summary(user['id'],day)['profile'] or user['profile'])
                personal.append(task)
    reports=[{**r,'metadata':json.loads(r['details'])} for r in db.list_reports(ids,start,end)]
    for summary in summaries:
        summary['count']=sum(d['count'] for d in summary['days'])
        summary['coverage']=round(summary['actual']/summary['expected']*100,1) if summary['expected'] else None
        summary['personal_pending']=sum(t['user_id']==summary['user']['id'] and t['task_state'] in {'PENDIENTE','POR CORREGIR'} for t in personal)
        summary['assigned_pending']=sum(t['user_id']==summary['user']['id'] and t['task_state'] in {'PENDIENTE','ATRASADA'} for t in tasks)
    return _render('reportes/tablero.html',users=_configured_users(),selected_user=filter_user,start=start.isoformat(),end=end.isoformat(),
      summaries=summaries,assignments=tasks,personal_tasks=personal,reports=reports,names=names,
      metrics={'expected':sum(s['expected'] for s in summaries),'actual':sum(s['actual'] for s in summaries),
      'late':sum(t['task_state']=='ATRASADA' for t in tasks),'ahead':sum(t['task_state']=='COMPLETADA' and t['trend']=='TERMINADA ANTES' for t in tasks),
      'pending':sum(t['task_state'] in {'PENDIENTE','ATRASADA'} for t in tasks)+sum(t['task_state'] in {'PENDIENTE','POR CORREGIR'} for t in personal),
      'unreviewed':sum(r['status']=='ENVIADO' for r in reports)})

@bp.post('/<report_id>/revisar')
@manager_required
def review(report_id):
    try:
        db.review_report(report_id,request.form,g.current_user)
        db.sync_reviewed_activity(report_id,g.current_user)
        _audit('REVISAR REPORTE',report_id); flash('Revisión registrada.','ok')
    except PermissionError: abort(403)
    except ValueError as e: flash(str(e),'error')
    except OSError: flash('La revisión quedó guardada, pero no se pudo actualizar el tablero. Cierre el Excel y vuelva a guardar la revisión para sincronizar.','error')
    return redirect(url_for('reportes.dashboard',desde=request.form.get('desde'),hasta=request.form.get('hasta'),usuario=request.form.get('usuario','')))

@bp.get('/exportar')
@manager_required
def export():
    start,end=_range(); users=_configured_users(); wanted=request.args.get('usuario','')
    if wanted:
        if not any(u['id']==wanted for u in users): abort(403)
        users=[u for u in users if u['id']==wanted]
    names={u['id']:u['name'] for u in users}; reports=db.list_reports(list(names),start,end)
    out=StringIO(); writer=csv.writer(out,delimiter=';')
    writer.writerow(['Fecha','Diseñador','Inicio','Fin','Horas netas','Tarea','Proyecto','Equipo','Número de equipo','Etapa','Centro de trabajo','Bloqueos','Estado','Comentario jefe'])
    for r in reports:
        metadata=json.loads(r['details'])
        values=[r['day'],names[r['user_id']],db.time_label(r['start_minute']),db.time_label(r['end_minute']),
        f"{r['minutes']/60:.2f}".replace('.',','),metadata.get('activity') or r['task_title'] or r['description'],metadata.get('project') or r['project_label'] or '',metadata.get('equipment',''),metadata.get('equipment_number',''),metadata.get('stage',''),metadata.get('workplace',''),r['blocker'],r['status'],r['review_note']]
        writer.writerow([("'"+str(v)) if str(v).lstrip().startswith(('=','+','-','@')) else v for v in values])
    _audit('EXPORTAR REPORTES',f'{start} a {end}')
    return Response('\ufeff'+out.getvalue(),mimetype='text/csv; charset=utf-8',headers={'Content-Disposition':f'attachment; filename="REPORTES_{start}_{end}.csv"'})


def _personal_range(user_id):
    with db.connection() as c:
        first=c.execute("SELECT MIN(day) FROM reports WHERE user_id=?",(user_id,)).fetchone()[0]
        configured=c.execute("SELECT MIN(effective_date) FROM profile_history WHERE user_id=?",(user_id,)).fetchone()[0]
    earliest=min([x for x in [first,configured] if x],default=db.today().isoformat())
    start=db.parse_day(request.args.get('desde') or earliest)
    end=db.parse_day(request.args.get('hasta') or db.today().isoformat())
    if end<start or end>db.today() or (end-start).days>3660:
        raise ValueError('Seleccione un día o rango válido, sin fechas futuras y hasta diez años.')
    return start,end

def _my_analysis(user_id):
    try: start,end=_personal_range(user_id)
    except ValueError as e: abort(400,description=str(e))
    result=db.personal_analysis(user_id,start,end)
    result.update(start=start.isoformat(),end=end.isoformat())
    return result

@bp.post('/mis-tareas/programar')
def plan_personal():
    try:
        task_id=db.plan_personal_task(request.form,g.current_user)
        _audit('PROGRAMAR TAREA PERSONAL',task_id)
        flash('Actividad programada. Complétala cuando hayas realizado el trabajo.','ok')
    except ValueError as e: flash(str(e),'error')
    return redirect(url_for('reportes.index',fecha=request.form.get('day') or db.today().isoformat()))

@bp.post('/mis-tareas/<task_id>/cancelar')
def cancel_personal(task_id):
    try:
        db.cancel_personal_task(task_id,g.current_user)
        _audit('CANCELAR TAREA PERSONAL',task_id);flash('Tarea cancelada.','ok')
    except ValueError as e: flash(str(e),'error')
    return redirect(url_for('reportes.index',fecha=request.form.get('day') or db.today().isoformat()))

@bp.get('/mis-reportes/excel')
def my_excel():
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.chart import LineChart,Reference
    from flask import send_file
    user=g.current_user
    try: start,end=_personal_range(user['id'])
    except ValueError as e: abort(400,description=str(e))
    analysis=db.personal_analysis(user['id'],start,end)
    reports=db.list_reports([user['id']],start,end)
    workbook=Workbook();summary=workbook.active;summary.title='Resumen'
    def safe(value):
        if isinstance(value,str) and value.lstrip().startswith(('=','+','-','@')): return "'"+value
        return value
    summary.append(['Mis reportes',safe(user['name'])])
    summary.append(['Desde',start]);summary.append(['Hasta',end])
    summary.append(['Horas registradas vigentes',analysis['actual']/60])
    summary.append(['Horas de jornada',analysis['expected']/60])
    summary.append(['Tareas completadas vigentes',analysis['count']])
    summary.append(['Días laborales del rango',analysis['working_days']])
    summary.append(['Tareas promedio por día laboral',analysis['average']])
    summary.append(['Cobertura de jornada',analysis['productivity']/100 if analysis['productivity'] is not None else None])
    summary.append(['Cobertura = horas reportadas / horas de jornada. Incluye reportes enviados y revisados; excluye devueltos.'])
    summary.append(['Fecha','Horas jornada','Horas reportadas','Tareas','Cobertura'])
    for d in analysis['days']:
        summary.append([db.parse_day(d['day']),d['expected']/60,d['actual']/60,d['count'],d['productivity']/100 if d['productivity'] is not None else None])
    for row in summary.iter_rows(min_row=12):
        row[0].number_format='dd/mm/yyyy';row[1].number_format=row[2].number_format='0.00';row[4].number_format='0.0%'
    summary['B2'].number_format=summary['B3'].number_format='dd/mm/yyyy';summary['B9'].number_format='0.0%'
    weekly=workbook.create_sheet('Semanal')
    weekly.append(['Semana inicio','Semana fin','Datos desde','Datos hasta','Horas jornada','Horas reportadas','Tareas','Promedio diario','Cobertura'])
    for w in analysis['weeks']:
        weekly.append([db.parse_day(w['start']),db.parse_day(w['end']),db.parse_day(w['period_start']),db.parse_day(w['period_end']),w['expected']/60,w['actual']/60,w['count'],w['average'],w['productivity']/100 if w['productivity'] is not None else None])
        for cell in weekly[weekly.max_row][:4]:cell.number_format='dd/mm/yyyy'
        weekly.cell(weekly.max_row,9).number_format='0.0%'
    trend=workbook.create_sheet('Tendencia')
    trend.append(['Fecha y hora final','Actividad','Horas de tarea','Horas acumuladas del día','Cobertura hasta la tarea','Horas laborales transcurridas'])
    for point in analysis['series']:
        trend.append([point['label'],safe(point['activity']),point['minutes']/60,point['accumulated']/60,point['productivity']/100 if point['productivity'] is not None else None,point['elapsed']/60])
        trend.cell(trend.max_row,5).number_format='0.0%'
    if analysis['series']:
        chart=LineChart();chart.title='Tendencia por tarea registrada';chart.y_axis.title='Cobertura hasta cada tarea';chart.x_axis.title='Fecha y hora final'
        chart.add_data(Reference(trend,min_col=5,min_row=1,max_row=1+len(analysis['series'])),titles_from_data=True)
        chart.set_categories(Reference(trend,min_col=1,min_row=2,max_row=1+len(analysis['series'])))
        chart.series[0].marker.symbol='circle';chart.series[0].marker.size=5;chart.series[0].smooth=False
        chart.display_blanks='gap';summary.add_chart(chart,'G2')
    summary['A10']='Tablas: cobertura diaria/semanal sobre la jornada completa. Gráfica: cobertura acumulada hasta cada tarea. Excluye devoluciones.'
    detail=workbook.create_sheet('Registro')
    detail.append(['PROYECTO','EQUIPO','NO. EQUIPO','ACTIVIDAD','ETAPA','FECHA','HORA INICIO','HORA FINAL','TOTAL HORAS NETAS','CENTRO DE TRABAJO','BLOQUEOS','AVANCE (%)','ESTADO','OBSERVACIÓN DEL JEFE'])
    for r in reversed(reports):
        metadata=json.loads(r['details']);values=[metadata.get('project') or r['project_label'] or '',metadata.get('equipment',''),metadata.get('equipment_number',''),metadata.get('activity') or r['task_title'] or r['description'] or 'Trabajo general',metadata.get('stage',''),db.parse_day(r['day']),r['start_minute']/1440,r['end_minute']/1440,r['minutes']/60,metadata.get('workplace',''),r['blocker'],r['progress'],r['status'],r['review_note']]
        detail.append([safe(v) for v in values]);n=detail.max_row
        detail.cell(n,6).number_format='dd/mm/yyyy';detail.cell(n,7).number_format=detail.cell(n,8).number_format='hh:mm';detail.cell(n,9).number_format='0.00'
    detail.freeze_panes='A2';detail.auto_filter.ref=detail.dimensions
    for sheet,header in [(summary,11),(detail,1),(weekly,1),(trend,1)]:
        for cell in sheet[header]:cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='164B7A')
        from openpyxl.utils import get_column_letter
        for col in range(1,sheet.max_column+1):sheet.column_dimensions[get_column_letter(col)].width=24
    summary.column_dimensions['A'].width=45
    detail.column_dimensions['K'].width=55
    for row in detail.iter_rows(min_row=2):
        for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
    buffer=BytesIO();workbook.save(buffer);buffer.seek(0)
    return send_file(buffer,as_attachment=True,download_name=f'Mis_reportes_{start}_{end}.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


def _my_report_rows(user_id,day):
    return [{**r,'metadata':json.loads(r['details'])} for r in db.list_reports([user_id],day,day)]

def _my_assignment_options(user_id,day):
    from proyectos_diseno.routes import storage
    stages={str(r['id']):r.get('etapa','') for r in storage.read('project_progress')}
    return [{**task,'stage_label':('ETAPA '+str(stages[task['source_id']])) if task['source_id'] in stages else ''}
            for task in db.assignments([user_id],day)]


@bp.post('/asignaciones/<task_id>/completar')
def complete_assigned_task(task_id):
    task=db.assignment(task_id)
    if not task or task['user_id']!=g.current_user['id'] or task['source_id']: abort(403)
    try:
        raw={ 'day':request.form.get('day'),'start':request.form.get('start'),'end':request.form.get('end'),
              'activity':task['title'],'description':task['title'],'assignment_id':task_id,'progress':100,
              'blocker':request.form.get('blocker',''),'project':task['project_label'],'_complete_assigned':True }
        report_id=db.submit_report(raw,g.current_user)
        _audit('COMPLETAR TAREA ASIGNADA',report_id)
        flash('Tarea completada. El trabajo está disponible para revisión del jefe.','ok')
    except ValueError as e: flash(str(e),'error')
    return redirect(url_for('reportes.tasks'))
