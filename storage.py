from __future__ import annotations

from contextvars import ContextVar
from datetime import datetime
from pathlib import Path
import hashlib
import json
import os
import re
import shutil
import threading
import uuid
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
BACKUP_MASTER_DIR = DATA_DIR / "BACKUP_MAESTRA"
USER_DATA_ROOT = DATA_DIR / "USUARIOS"
LEGACY_WORK_DIR = DATA_DIR / "TRABAJO_ACTUAL"
LEGACY_HIST_DIR = DATA_DIR / "HISTORICO"
LEGACY_AUTOSAVE_DIR = DATA_DIR / "AUTOGUARDADO"
LEGACY_HISTORY_FILE = DATA_DIR / "HISTORICO.xlsx"
MASTER_FILE = DATA_DIR / "LISTA_MAESTRA_FACTORY.xlsx"
CONFIG_FILE = DATA_DIR / "CONFIGURACION.xlsx"
HOLIDAYS_FILE = DATA_DIR / "FESTIVOS.xlsx"
AUTOSAVE_LIMIT = 10
_AUTOSAVE_LOCK = threading.RLock()
_EXCEL_LOCK = threading.RLock()
_CURRENT_USER_SCOPE = ContextVar("oga_current_user_scope", default=None)

BLUE = "0B64B4"
LIGHT_BLUE = "DDEBF7"
GREEN = "C6EFCE"
GREEN_TEXT = "006100"
RED = "FFC7CE"
RED_TEXT = "9C0006"
CYAN = "BDD7EE"
YELLOW = "FFF2CC"
GRAY = "E7E6E6"
WHITE = "FFFFFF"


def set_user_scope(scope):
    """Selecciona el espacio de datos del usuario para la solicitud actual."""
    if scope in (None, ""):
        _CURRENT_USER_SCOPE.set(None)
        return
    safe = re.sub(r"[^A-Za-z0-9_-]+", "_", str(scope)).strip("_")
    if not safe:
        raise ValueError("Identificador de usuario no válido.")
    _CURRENT_USER_SCOPE.set(safe[:80])


def get_user_scope(required=True):
    scope = _CURRENT_USER_SCOPE.get()
    if required and not scope:
        raise RuntimeError("No hay un espacio de usuario activo.")
    return scope


def get_user_data_dir():
    return USER_DATA_ROOT / get_user_scope()


def get_work_dir():
    return get_user_data_dir() / "TRABAJO_ACTUAL"


def get_hist_dir():
    return get_user_data_dir() / "HISTORICO"


def get_autosave_dir():
    return get_user_data_dir() / "AUTOGUARDADO"


def get_autosave_index():
    return get_autosave_dir() / "copias.json"


def get_history_file():
    return get_user_data_dir() / "HISTORICO.xlsx"


def ensure_dirs(include_user=True):
    paths = [DATA_DIR, BACKUP_MASTER_DIR, USER_DATA_ROOT]
    if include_user and get_user_scope(required=False):
        paths.extend((get_user_data_dir(), get_work_dir(), get_hist_dir(), get_autosave_dir()))
    for p in paths:
        p.mkdir(parents=True, exist_ok=True)


def migrate_legacy_data_to_current_user():
    """Copia una sola vez los datos de la V12 al primer administrador."""
    ensure_dirs()
    root = get_user_data_dir()
    marker = root / ".legacy_v12_migrated"
    if marker.exists():
        return False
    mappings = (
        (LEGACY_WORK_DIR, get_work_dir()),
        (LEGACY_HIST_DIR, get_hist_dir()),
        (LEGACY_AUTOSAVE_DIR, get_autosave_dir()),
    )
    copied = False
    for source, destination in mappings:
        if source.exists() and any(source.iterdir()):
            shutil.copytree(source, destination, dirs_exist_ok=True)
            copied = True
    if LEGACY_HISTORY_FILE.exists() and not get_history_file().exists():
        shutil.copy2(LEGACY_HISTORY_FILE, get_history_file())
        copied = True
    planos_root = root / "PLANOS"
    legacy_planos = (
        (BASE_DIR / "planos_uploads", planos_root / "CARGAS"),
        (BASE_DIR / "planos_revisados", planos_root / "REPORTES"),
    )
    for source, destination in legacy_planos:
        if source.exists() and any(source.iterdir()):
            shutil.copytree(source, destination, dirs_exist_ok=True)
            copied = True
    legacy_state = BASE_DIR / "planos_state.json"
    target_state = planos_root / "estado.json"
    if legacy_state.exists() and not target_state.exists():
        target_state.parent.mkdir(parents=True, exist_ok=True)
        try:
            legacy_data = json.loads(legacy_state.read_text(encoding="utf-8"))

            def relocate(value, key=""):
                if isinstance(value, dict):
                    return {item_key: relocate(item_value, item_key) for item_key, item_value in value.items()}
                if isinstance(value, list):
                    return [relocate(item, key) for item in value]
                if isinstance(value, str) and (key.endswith("_path") or key == "last_excel"):
                    filename = re.split(r"[\\/]", value)[-1]
                    destination = planos_root / ("REPORTES" if key == "last_excel" else "CARGAS")
                    return str(destination / filename)
                return value

            target_state.write_text(
                json.dumps(relocate(legacy_data), ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
        except (OSError, json.JSONDecodeError):
            shutil.copy2(legacy_state, target_state)
        copied = True
    marker.write_text(datetime.now().isoformat(timespec="seconds"), encoding="utf-8")
    return copied


def normalize_text(v):
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v).strip()).upper()


def clean_code(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    s = str(v).strip()
    if re.fullmatch(r"\d+\.0", s):
        return s[:-2]
    return s


def safe_float(v, default=0.0):
    if v is None or str(v).strip() == "":
        return default
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("$", "").replace(" ", "")
    # Colombian/US thousands handling
    if "," in s and "." in s:
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif s.count(".") > 1:
        s = s.replace(".", "")
    elif s.count(",") > 1:
        s = s.replace(",", "")
    elif "," in s:
        left, right = s.rsplit(",", 1)
        s = left + ("." + right if len(right) <= 2 else right)
    try:
        return float(s)
    except Exception:
        return default


def style_sheet(ws, currency_cols=None, status_col=None, warning_col=None):
    currency_cols = set(currency_cols or [])
    thin = Side(style="thin", color="D9E2F3")
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.border = Border(bottom=thin)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if c.column in currency_cols and isinstance(c.value, (int, float)):
                c.number_format = '$ #,##0'
        if status_col:
            idx = None
            for c in ws[1]:
                if normalize_text(c.value) == normalize_text(status_col):
                    idx = c.column
                    break
            if idx:
                val = normalize_text(row[idx - 1].value)
                fill = None
                if val in {"CORRECTO", "GESTIONADO", "COMPLETO"}:
                    fill = GREEN
                elif val in {"FALTANTE", "PENDIENTE", "CODIGO NO EXISTENTE", "CODIGO INVALIDO", "ERROR"}:
                    fill = RED
                elif val in {"SOBRANTE"}:
                    fill = CYAN
                if fill:
                    row[idx - 1].fill = PatternFill("solid", fgColor=fill)
        # Cualquier referencia sin costo debe quedar claramente marcada en rojo.
        if any(normalize_text(c.value) == "VALOR NO DISPONIBLE" for c in row):
            for c in row:
                c.font = Font(color=RED_TEXT, underline="single", bold=c.font.bold)
        if warning_col:
            idx = None
            for c in ws[1]:
                if normalize_text(c.value) == normalize_text(warning_col):
                    idx = c.column
                    break
            if idx and row[idx - 1].value:
                row[idx - 1].fill = PatternFill("solid", fgColor=YELLOW)
    for col in range(1, ws.max_column + 1):
        max_len = 0
        for c in ws.iter_cols(min_col=col, max_col=col, min_row=1, max_row=min(ws.max_row, 250)):
            for cell in c:
                max_len = max(max_len, len(str(cell.value or "")))
        ws.column_dimensions[get_column_letter(col)].width = min(max(max_len + 2, 11), 48)
    ws.row_dimensions[1].height = 24


def write_rows(path: Path, headers, rows, sheet_name="DATOS", currency_headers=None, status_col=None, warning_col=None):
    with _EXCEL_LOCK:
        path.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name
        ws.append(list(headers))
        for r in rows:
            ws.append([r.get(h, "") for h in headers])
        currency_cols = []
        for h in currency_headers or []:
            if h in headers:
                currency_cols.append(headers.index(h) + 1)
        style_sheet(ws, currency_cols=currency_cols, status_col=status_col, warning_col=warning_col)
        temporary = path.with_name(f".{path.stem}.{uuid.uuid4().hex}.tmp{path.suffix}")
        try:
            wb.save(temporary)
            os.replace(temporary, path)
        finally:
            wb.close()
            temporary.unlink(missing_ok=True)


def read_rows(path: Path, sheet_name=None):
    with _EXCEL_LOCK:
        if not path.exists():
            return []
        wb = load_workbook(path, data_only=True, read_only=True)
        try:
            ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
            it = ws.iter_rows(values_only=True)
            try:
                headers = [str(v).strip() if v is not None else "" for v in next(it)]
            except StopIteration:
                return []
            out = []
            for vals in it:
                if not any(v is not None and str(v).strip() != "" for v in vals):
                    continue
                out.append({headers[i]: vals[i] if i < len(vals) else "" for i in range(len(headers))})
            return out
        finally:
            wb.close()


def ensure_simple_workbooks():
    ensure_dirs(include_user=False)
    if not CONFIG_FILE.exists():
        write_rows(CONFIG_FILE, ["CLAVE", "VALOR"], [
            {"CLAVE": "DIAS_C.TEMPRANA", "VALOR": 30},
            {"CLAVE": "DIAS_PRODUCCION", "VALOR": 8},
            {"CLAVE": "DIAS_ENSAMBLE", "VALOR": 8},
        ], "CONFIGURACION")
    if not HOLIDAYS_FILE.exists():
        write_rows(HOLIDAYS_FILE, ["FECHA", "DESCRIPCION"], [], "FESTIVOS")
    if get_user_scope(required=False):
        ensure_dirs()
        history_file = get_history_file()
        if not history_file.exists():
            write_rows(history_file, ["ID", "NOMBRE", "FECHA", "ARCHIVO"], [], "HISTORICO")


def save_upload(file_storage, destination: Path):
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = destination.with_name(f".{destination.name}.{uuid.uuid4().hex}.uploading")
    try:
        file_storage.save(temporary)
        os.replace(temporary, destination)
    finally:
        temporary.unlink(missing_ok=True)


def clear_workspace():
    ensure_dirs()
    work_dir = get_work_dir()
    if work_dir.exists():
        shutil.rmtree(work_dir)
    work_dir.mkdir(parents=True, exist_ok=True)


def workspace_dirty():
    work_dir = get_work_dir()
    if not work_dir.exists():
        return False
    ignored = {".gitkeep"}
    for p in work_dir.rglob("*"):
        if p.is_file() and p.name not in ignored:
            return True
    return False


def workspace_signature():
    """Firma rápida del trabajo actual para evitar copias idénticas."""
    work_dir = get_work_dir()
    if not work_dir.exists():
        return ""
    digest = hashlib.sha256()
    for path in sorted((p for p in work_dir.rglob("*") if p.is_file()), key=lambda p: str(p.relative_to(work_dir))):
        try:
            stat = path.stat()
        except OSError:
            continue
        relative = str(path.relative_to(work_dir)).replace("\\", "/")
        digest.update(relative.encode("utf-8", errors="ignore"))
        digest.update(str(stat.st_size).encode("ascii"))
        digest.update(str(stat.st_mtime_ns).encode("ascii"))
    return digest.hexdigest()


def _read_autosave_index():
    autosave_index = get_autosave_index()
    autosave_dir = get_autosave_dir()
    if not autosave_index.exists():
        return []
    try:
        rows = json.loads(autosave_index.read_text(encoding="utf-8"))
        if not isinstance(rows, list):
            return []
    except Exception:
        return []
    valid = []
    for row in rows:
        if isinstance(row, dict) and (autosave_dir / str(row.get("archivo", ""))).is_file():
            valid.append(row)
    return sorted(valid, key=lambda row: str(row.get("fecha", "")), reverse=True)


def _write_autosave_index(rows):
    autosave_dir = get_autosave_dir()
    autosave_index = get_autosave_index()
    autosave_dir.mkdir(parents=True, exist_ok=True)
    temp = autosave_index.with_suffix(".tmp")
    temp.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")
    temp.replace(autosave_index)


def autosave_rows():
    with _AUTOSAVE_LOCK:
        return _read_autosave_index()


def create_autosave(reason="Autoguardado automático", force=False):
    """Crea una copia ZIP recuperable del trabajo actual y conserva las 10 últimas."""
    with _AUTOSAVE_LOCK:
        ensure_dirs()
        if not workspace_dirty():
            return None

        signature = workspace_signature()
        rows = _read_autosave_index()
        if not force and rows and rows[0].get("firma") == signature:
            return None

        work_dir = get_work_dir()
        autosave_dir = get_autosave_dir()
        now = datetime.now()
        autosave_id = now.strftime("%Y%m%d_%H%M%S_%f")
        safe_reason = re.sub(r"[^A-Za-z0-9_-]+", "_", reason.strip()).strip("_")[:45] or "AUTOGUARDADO"
        filename = f"{autosave_id}_{safe_reason}.zip"
        destination = autosave_dir / filename
        temporary = autosave_dir / f".{filename}.tmp"

        with zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as zf:
            for path in work_dir.rglob("*"):
                if path.is_file():
                    zf.write(path, path.relative_to(work_dir))
        temporary.replace(destination)

        meta = read_meta()
        row = {
            "id": autosave_id,
            "fecha": now.strftime("%Y-%m-%d %H:%M:%S"),
            "motivo": reason.strip() or "Autoguardado automático",
            "archivo": filename,
            "tamano": destination.stat().st_size,
            "firma": signature,
            "procesado": bool(meta.get("procesado")),
        }
        rows.insert(0, row)
        while len(rows) > AUTOSAVE_LIMIT:
            old = rows.pop()
            old_path = autosave_dir / str(old.get("archivo", ""))
            if old_path.exists():
                old_path.unlink()
        _write_autosave_index(rows)
        return row


def recover_autosave(autosave_id):
    with _AUTOSAVE_LOCK:
        target = next((row for row in _read_autosave_index() if str(row.get("id")) == str(autosave_id)), None)
        if not target:
            raise FileNotFoundError("La copia de seguridad no existe.")
        work_dir = get_work_dir()
        archive = get_autosave_dir() / str(target.get("archivo", ""))
        if not archive.exists():
            raise FileNotFoundError("No se encontró el archivo de la copia de seguridad.")

        with zipfile.ZipFile(archive, "r") as zf:
            work_root = work_dir.resolve()
            for member in zf.infolist():
                resolved = (work_dir / member.filename).resolve()
                if resolved != work_root and work_root not in resolved.parents:
                    raise ValueError("La copia contiene una ruta no permitida.")
            clear_workspace()
            zf.extractall(work_dir)
        return target


def delete_autosave(autosave_id):
    with _AUTOSAVE_LOCK:
        rows = _read_autosave_index()
        keep = []
        deleted = False
        for row in rows:
            if str(row.get("id")) == str(autosave_id):
                path = get_autosave_dir() / str(row.get("archivo", ""))
                if path.exists():
                    path.unlink()
                deleted = True
            else:
                keep.append(row)
        _write_autosave_index(keep)
        return deleted


def meta_path():
    return get_work_dir() / "META.json"


def read_meta():
    p = meta_path()
    if not p.exists():
        return {}
    try:
        return json.loads(p.read_text(encoding="utf-8"))
    except Exception:
        return {}


def write_meta(data):
    get_work_dir().mkdir(parents=True, exist_ok=True)
    path = meta_path()
    temporary = path.with_suffix(".tmp")
    temporary.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(path)


def set_meta(key, value):
    d = read_meta()
    d[key] = value
    write_meta(d)


def history_rows():
    rows = read_rows(get_history_file())
    return sorted(rows, key=lambda r: str(r.get("FECHA", "")), reverse=True)


def archive_workspace(name: str):
    ensure_simple_workbooks()
    if not workspace_dirty():
        raise ValueError("No hay una lista en proceso para archivar.")
    safe_name = re.sub(r"[^A-Za-z0-9_-]+", "_", name.strip()).strip("_") or "TRABAJO"
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    archive_id = stamp
    filename = f"{stamp}_{safe_name}.zip"
    work_dir = get_work_dir()
    hist_dir = get_hist_dir()
    zip_path = hist_dir / filename
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for p in work_dir.rglob("*"):
            if p.is_file():
                zf.write(p, p.relative_to(work_dir))
    rows = history_rows()
    rows.append({"ID": archive_id, "NOMBRE": name.strip(), "FECHA": datetime.now().strftime("%Y-%m-%d %H:%M:%S.%f"), "ARCHIVO": filename})
    rows = sorted(rows, key=lambda r: str(r.get("FECHA", "")), reverse=True)
    while len(rows) > 5:
        old = rows.pop()
        old_path = hist_dir / str(old.get("ARCHIVO", ""))
        if old_path.exists():
            old_path.unlink()
    write_rows(get_history_file(), ["ID", "NOMBRE", "FECHA", "ARCHIVO"], rows, "HISTORICO")
    clear_workspace()
    return archive_id


def recover_history(history_id: str):
    rows = history_rows()
    target = next((r for r in rows if str(r.get("ID")) == str(history_id)), None)
    if not target:
        raise FileNotFoundError("Histórico no encontrado")
    work_dir = get_work_dir()
    zp = get_hist_dir() / str(target.get("ARCHIVO", ""))
    if not zp.exists():
        raise FileNotFoundError("Archivo del histórico no encontrado")
    clear_workspace()
    with zipfile.ZipFile(zp, "r") as zf:
        work_root = work_dir.resolve()
        for member in zf.infolist():
            resolved = (work_dir / member.filename).resolve()
            if resolved != work_root and work_root not in resolved.parents:
                raise ValueError("El histórico contiene una ruta no permitida.")
        zf.extractall(work_dir)
    return target


def delete_history(history_id: str):
    rows = history_rows()
    keep = []
    for r in rows:
        if str(r.get("ID")) == str(history_id):
            p = get_hist_dir() / str(r.get("ARCHIVO", ""))
            if p.exists():
                p.unlink()
        else:
            keep.append(r)
    write_rows(get_history_file(), ["ID", "NOMBRE", "FECHA", "ARCHIVO"], keep, "HISTORICO")


def backup_master():
    if MASTER_FILE.exists():
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        dest = BACKUP_MASTER_DIR / f"LISTA_MAESTRA_FACTORY_{stamp}.xlsx"
        shutil.copy2(MASTER_FILE, dest)
        return dest
    return None
